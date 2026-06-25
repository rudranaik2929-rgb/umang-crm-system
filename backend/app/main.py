"""Umang Hometech LLP – Real Estate CRM Backend (Supabase Production)"""
from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File
from fastapi.responses import JSONResponse, PlainTextResponse
from starlette.middleware.cors import CORSMiddleware
from starlette.middleware.gzip import GZipMiddleware
from concurrent.futures import ThreadPoolExecutor
import asyncio
import threading
import uuid, logging, random, os, httpx, csv, io, openpyxl
import base64, hashlib, hmac, json, re, time
from io import BytesIO
from passlib.context import CryptContext
from pydantic import BaseModel
from typing import List, Optional, Dict, Any, Iterable, Tuple
from datetime import datetime, timezone, timedelta
from pathlib import Path
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parents[1] / ".env")
load_dotenv(Path(__file__).resolve().parents[2] / ".env")

# Password hashing configuration
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

def verify_password(plain_password, hashed_password):
    if not hashed_password: return False
    # Check if it looks like a bcrypt hash (starts with $2b$ or $2a$)
    if hashed_password.startswith("$2"):
        return pwd_context.verify(plain_password, hashed_password)
    # Fallback to plain text for legacy users
    return plain_password == hashed_password

def get_password_hash(password: str) -> str:
    return pwd_context.hash((password or "").strip())


def _password_db_fields(plain_password: str) -> Dict[str, str]:
    """Write both password_hash and legacy password column when present in Supabase."""
    hashed = get_password_hash(plain_password)
    return {"password_hash": hashed, "password": hashed}


def _verify_login_password_saved(user_id: str, plain_password: str) -> None:
    """Confirm the password was persisted and verifies — catches missing DB columns."""
    rows = sb_select(
        "users",
        {"user_id": f"eq.{user_id}", "select": "password_hash,password", "limit": "1"},
    )
    if not rows:
        raise HTTPException(status_code=500, detail="Login record missing after save.")
    stored = rows[0].get("password_hash") or rows[0].get("password")
    if not stored or not verify_password(plain_password.strip(), stored):
        raise HTTPException(
            status_code=500,
            detail=(
                "Password was not saved correctly. Run supabase/employee_login_migration.sql "
                "in the Supabase SQL Editor, then create the employee again."
            ),
        )


def _resolve_user_id_for_employee(employee: Dict[str, Any], email: str) -> Optional[str]:
    """Find users row by employee.user_id or by login email (fixes orphaned links)."""
    uid = employee.get("user_id")
    if uid:
        found = sb_select("users", {"user_id": f"eq.{uid}", "select": "user_id", "limit": "1"})
        if found:
            return uid
    norm = normalize_email(email)
    if not norm:
        return None
    by_email = sb_select("users", {"email": f"eq.{norm}", "select": "user_id", "limit": "1"})
    return by_email[0]["user_id"] if by_email else None

def normalize_email(email: Optional[str]) -> str:
    return (email or "").strip().lower()

app = FastAPI(title="Umang Hometech LLP CRM")

# Compress large JSON responses (e.g. full lead lists) to cut transfer time.
app.add_middleware(GZipMiddleware, minimum_size=1024)

RATE_LIMIT_WINDOW_SECONDS = int(os.environ.get("RATE_LIMIT_WINDOW_SECONDS", "60"))
RATE_LIMIT_MAX_REQUESTS = int(os.environ.get("RATE_LIMIT_MAX_REQUESTS", "240"))
WEBHOOK_RATE_LIMIT_MAX_REQUESTS = int(os.environ.get("WEBHOOK_RATE_LIMIT_MAX_REQUESTS", "120"))
RATE_LIMIT_ENABLED = os.environ.get("RATE_LIMIT_ENABLED", "true").lower() not in {"0", "false", "no"}
_RATE_LIMIT_BUCKETS: Dict[str, List[float]] = {}

# ---- Env config (must load before CORS middleware) ----
JWT_SECRET = os.environ.get("JWT_SECRET") or os.environ.get("SUPABASE_JWT_SECRET") or "change-me-in-production"
SESSION_TTL_DAYS = int(os.environ.get("SESSION_TTL_DAYS", "7"))
COOKIE_SECURE = os.environ.get("COOKIE_SECURE", "true").lower() not in {"0", "false", "no"}
_DEFAULT_CORS_ORIGINS = [
    "https://umanghometechllp.in",
    "https://www.umanghometechllp.in",
    "https://umang-home-tech.vercel.app",
    "http://localhost:8081",
    "http://localhost:19006",
    "http://localhost:3000",
]
_cors_env = (os.environ.get("CORS_ORIGINS") or "").strip()
CORS_ORIGINS = [o.strip() for o in _cors_env.split(",") if o.strip()] if _cors_env else _DEFAULT_CORS_ORIGINS


def _ensure_www_cors_pair(origins: List[str]) -> List[str]:
    """Vercel often serves www.umanghometechllp.in while env only lists the apex — allow both."""
    out = list(origins)
    pairs = [
        ("https://umanghometechllp.in", "https://www.umanghometechllp.in"),
        ("https://www.umanghometechllp.in", "https://umanghometechllp.in"),
    ]
    for a, b in pairs:
        if a in out and b not in out:
            out.append(b)
    return out


CORS_ORIGINS = _ensure_www_cors_pair(CORS_ORIGINS)
FRONTEND_URL = os.environ.get("FRONTEND_URL", "https://umanghometechllp.in").rstrip("/")
SUPABASE_URL = os.environ.get("SUPABASE_URL", "https://xlaiwmyyldxmuvopqomi.supabase.co")
SUPABASE_KEY = (
    os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    or os.environ.get("SUPABASE_KEY")
    or os.environ.get("SUPABASE_ANON_KEY")
    or ""
)

@app.middleware("http")
async def rate_limit_middleware(request: Request, call_next):
    if not RATE_LIMIT_ENABLED:
        return await call_next(request)

    path = request.url.path
    if path in {"/", "/debug-config"}:
        return await call_next(request)

    forwarded_for = request.headers.get("x-forwarded-for", "")
    client_ip = forwarded_for.split(",")[0].strip() or (request.client.host if request.client else "unknown")
    bucket_key = f"{client_ip}:{path}"
    now_ts = time.time()
    window_start = now_ts - RATE_LIMIT_WINDOW_SECONDS
    bucket = [ts for ts in _RATE_LIMIT_BUCKETS.get(bucket_key, []) if ts >= window_start]
    limit = WEBHOOK_RATE_LIMIT_MAX_REQUESTS if "webhook" in path else RATE_LIMIT_MAX_REQUESTS

    if len(bucket) >= limit:
        return JSONResponse(
            status_code=429,
            content={"detail": "Too many requests. Please retry shortly."},
            headers={"Retry-After": str(RATE_LIMIT_WINDOW_SECONDS)},
        )

    bucket.append(now_ts)
    _RATE_LIMIT_BUCKETS[bucket_key] = bucket
    return await call_next(request)

@app.get("/")
async def root_health():
    return {"status": "online", "message": "Umang Hometech LLP CRM Backend is running", "timestamp": datetime.now().isoformat()}

@app.get("/debug-config")
async def debug_config():
    return {
        "url_configured": bool(os.environ.get("SUPABASE_URL")),
        "service_role_configured": bool(os.environ.get("SUPABASE_SERVICE_ROLE_KEY")),
        "key_configured": bool(SUPABASE_KEY),
        "jwt_secret_ok": bool(JWT_SECRET) and JWT_SECRET not in ("change-me-in-production", "YOUR_RANDOM_SECRET"),
        "cors_origins": CORS_ORIGINS,
        "frontend_url": FRONTEND_URL,
    }

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
api_router = APIRouter(prefix="/api")

STAGES = ["new","assigned","positive","site_visit","booking","loan","registration","closed"]
ROLES = ["admin","manager","telecaller","site_visit","sales_executive","booking","loan","marketing"]

# ---- Integration Config (from .env) ----
INTERAKT_API_KEY = os.environ.get("INTERAKT_API_KEY", "")
OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "")
FACEBOOK_VERIFY_TOKEN = os.environ.get("FACEBOOK_VERIFY_TOKEN", "UMANGCRM123")
FACEBOOK_PAGE_ACCESS_TOKEN = os.environ.get("FACEBOOK_PAGE_ACCESS_TOKEN", "")
FACEBOOK_GRAPH_VERSION = os.environ.get("FACEBOOK_GRAPH_VERSION", "v20.0")
FACEBOOK_PAGE_ID = os.environ.get("FACEBOOK_PAGE_ID", "")
FACEBOOK_FORM_ID = os.environ.get("FACEBOOK_FORM_ID", "")
META_FAKE_LEADGEN_IDS = {"444444444444", "0", "test"}
META_FAKE_PAGE_IDS = {"0", "test"}
_FACEBOOK_PAGE_TOKEN_CACHE: Dict[str, Any] = {"tokens": {}, "fetched_at": 0.0}
FACEBOOK_AUTO_SYNC_ENABLED = os.environ.get("FACEBOOK_AUTO_SYNC_ENABLED", "true").lower() in ("1", "true", "yes")
FACEBOOK_AUTO_SYNC_INTERVAL_SEC = int(os.environ.get("FACEBOOK_AUTO_SYNC_INTERVAL_SEC", "900"))  # every 15 min
HOUSING_PROFILE_ID = os.environ.get("HOUSING_PROFILE_ID", "")
HOUSING_ENCRYPTION_KEY = os.environ.get("HOUSING_ENCRYPTION_KEY", "")
HOUSING_INTEGRATION_UUID = os.environ.get("HOUSING_INTEGRATION_UUID", "")
HOUSING_API_URL = os.environ.get("HOUSING_API_URL", "https://leads.housing.com/api/v0/get-builder-leads")
HOUSING_WEBHOOK_SECRET = os.environ.get("HOUSING_WEBHOOK_SECRET", HOUSING_INTEGRATION_UUID)
# Housing API pull: only import leads inside this window (not old backlog).
HOUSING_POLL_INITIAL_WINDOW_SEC = int(os.environ.get("HOUSING_POLL_INITIAL_WINDOW_SEC", "300"))  # 5 min
HOUSING_POLL_OVERLAP_SEC = int(os.environ.get("HOUSING_POLL_OVERLAP_SEC", "300"))  # 5 min overlap
HOUSING_MANUAL_DEFAULT_WINDOW_SEC = int(os.environ.get("HOUSING_MANUAL_DEFAULT_WINDOW_SEC", "7200"))  # 2 hours
HOUSING_API_MAX_RANGE_SEC = 2 * 86400  # Housing.com rejects ranges wider than 2 days
HOUSING_AUTO_SYNC_ENABLED = os.environ.get("HOUSING_AUTO_SYNC_ENABLED", "true").lower() in ("1", "true", "yes")
HOUSING_AUTO_SYNC_INTERVAL_SEC = int(os.environ.get("HOUSING_AUTO_SYNC_INTERVAL_SEC", "300"))  # every 5 min
CRON_SECRET = os.environ.get("CRON_SECRET", "")

def sb_headers():
    return {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
    }

def sb_url(table: str) -> str:
    return f"{SUPABASE_URL}/rest/v1/{table}"

_http = httpx.Client(timeout=20, limits=httpx.Limits(max_connections=20, max_keepalive_connections=10))

# Short-lived caches — cut repeated full-table Supabase reads on dashboard/login.
LEADS_CACHE_TTL_SEC = int(os.environ.get("LEADS_CACHE_TTL_SEC", "300"))
_USER_CACHE_TTL_SEC = int(os.environ.get("USER_CACHE_TTL_SEC", "300"))
_leads_cache: Dict[str, Any] = {"ts": 0.0, "select": "", "data": []}
_user_profile_cache: Dict[str, Dict[str, Any]] = {}
_employees_cache: Dict[str, Any] = {"ts": 0.0, "data": []}
_EMPLOYEES_CACHE_TTL = 60
_assignment_stats_cache: Dict[str, Any] = {"ts": 0.0, "data": None}
_employee_stats_cache: Dict[str, Any] = {"ts": 0.0, "data": None}
_dashboard_stats_cache: Dict[str, Any] = {"ts": 0.0, "data": None}
_STATS_CACHE_TTL_SEC = int(os.environ.get("STATS_CACHE_TTL_SEC", "120"))
_leads_cache_lock = threading.Lock()
_leads_cache_loading = threading.Event()

# One wide fetch shared by dashboard / assignment / employee stats (avoids 3× full-table reads).
LEADS_CANONICAL_SELECT = (
    "lead_id,name,phone,email,external_lead_id,source,stage,status,lead_type,"
    "priority,call_status,budget,location,property_type,assigned_to,assigned_at,last_employee_action_at,"
    "follow_up_at,created_at,updated_at,starred,brokerage_amount"
)

# Fields required for employee performance + missed-lead counts on team dashboard.
EMPLOYEE_WORKFLOW_LEAD_SELECT = (
    "lead_id,name,phone,email,source,assigned_to,status,stage,priority,call_status,"
    "follow_up_at,assigned_at,last_employee_action_at,created_at,updated_at,lead_type,"
    "budget,location,property_type,raw_payload"
)

def _b64url_encode(raw: bytes) -> str:
    return base64.urlsafe_b64encode(raw).rstrip(b"=").decode("ascii")

def _b64url_decode(value: str) -> bytes:
    padding = "=" * (-len(value) % 4)
    return base64.urlsafe_b64decode((value + padding).encode("ascii"))

def create_jwt(payload: Dict[str, Any], ttl_days: int = SESSION_TTL_DAYS) -> Tuple[str, str]:
    issued_at = int(now_utc().timestamp())
    expires_at = now_utc() + timedelta(days=ttl_days)
    body = {
        **payload,
        "iat": issued_at,
        "exp": int(expires_at.timestamp()),
        "iss": "umang-crm",
    }
    header = {"alg": "HS256", "typ": "JWT"}
    head = _b64url_encode(json.dumps(header, separators=(",", ":")).encode("utf-8"))
    data = _b64url_encode(json.dumps(body, separators=(",", ":"), default=str).encode("utf-8"))
    signing_input = f"{head}.{data}"
    signature = hmac.new(JWT_SECRET.encode("utf-8"), signing_input.encode("ascii"), hashlib.sha256).digest()
    return f"{signing_input}.{_b64url_encode(signature)}", expires_at.isoformat()

def decode_jwt(token: str) -> Optional[Dict[str, Any]]:
    try:
        head, body, signature = token.split(".", 2)
        signing_input = f"{head}.{body}"
        expected = hmac.new(JWT_SECRET.encode("utf-8"), signing_input.encode("ascii"), hashlib.sha256).digest()
        if not hmac.compare_digest(_b64url_decode(signature), expected):
            return None
        payload = json.loads(_b64url_decode(body).decode("utf-8"))
        if int(payload.get("exp", 0)) <= int(now_utc().timestamp()):
            return None
        return payload
    except Exception:
        return None

def sb_select(table, params=None):
    r = _http.get(sb_url(table), headers=sb_headers(), params=params or {})
    if r.status_code >= 400:
        logging.error(f"Supabase SELECT {table}: {r.status_code} {r.text[:300]}")
    return r.json() if r.status_code < 400 else []


# Shared thread pool to run independent Supabase reads concurrently. httpx.Client
# is thread-safe for issuing requests, so endpoints that need several unrelated
# tables (e.g. the dashboard) can fetch them in parallel instead of serially.
_read_pool = ThreadPoolExecutor(max_workers=8, thread_name_prefix="sb-read")


def sb_select_parallel(specs: Dict[str, Tuple[str, Optional[dict]]]) -> Dict[str, list]:
    """Run multiple sb_select calls concurrently. specs: key -> (table, params)."""
    futures = {
        key: _read_pool.submit(sb_select, table, params)
        for key, (table, params) in specs.items()
    }
    out: Dict[str, list] = {}
    for key, fut in futures.items():
        try:
            out[key] = fut.result()
        except Exception as exc:  # pragma: no cover - defensive
            logging.error(f"Supabase parallel SELECT {key}: {exc}")
            out[key] = []
    return out


def sb_insert(table, data):
    h = {**sb_headers(), "Prefer": "return=representation"}
    r = _http.post(sb_url(table), headers=h, json=data)
    if r.status_code >= 400:
        logging.error(f"Supabase INSERT {table}: {r.status_code} {r.text[:300]}")
        return None
    if not r.text or not r.text.strip():
        return data  # Return the input data if Supabase returns empty body
    try:
        rows = r.json()
        return rows[0] if isinstance(rows, list) and rows else rows
    except Exception:
        return data

def sb_update(table, pk_col, pk_val, data):
    h = {**sb_headers(), "Prefer": "return=representation"}
    r = _http.patch(f"{sb_url(table)}?{pk_col}=eq.{pk_val}", headers=h, json=data)
    if r.status_code >= 400:
        logging.error(f"Supabase UPDATE {table}: {r.status_code} {r.text[:300]}")
        return None
    rows = r.json()
    return rows[0] if isinstance(rows, list) and rows else rows

def sb_delete(table, pk_col, pk_val):
    r = _http.delete(f"{sb_url(table)}?{pk_col}=eq.{pk_val}", headers=sb_headers())
    return r.status_code < 400

def now_utc(): return datetime.now(timezone.utc)
def gen_id(p="id"): return f"{p}_{uuid.uuid4().hex[:12]}"
def model_payload(model: BaseModel) -> dict:
    data = {}
    for key, value in model.model_dump().items():
        if value is None:
            continue
        data[key] = value.isoformat() if isinstance(value, datetime) else value
    return data

def clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None

def normalize_phone(value: Any) -> str:
    raw = clean_text(value) or ""
    if not raw:
        return ""
    has_plus = raw.startswith("+")
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return ""
    if has_plus:
        return f"+{digits}"
    if len(digits) == 10:
        return f"+91{digits}"
    if len(digits) == 12 and digits.startswith("91"):
        return f"+{digits}"
    return digits

def pick_first(payload: Dict[str, Any], keys: Iterable[str]) -> Optional[Any]:
    for key in keys:
        if "." in key:
            cur: Any = payload
            for part in key.split("."):
                if not isinstance(cur, dict):
                    cur = None
                    break
                cur = cur.get(part)
            if clean_text(cur):
                return cur
            continue
        if clean_text(payload.get(key)):
            return payload.get(key)
    return None

def as_list_payload(value: Any) -> List[Dict[str, Any]]:
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    if isinstance(value, dict):
        for key in ("leads", "data", "items", "records"):
            nested = value.get(key)
            if isinstance(nested, list):
                return [item for item in nested if isinstance(item, dict)]
        return [value]
    return []

def is_brokerage_lead(source: Optional[str], payload: Optional[Dict[str, Any]] = None) -> bool:
    normalized = re.sub(r"[\s_.-]+", "", (source or "").strip().lower())
    if "broker" in normalized:
        return True
    payload = payload or {}
    if str(clean_text(payload.get("lead_type")) or "").lower() == "brokerage":
        return True
    if str(clean_text(payload.get("intake_type")) or "").lower() == "brokerage":
        return True
    return False

def is_broker_pool_lead(lead: Dict[str, Any]) -> bool:
    return lead.get("stage") == "broker" or str(lead.get("lead_type") or "").lower() == "brokerage"

def is_pipeline_lead(lead: Dict[str, Any]) -> bool:
    return not is_broker_pool_lead(lead)

def classify_lead_platform(source: Optional[str]) -> str:
    """Group lead sources into manual | housing | meta | other for dashboard breakdown."""
    if is_brokerage_lead(source):
        return "brokerage"
    normalized = re.sub(r"[\s_.-]+", "", (source or "").strip().lower())
    if "housing" in normalized:
        return "housing"
    if (
        normalized in {"facebook", "instagram", "meta", "fb"}
        or "facebook" in normalized
        or "instagram" in normalized
        or normalized.startswith("meta")
        or "metalead" in normalized
    ):
        return "meta"
    # "Database" = leads entered/owned by the team: manual add, walk-in,
    # referral, website enquiry, and Excel/CSV bulk uploads.
    MANUAL_SOURCES = {
        "manual", "manualentry", "walkin", "walkins", "referral", "direct", "call",
        "database", "db", "bulkimport", "bulk", "import", "excel", "excelimport",
        "csv", "csvimport", "spreadsheet", "offline", "website", "websiteenquiry",
        "webenquiry", "enquiry", "inquiry", "contactform",
    }
    if (
        normalized in MANUAL_SOURCES
        or normalized.startswith("manual")
        or normalized.startswith("bulkimport")
        or normalized.startswith("import")
        or normalized.startswith("excel")
        or normalized.startswith("csv")
    ):
        return "manual"
    return "other"

def lead_matches_platform(lead: Dict[str, Any], platform_key: str) -> bool:
    platform = classify_lead_platform(lead.get("source"))
    if platform == "brokerage":
        return platform_key == "brokerage"
    return platform == platform_key


def is_fake_meta_leadgen_id(leadgen_id: Optional[str]) -> bool:
    lid = clean_text(leadgen_id)
    return not lid or lid.lower() in {x.lower() for x in META_FAKE_LEADGEN_IDS}


def normalize_meta_page_id(*candidates: Optional[str]) -> Optional[str]:
    """Resolve a real Facebook Page id (Meta test webhooks send entry.id=0)."""
    for candidate in candidates:
        pid = clean_text(candidate)
        if pid and pid.lower() not in {x.lower() for x in META_FAKE_PAGE_IDS}:
            return pid
    env_pid = clean_text(FACEBOOK_PAGE_ID)
    if env_pid and env_pid.lower() not in {x.lower() for x in META_FAKE_PAGE_IDS}:
        return env_pid
    return None


def is_real_meta_lead(lead: Dict[str, Any]) -> bool:
    """A genuine Meta lead must come from a real leadgen id (not a Meta test id)
    and carry real contact details. Filters out the '444444444444' sample
    submissions Meta sends from the webhook test button."""
    external = str(lead.get("external_lead_id") or "").strip().lower()
    if external in {x.lower() for x in META_FAKE_LEADGEN_IDS}:
        return False
    phone = normalize_phone(lead.get("phone"))
    email = clean_text(lead.get("email"))
    if not phone and not email:
        return False
    # Reject obvious placeholder phone numbers (all same digit / too short).
    if phone and (len(set(phone)) <= 1 or len(phone) < 6):
        return bool(email)
    return True


def dedupe_leads(leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Collapse duplicate leads that share an external id, phone or email.
    Keeps the first (most recent, since callers pre-sort desc) occurrence."""
    seen: set = set()
    out: List[Dict[str, Any]] = []
    for lead in leads:
        keys = []
        ext = str(lead.get("external_lead_id") or "").strip().lower()
        if ext:
            keys.append(f"ext:{ext}")
        phone = normalize_phone(lead.get("phone"))
        if phone:
            keys.append(f"phone:{phone}")
        email = str(clean_text(lead.get("email")) or "").lower()
        if email:
            keys.append(f"email:{email}")
        if not keys:
            out.append(lead)
            continue
        if any(k in seen for k in keys):
            continue
        for k in keys:
            seen.add(k)
        out.append(lead)
    return out

def merge_leads_with_cache(db_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    cache_ids = {l.get("lead_id") for l in SESSION_CACHE["leads"]}
    return SESSION_CACHE["leads"] + [l for l in db_leads if l.get("lead_id") not in cache_ids]

PLATFORM_LABELS = {
    "manual": "Database",
    "housing": "Housing.com",
    "meta": "Meta (Facebook)",
    "other": "Other Sources",
    "brokerage": "Broker Pool",
}


def sb_select_all(table: str, params: Optional[dict] = None, page_size: int = 1000) -> List[Dict[str, Any]]:
    """Paginate through PostgREST default row cap so counts match opened lists."""
    base = dict(params or {})
    out: List[Dict[str, Any]] = []
    offset = 0
    while True:
        batch = sb_select(table, {**base, "limit": str(page_size), "offset": str(offset)})
        if not batch:
            break
        out.extend(batch)
        if len(batch) < page_size:
            break
        offset += page_size
    return out


def invalidate_leads_cache() -> None:
    with _leads_cache_lock:
        _leads_cache["ts"] = 0.0
    _assignment_stats_cache["ts"] = 0.0
    _employee_stats_cache["ts"] = 0.0
    _dashboard_stats_cache["ts"] = 0.0


def invalidate_employees_cache() -> None:
    _employees_cache["ts"] = 0.0


def _project_lead_fields(row: Dict[str, Any], select: str) -> Dict[str, Any]:
    if select in ("*", LEADS_CANONICAL_SELECT):
        return row
    keys = [k.strip() for k in select.split(",") if k.strip()]
    return {k: row.get(k) for k in keys}


def fetch_all_leads_merged(select: str = "*") -> List[Dict[str, Any]]:
    """Full leads table + session cache — one canonical fetch shared across dashboard endpoints."""
    now = time.time()

    def _from_cache() -> Optional[List[Dict[str, Any]]]:
        with _leads_cache_lock:
            if (
                _leads_cache["ts"]
                and (now - float(_leads_cache["ts"])) < LEADS_CACHE_TTL_SEC
                and _leads_cache.get("select") == LEADS_CANONICAL_SELECT
            ):
                cached = _leads_cache["data"]
                if select in ("*", LEADS_CANONICAL_SELECT):
                    return cached
                return [_project_lead_fields(l, select) for l in cached]
        return None

    hit = _from_cache()
    if hit is not None:
        return hit

    is_loader = False
    with _leads_cache_lock:
        if not _leads_cache_loading.is_set() and not _leads_cache.get("_loading"):
            _leads_cache["_loading"] = True
            _leads_cache_loading.clear()
            is_loader = True

    if not is_loader:
        _leads_cache_loading.wait(timeout=90)
        hit = _from_cache()
        if hit is not None:
            return hit

    try:
        db_leads = sb_select_all("leads", {"select": LEADS_CANONICAL_SELECT, "order": "created_at.desc"})
        merged = merge_leads_with_cache(db_leads)
        with _leads_cache_lock:
            _leads_cache.update({"ts": time.time(), "select": LEADS_CANONICAL_SELECT, "data": merged})
            _leads_cache.pop("_loading", None)
        _leads_cache_loading.set()
        if select in ("*", LEADS_CANONICAL_SELECT):
            return merged
        return [_project_lead_fields(l, select) for l in merged]
    except Exception:
        with _leads_cache_lock:
            _leads_cache.pop("_loading", None)
        _leads_cache_loading.set()
        raise


def clean_leads_for_platform_stats(leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Same pre-processing as compute_platform_breakdown so totals match platform lists."""
    cleaned: List[Dict[str, Any]] = []
    for lead in leads:
        if classify_lead_platform(lead.get("source")) == "meta" and not is_real_meta_lead(lead):
            continue
        if is_broker_pool_lead(lead) or classify_lead_platform(lead.get("source")) == "brokerage":
            continue
        cleaned.append(lead)
    return dedupe_leads(cleaned)


WORKSPACE_QUEUE_STAGES: Dict[str, List[str]] = {
    "telecaller": ["new", "assigned"],
    "sales_executive": ["new", "assigned", "positive"],
    "site_visit": ["new", "assigned", "positive"],
}


def workspace_queue_stages(role: Optional[str]) -> List[str]:
    key = (role or "telecaller").strip().lower()
    return WORKSPACE_QUEUE_STAGES.get(key, WORKSPACE_QUEUE_STAGES["telecaller"])


MISSED_LEAD_HOURS = 24


def parse_lead_ts(value: Any) -> Optional[datetime]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    s = str(value).strip()
    if not s:
        return None
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except ValueError:
        return None


def effective_assigned_at(lead: Dict[str, Any]) -> Optional[datetime]:
    """When assigned_at is missing on legacy rows, fall back like SQL backfill."""
    if not lead.get("assigned_to"):
        return None
    for key in ("assigned_at", "updated_at", "created_at"):
        ts = parse_lead_ts(lead.get(key))
        if ts:
            return ts
    return None


def is_missed_lead(lead: Dict[str, Any], hours: int = MISSED_LEAD_HOURS) -> bool:
    """Assigned lead with no employee workflow action within N hours."""
    if not lead.get("assigned_to"):
        return False
    if lead.get("status") != "active":
        return False
    if lead.get("stage") not in ("new", "assigned"):
        return False
    if clean_text(lead.get("call_status")):
        return False
    if lead.get("follow_up_at"):
        return False
    assigned_at = effective_assigned_at(lead)
    if not assigned_at:
        return False
    last_action = parse_lead_ts(lead.get("last_employee_action_at"))
    if last_action and assigned_at:
        # Ignore assignment noise (lead_assigned activity at same time as assign).
        if (last_action - assigned_at).total_seconds() < 120:
            last_action = None
    if last_action and last_action >= assigned_at:
        return False
    cutoff = now_utc() - timedelta(hours=hours)
    return assigned_at <= cutoff


def filter_employee_missed_leads(emp_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Leads assigned 24h+ ago with no employee status update — My Dashboard Missed Lead box."""
    rows = normalize_employee_leads(emp_leads)
    return dedupe_leads([l for l in rows if is_missed_lead(l)])


def filter_employee_my_queue_leads(emp_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """My Queue — assigned pipeline leads excluding missed (missed has its own box)."""
    rows = normalize_employee_leads(emp_leads)
    return dedupe_leads([l for l in rows if not is_missed_lead(l)])


def filter_employee_queue_leads(emp_leads: List[Dict[str, Any]], role: Optional[str]) -> List[Dict[str, Any]]:
    """Exact filter used by Telecaller / Sales Executive queue tabs."""
    stages = workspace_queue_stages(role)
    return dedupe_leads([
        l for l in emp_leads
        if l.get("status") == "active"
        and l.get("stage") in stages
        and not clean_text(l.get("call_status"))
        and not l.get("follow_up_at")
        and not is_missed_lead(l)
    ])


def filter_employee_follow_up_leads(emp_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Exact filter used by Follow Ups tab + employee follow-up KPI."""
    rows = normalize_employee_leads(emp_leads)
    return [
        l for l in rows
        if l.get("follow_up_at") and l.get("status") != "negative"
    ]


def _lead_priority(lead: Dict[str, Any]) -> str:
    return str(lead.get("priority") or "").strip().lower()


HANDOFF_BOOKING = "handoff_booking"
HANDOFF_LOAN = "handoff_loan"
HOT_PRIORITY = "hot"


def lead_ready_for_booking_queue(lead: Dict[str, Any]) -> bool:
    """Leads telecaller marks Hot, or sales marks handoff_booking, or legacy stage=booking."""
    if lead.get("status") == "negative":
        return False
    pr = _lead_priority(lead)
    if pr in (HANDOFF_BOOKING, HOT_PRIORITY):
        return True
    return lead.get("stage") == "booking"


def lead_ready_for_loan_queue(lead: Dict[str, Any]) -> bool:
    """Hot leads and loan handoffs — visible in Loan department New Application list."""
    if lead.get("status") == "negative":
        return False
    pr = _lead_priority(lead)
    if pr in (HANDOFF_LOAN, HOT_PRIORITY):
        return True
    return lead.get("stage") == "loan"


def normalize_employee_leads(emp_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Assigned leads for performance stats — pipeline only, deduped."""
    return dedupe_leads([l for l in emp_leads if is_pipeline_lead(l)])


def classify_employee_performance_metric(lead: Dict[str, Any]) -> Optional[str]:
    """One bucket per lead — same priority as assign-workspace inquiry filters."""
    if lead.get("status") == "negative":
        return "not_interested"
    if _lead_priority(lead) == "low_budget":
        return "low_budget"
    if clean_text(lead.get("call_status")):
        return "ringing"
    if lead.get("stage") in ["booking", "loan", "registration"] or _lead_priority(lead) in [HANDOFF_BOOKING, HANDOFF_LOAN]:
        return "booking_done"
    if _lead_priority(lead) == "hot":
        return "hot"
    if lead.get("stage") in ["site_visit", "positive"]:
        return "visited"
    return None


def filter_employee_metric_leads(emp_leads: List[Dict[str, Any]], metric_key: str) -> List[Dict[str, Any]]:
    """Single source of truth for employee performance boxes + drill-down lists."""
    rows = normalize_employee_leads(emp_leads)
    key = (metric_key or "").strip().lower()
    if key == "active":
        return [l for l in rows if l.get("status") != "negative" and l.get("stage") != "closed"]
    if key in ("follow_ups", "follow_up", "assigned_follow_ups"):
        return filter_employee_follow_up_leads(emp_leads)
    if key == "ringing":
        return [
            l for l in rows
            if l.get("status") != "negative" and clean_text(l.get("call_status"))
        ]
    if key in ("missed_leads", "missed"):
        return filter_employee_missed_leads(emp_leads)
    return [l for l in rows if classify_employee_performance_metric(l) == key]


def compute_employee_workflow_stats(emp_leads: List[Dict[str, Any]]) -> Dict[str, int]:
    """Dashboard employee performance boxes — counts match filter_employee_metric_leads."""
    return {
        "emp_active": len(filter_employee_metric_leads(emp_leads, "active")),
        "emp_hot": len(filter_employee_metric_leads(emp_leads, "hot")),
        "emp_visited": len(filter_employee_metric_leads(emp_leads, "visited")),
        "emp_not_interested": len(filter_employee_metric_leads(emp_leads, "not_interested")),
        "emp_booking_done": len(filter_employee_metric_leads(emp_leads, "booking_done")),
        "emp_low_budget": len(filter_employee_metric_leads(emp_leads, "low_budget")),
        "emp_ringing": len(filter_employee_metric_leads(emp_leads, "ringing")),
        "emp_follow_ups": len(filter_employee_follow_up_leads(emp_leads)),
        "emp_missed_leads": len(filter_employee_missed_leads(emp_leads)),
    }


def compute_employee_assignment_stats(emp_leads: List[Dict[str, Any]], role: Optional[str] = None) -> Dict[str, int]:
    """Per-employee assignment counts — dedupe within this employee's leads only."""
    rows = normalize_employee_leads(emp_leads)
    queue_rows = filter_employee_queue_leads(rows, role)
    follow_rows = filter_employee_follow_up_leads(rows)
    in_progress = sum(
        1 for l in rows
        if l.get("status") == "active"
        and l.get("stage") not in workspace_queue_stages(role)
        and l.get("stage") != "closed"
    )
    workflow = compute_employee_workflow_stats(rows)
    return {
        "assigned_total": len(rows),
        "assigned_queue": len(queue_rows),
        "assigned_in_progress": in_progress,
        "assigned_active": workflow["emp_active"],
        "assigned_completed": sum(1 for l in rows if l.get("stage") == "closed"),
        "assigned_positive": sum(
            1 for l in rows
            if l.get("stage") in ["positive", "site_visit", "booking", "loan", "registration"]
            and l.get("status") != "negative"
        ),
        "assigned_not_interested": workflow["emp_not_interested"],
        "assigned_new": sum(1 for l in rows if l.get("stage") in ["new", "assigned"] and l.get("status") != "negative"),
        "assigned_follow_ups": len(follow_rows),
        **workflow,
    }


CALL_STATUS_LABELS: Dict[str, str] = {
    "ringing": "Ringing",
    "out_of_service": "Out of Service",
    "call_back": "Call Back",
    "disconnect": "Disconnect",
}

NEGATIVE_PRIORITY_LABELS: Dict[str, str] = {
    "low_budget": "Low Budget",
    "other_location": "Other Location",
    "already_purchased": "Already Purchased",
}


def workflow_status_label(lead: Dict[str, Any]) -> str:
    """Human-readable status for lists — matches performance box buckets."""
    if is_missed_lead(lead):
        return "Missed Lead"
    if lead.get("status") == "negative":
        return NEGATIVE_PRIORITY_LABELS.get(_lead_priority(lead), "Not Interested")
    if _lead_priority(lead) == "low_budget":
        return "Low Budget"
    cs = clean_text(lead.get("call_status"))
    if cs:
        return CALL_STATUS_LABELS.get(cs, cs.replace("_", " ").title())
    if lead.get("stage") in ["booking", "loan", "registration"] or _lead_priority(lead) in [HANDOFF_BOOKING, HANDOFF_LOAN]:
        return "Booking Done"
    if _lead_priority(lead) == "hot":
        return "Hot"
    if lead.get("stage") in ["site_visit", "positive"]:
        return "Visited"
    if lead.get("follow_up_at") and lead.get("status") != "negative":
        return "Follow Up"
    if lead.get("stage") == "closed":
        return "Closed"
    if lead.get("stage") == "new":
        return "New Lead"
    if lead.get("stage") == "assigned":
        return "Assigned"
    if not lead.get("assigned_to") and lead.get("stage") in ["new", "assigned"]:
        return "New Lead"
    return "Active"


def enrich_lead_display_fields(lead: Dict[str, Any]) -> Dict[str, Any]:
    row = dict(lead)
    row["inquiry_status"] = classify_inquiry_status(lead)
    row["workflow_status"] = row["inquiry_status"]
    row["is_missed"] = is_missed_lead(lead)
    row["workflow_status_label"] = workflow_status_label(lead)
    return row


def enrich_leads_display_fields(leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [enrich_lead_display_fields(l) for l in leads]


def classify_inquiry_status(lead: Dict[str, Any]) -> str:
    """Primary inquiry bucket for manager assign workspace filters."""
    if lead.get("status") == "negative":
        return "not_interested"
    if _lead_priority(lead) == "low_budget":
        return "low_budget"
    if clean_text(lead.get("call_status")):
        return "ringing"
    if lead.get("stage") in ["booking", "loan", "registration"] or _lead_priority(lead) in [HANDOFF_BOOKING, HANDOFF_LOAN]:
        return "booked"
    if _lead_priority(lead) == "hot":
        return "hot"
    if lead.get("stage") in ["site_visit", "positive"]:
        return "visited"
    if not lead.get("assigned_to") and lead.get("stage") in ["new", "assigned"]:
        return "new"
    if lead.get("status") != "negative" and lead.get("stage") != "closed":
        return "active"
    return "active"


def lead_matches_inquiry_filter(lead: Dict[str, Any], inquiry_status: str) -> bool:
    key = (inquiry_status or "all").strip().lower()
    if key in ("", "all"):
        return True
    if key == "unassigned":
        return not lead.get("assigned_to")
    if key == "booked":
        return lead.get("stage") in ["booking", "loan", "registration"] or _lead_priority(lead) in [HANDOFF_BOOKING, HANDOFF_LOAN]
    return classify_inquiry_status(lead) == key


def lead_matches_assign_source(lead: Dict[str, Any], source_key: str) -> bool:
    key = (source_key or "all").strip().lower()
    if key in ("", "all"):
        return True
    platform = classify_lead_platform(lead.get("source"))
    if key == "manual":
        return platform == "manual"
    if key == "meta":
        return platform == "meta" and is_real_meta_lead(lead)
    if key == "housing":
        return platform == "housing"
    return platform == "other"


def lead_matches_search_query(lead: Dict[str, Any], q: str) -> bool:
    needle = (q or "").strip().lower()
    if not needle:
        return True
    hay = " ".join(
        str(lead.get(k) or "")
        for k in ("name", "phone", "email", "source", "location", "budget", "property_type", "notes")
    ).lower()
    return needle in hay


def filter_assign_workspace_leads(
    all_leads: List[Dict[str, Any]],
    inquiry_status: str = "all",
    source: str = "all",
    assigned_to: str = "all",
    q: str = "",
    location: str = "",
) -> List[Dict[str, Any]]:
    rows = [
        l for l in dedupe_leads(all_leads)
        if is_pipeline_lead(l)
        and lead_matches_inquiry_filter(l, inquiry_status)
        and lead_matches_assign_source(l, source)
        and lead_matches_search_query(l, q)
        and lead_matches_search_query(l, location)
    ]
    assignee = (assigned_to or "all").strip().lower()
    if assignee == "unassigned":
        rows = [l for l in rows if not l.get("assigned_to")]
    elif assignee not in ("", "all"):
        rows = [l for l in rows if l.get("assigned_to") == assigned_to]
    rows.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    return rows


def compute_assign_workspace_facets(all_leads: List[Dict[str, Any]]) -> Dict[str, Any]:
    pipeline = [l for l in dedupe_leads(all_leads) if is_pipeline_lead(l)]
    inquiry: Dict[str, int] = {k: 0 for k in ASSIGN_INQUIRY_STATUSES}
    source: Dict[str, int] = {k: 0 for k in ASSIGN_SOURCE_FILTERS}
    for lead in pipeline:
        inquiry["all"] += 1
        bucket = classify_inquiry_status(lead)
        if bucket in inquiry:
            inquiry[bucket] += 1
        if not lead.get("assigned_to"):
            inquiry["unassigned"] += 1
        plat = classify_lead_platform(lead.get("source"))
        source["all"] += 1
        if plat == "housing":
            source["housing"] += 1
        elif plat == "meta" and is_real_meta_lead(lead):
            source["meta"] += 1
        elif plat == "manual":
            source["manual"] += 1
        else:
            source["other"] += 1
    return {"inquiry_status": inquiry, "source": source}


def enrich_leads_with_employee_names(
    leads: List[Dict[str, Any]],
    employees: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    emp_map = {e.get("employee_id"): e.get("name") for e in employees if e.get("employee_id")}
    out: List[Dict[str, Any]] = []
    for lead in leads:
        row = enrich_lead_display_fields(lead)
        eid = lead.get("assigned_to")
        row["employee_name"] = emp_map.get(eid) if eid else None
        row["platform"] = classify_lead_platform(lead.get("source"))
        out.append(row)
    return out


def compute_unassigned_queue(all_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    queue = [l for l in all_leads if is_unassigned_new_lead(l)]
    return dedupe_leads(queue)


def is_unassigned_new_lead(lead: Dict[str, Any]) -> bool:
    """Incoming lead still waiting for manager assignment — not on any employee yet."""
    if not is_pipeline_lead(lead):
        return False
    if lead.get("status") == "negative":
        return False
    if clean_text(lead.get("assigned_to")):
        return False
    return lead.get("stage") in ("new", "assigned")


def is_new_lead_stage(lead: Dict[str, Any]) -> bool:
    """Dashboard / pipeline 'New Lead' column — unassigned only."""
    if not is_pipeline_lead(lead):
        return False
    if lead.get("status") == "negative":
        return False
    if clean_text(lead.get("assigned_to")):
        return False
    return lead.get("stage") == "new"


def actor_activity_keys(cu: User) -> set:
    """Activity rows may store employee_id or legacy user_id — match both."""
    keys = {cu.user_id}
    if cu.acting_as_employee_id:
        keys.add(cu.acting_as_employee_id)
    if cu.employee_id:
        keys.add(cu.employee_id)
    return {k for k in keys if k}


def resolve_employee_id(cu: User) -> Optional[str]:
    if cu.acting_as_employee_id:
        return cu.acting_as_employee_id
    if cu.employee_id:
        return cu.employee_id
    if cu.email:
        emps = sb_select("employees", {
            "email": f"eq.{cu.email}",
            "select": "employee_id",
            "limit": "1",
        })
        if emps and emps[0].get("employee_id"):
            return emps[0]["employee_id"]
    if cu.user_id:
        emps = sb_select("employees", {
            "user_id": f"eq.{cu.user_id}",
            "select": "employee_id",
            "limit": "1",
        })
        if emps and emps[0].get("employee_id"):
            return emps[0]["employee_id"]
    return None


def employee_assignee_ids(cu: User) -> List[str]:
    """IDs that may appear in leads.assigned_to for this login (employee_id, user_id, or legacy name)."""
    ids: List[str] = []
    for candidate in (resolve_employee_id(cu), cu.acting_as_employee_id, cu.employee_id, cu.user_id):
        if candidate and candidate not in ids:
            ids.append(candidate)
    emp_id = resolve_employee_id(cu) or cu.acting_as_employee_id or cu.employee_id
    if emp_id:
        emps = sb_select("employees", {
            "employee_id": f"eq.{emp_id}",
            "select": "name,user_id",
            "limit": "1",
        })
        if emps:
            uid = clean_text(emps[0].get("user_id"))
            if uid and uid not in ids:
                ids.append(uid)
            name = clean_text(emps[0].get("name"))
            if name and name not in ids:
                ids.append(name)
    return ids


def employee_record_match_values(employee: Dict[str, Any]) -> List[str]:
    """All values that may appear in leads.assigned_to for one employee row."""
    values: List[str] = []
    for candidate in (employee.get("employee_id"), employee.get("user_id"), clean_text(employee.get("name"))):
        if candidate and candidate not in values:
            values.append(candidate)
    return values


def lead_assigned_to_employee(lead: Dict[str, Any], employee: Dict[str, Any]) -> bool:
    assigned = clean_text(lead.get("assigned_to"))
    if not assigned:
        return False
    match_values = employee_record_match_values(employee)
    if assigned in match_values:
        return True
    emp_name = clean_text(employee.get("name"))
    return bool(emp_name and assigned.lower() == emp_name.lower())


def leads_for_employee_record(employee: Dict[str, Any], all_leads: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return [l for l in all_leads if lead_assigned_to_employee(l, employee)]


def fetch_employee_assigned_leads(cu: User, select: str) -> List[Dict[str, Any]]:
    """Leads assigned to the logged-in employee (My Dashboard / workspace)."""
    assignee_ids = employee_assignee_ids(cu)
    if not assignee_ids:
        return []
    merged: List[Dict[str, Any]] = []
    seen: set = set()
    for emp_id in assignee_ids:
        db_leads = sb_select_all("leads", {
            "select": select,
            "assigned_to": f"eq.{emp_id}",
            "order": "created_at.desc",
        })
        for lead in db_leads:
            lid = lead.get("lead_id")
            if lid and lid not in seen:
                seen.add(lid)
                merged.append(lead)
    cache_slice = [l for l in SESSION_CACHE["leads"] if l.get("assigned_to") in assignee_ids]
    cache_ids = {l.get("lead_id") for l in cache_slice}
    return cache_slice + [l for l in merged if l.get("lead_id") not in cache_ids]


def compute_platform_breakdown(leads: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Shared Database / housing / meta / other / broker counts for dashboard + modal."""
    platform_defs = [
        ("manual", PLATFORM_LABELS["manual"]),
        ("housing", PLATFORM_LABELS["housing"]),
        ("meta", PLATFORM_LABELS["meta"]),
    ]
    buckets = {
        key: {"platform": key, "label": label, "count": 0, "active": 0, "negative": 0, "sources": []}
        for key, label in platform_defs
    }
    buckets["other"] = {"platform": "other", "label": "Other Sources", "count": 0, "active": 0, "negative": 0, "sources": []}
    source_counts: Dict[str, int] = {}
    broker_pool_count = 0

    # Drop Meta test/placeholder submissions and collapse duplicates so the
    # headline counts match the per-platform lead lists shown in the UI.
    cleaned: List[Dict[str, Any]] = []
    for lead in leads:
        if classify_lead_platform(lead.get("source")) == "meta" and not is_real_meta_lead(lead):
            continue
        cleaned.append(lead)
    leads = dedupe_leads(cleaned)

    for lead in leads:
        raw_source = (lead.get("source") or "direct").strip()
        source_key = raw_source.lower()
        source_counts[source_key] = source_counts.get(source_key, 0) + 1
        platform = classify_lead_platform(raw_source)
        if platform == "brokerage":
            broker_pool_count += 1
            continue
        if is_broker_pool_lead(lead) and platform not in {"manual", "housing", "meta"}:
            broker_pool_count += 1
            continue
        bucket = buckets.get(platform) or buckets["other"]
        bucket["count"] += 1
        if lead.get("status") == "negative":
            bucket["negative"] += 1
        else:
            bucket["active"] += 1

    for source_key, count in sorted(source_counts.items(), key=lambda item: item[1], reverse=True):
        platform = classify_lead_platform(source_key)
        if platform in buckets and platform != "brokerage":
            buckets[platform]["sources"].append({"source": source_key, "count": count})

    platforms = [buckets["manual"], buckets["housing"], buckets["meta"]]
    if buckets["other"]["count"]:
        platforms.append(buckets["other"])

    pipeline_total = sum(p["count"] for p in platforms)
    return {
        "total": pipeline_total,
        "broker_pool": broker_pool_count,
        "platforms": platforms,
        "housing": buckets["housing"]["count"],
        "meta": buckets["meta"]["count"],
        "manual": buckets["manual"]["count"],
        "other": buckets["other"]["count"],
    }

def safe_json(value: Any) -> Any:
    try:
        json.dumps(value, default=str)
        return value
    except Exception:
        return {"raw": str(value)}

# ---- Pydantic Models ----
class User(BaseModel):
    user_id: str; email: str; name: str; picture: Optional[str]=None
    role: Optional[str]=None; acting_as_employee_id: Optional[str]=None; created_at: datetime
    employee_id: Optional[str]=None
    allowed_pages: Optional[List[str]] = None
    dashboard_type: Optional[str] = None
class RoleSet(BaseModel): role: str
class ActAs(BaseModel): employee_id: Optional[str]=None
class LeadCreatePublic(BaseModel):
    name: str; phone: str; email: Optional[str]=None; budget: Optional[str]=None
    location: Optional[str]=None; property_type: Optional[str]=None; notes: Optional[str]=None
    source: Optional[str]=None; starred: Optional[bool]=None
class LeadUpdate(BaseModel):
    stage: Optional[str]=None; status: Optional[str]=None; assigned_to: Optional[str]=None
    phone: Optional[str]=None; email: Optional[str]=None; budget: Optional[str]=None
    location: Optional[str]=None; property_type: Optional[str]=None; notes: Optional[str]=None
    source: Optional[str]=None; follow_up_at: Optional[datetime]=None; priority: Optional[str]=None
    starred: Optional[bool]=None
    lead_type: Optional[str]=None
    brokerage_amount: Optional[float]=None
    call_status: Optional[str]=None
class NoteCreate(BaseModel): text: str; type: str="call_note"
class SiteVisitCreate(BaseModel): lead_id: str; scheduled_at: datetime; assigned_to: Optional[str]=None
class SiteVisitUpdate(BaseModel):
    status: Optional[str]=None; feedback: Optional[str]=None; interested: Optional[bool]=None
    scheduled_at: Optional[datetime]=None; assigned_to: Optional[str]=None
    property_details: Optional[str]=None; interest_level: Optional[str]=None
class SiteVisitFollowUpCreate(BaseModel):
    visit_id: str; follow_up_date: str; follow_up_time: str; follow_up_day: str
    notes: Optional[str]=None
class LeadFollowUpCreate(BaseModel):
    follow_up_date: Optional[str]=None; follow_up_time: Optional[str]=None
    follow_up_day: Optional[str]=None; reason: Optional[str]=None; notes: Optional[str]=None
class BulkAssignRequest(BaseModel):
    lead_ids: List[str]
    assigned_to: str


class BulkLeadManageRequest(BaseModel):
    lead_ids: List[str]
    assigned_to: Optional[str] = None
    inquiry_action: Optional[str] = None
    stage: Optional[str] = None
    status: Optional[str] = None
    priority: Optional[str] = None
    call_status: Optional[str] = None
    reactivate: bool = True


ASSIGN_INQUIRY_STATUSES = [
    "all", "active", "new", "hot", "visited", "booked", "ringing",
    "not_interested", "low_budget", "unassigned",
]
ASSIGN_SOURCE_FILTERS = ["all", "housing", "meta", "manual", "other"]
INQUIRY_ACTION_PRESETS: Dict[str, Dict[str, Any]] = {
    "active": {"status": "active", "stage": "assigned"},
    "new": {"status": "active", "stage": "new"},
    "visited": {"status": "active", "stage": "site_visit"},
    "positive": {"status": "active", "stage": "positive"},
    "booked": {"status": "active", "stage": "booking"},
    "booking": {"status": "active", "stage": "booking"},
    "not_interested": {"status": "negative"},
    "ringing": {"status": "active", "call_status": "ringing"},
    "hot": {"status": "active", "stage": "positive", "priority": "hot"},
    "low_budget": {"status": "negative", "priority": "low_budget"},
}
class BookingCreate(BaseModel):
    lead_id: str; property_name: str; booking_amount: float=0; token_received: float=0
    unit_number: Optional[str]=None; tower: Optional[str]=None
    flat_cost: Optional[float]=None; agreement_value: Optional[float]=None
    stamp_duty: Optional[float]=None; registration_fees: Optional[float]=None
    gst: Optional[float]=None; society_charges: Optional[float]=None
    payment_status: Optional[str]=None; payment_progress: Optional[int]=None; booking_date: Optional[datetime]=None
    starred: Optional[bool]=None; completed_tasks: Optional[List[str]]=None
class BookingUpdate(BaseModel):
    token_received: Optional[float]=None; agreement_status: Optional[str]=None; status: Optional[str]=None
    property_name: Optional[str]=None; booking_amount: Optional[float]=None
    unit_number: Optional[str]=None; tower: Optional[str]=None
    flat_cost: Optional[float]=None; agreement_value: Optional[float]=None
    stamp_duty: Optional[float]=None; registration_fees: Optional[float]=None
    gst: Optional[float]=None; society_charges: Optional[float]=None
    brokerage_amount: Optional[float]=None
    payment_status: Optional[str]=None; payment_progress: Optional[int]=None; booking_date: Optional[datetime]=None
    starred: Optional[bool]=None; completed_tasks: Optional[List[str]]=None
class LoanCreate(BaseModel):
    lead_id: str; amount: float; bank_name: Optional[str]=None
    documents_status: Optional[str]=None; pending_documents: Optional[List[str]]=None
    starred: Optional[bool]=None
class LoanUpdate(BaseModel):
    bank_name: Optional[str]=None; application_status: Optional[str]=None; progress: Optional[int]=None; bank_stage: Optional[str]=None
    documents_status: Optional[str]=None; pending_documents: Optional[List[str]]=None; emi_eligible: Optional[float]=None
    amount: Optional[float]=None; starred: Optional[bool]=None
class EmployeeCreate(BaseModel):
    name: str; email: str; phone: Optional[str]=None; role: str; department: Optional[str]=None
    password: str
    allowed_pages: Optional[List[str]]=None
class EmployeeUpdate(BaseModel):
    name: Optional[str]=None; email: Optional[str]=None; phone: Optional[str]=None; role: Optional[str]=None; active: Optional[bool]=None
    allowed_pages: Optional[List[str]]=None; password: Optional[str]=None
class TemplateCreate(BaseModel): name: str; body: str
class CampaignCreate(BaseModel):
    name: str; template_id: Optional[str]=None; audience: str="all"; scheduled_at: Optional[datetime]=None
class HousingSyncRequest(BaseModel):
    start_date: Optional[int]=None
    end_date: Optional[int]=None
    allow_historical: bool = False

class FacebookImportRequest(BaseModel):
    page_id: Optional[str] = None
    form_id: Optional[str] = None
    days: int = 90
    limit: int = 500

# ---- Auth Helpers ----
LOCAL_SESSIONS = {}
SESSION_CACHE = {"leads": [], "bookings": [], "visits": [], "followups": [], "loans": [], "activities": [], "customers": [], "notifications": []}

DEMO_LEADS = []
DEMO_BOOKINGS = []
DEMO_VISITS = []
DEMO_LOANS = []

async def get_session_token(request: Request):
    t = request.cookies.get("session_token")
    if t: return t
    auth = request.headers.get("Authorization")
    if auth and auth.startswith("Bearer "): return auth[7:]
    return None

# Hardcoded user registry — map user_id to their user object for cold-start recovery
HARDCODED_USERS = {
    "user_admin001": {
        "user_id": "user_admin001", "email": "htshpatil13@gmail.com",
        "name": "Umang Admin", "role": "admin", "created_at": "2026-01-01T00:00:00+00:00",
    },
    "user_mukesh001": {
        "user_id": "user_mukesh001", "email": "mukesh@umang.com",
        "name": "Mukesh Sharma", "role": "telecaller",
        "employee_id": "emp_1b7760567ae6", "acting_as_employee_id": "emp_1b7760567ae6",
        "created_at": "2026-01-01T00:00:00+00:00",
    },
    "user_manager001": {
        "user_id": "user_manager001", "email": "rohitsingh241993@gmail.com",
        "name": "Rohit Singh", "role": "manager",
        "dashboard_type": "manager",
        "allowed_pages": [
            "dashboard", "my-dashboard", "pipeline", "assign-leads", "bookings", "loans",
            "integrations", "broker", "employees",
        ],
        "created_at": "2026-01-01T00:00:00+00:00",
    },
}

def public_user_payload(user: Dict[str, Any]) -> Dict[str, Any]:
    return User(**_finalize_user_session(dict(user))).model_dump(mode="json")

def _finalize_user_session(u: Dict[str, Any]) -> Dict[str, Any]:
    """Normalize role, sidebar pages, and dashboard landing for session/JWT payloads."""
    role = (u.get("role") or "telecaller").strip()
    u["role"] = role
    pages = u.get("allowed_pages")
    if pages is None:
        pages = _default_pages_for_role(role)
    else:
        pages = _coerce_allowed_pages(pages)
    if role == "manager":
        u["dashboard_type"] = "manager"
        if "dashboard" not in pages:
            pages = ["dashboard", *pages]
        if "my-dashboard" not in pages:
            pages = ["my-dashboard", *pages]
    elif role == "admin":
        u["dashboard_type"] = u.get("dashboard_type") or "admin"
        if "dashboard" not in pages:
            pages = ["dashboard", *pages]
    u["allowed_pages"] = pages
    return u

def issue_session(user: Dict[str, Any], response: Response):
    user = _finalize_user_session(dict(user))
    token, expires = create_jwt({
        "sub": user["user_id"],
        "email": user.get("email"),
        "role": user.get("role"),
        "name": user.get("name"),
        "employee_id": user.get("employee_id"),
        "allowed_pages": user.get("allowed_pages"),
        "dashboard_type": user.get("dashboard_type"),
    })
    LOCAL_SESSIONS[token] = {"user": user, "expires_at": expires}
    # Persist session in background — login response returns immediately (JWT is authoritative).
    def _persist_session():
        try:
            sb_insert("sessions", {
                "session_token": token,
                "user_id": user["user_id"],
                "created_at": now_utc().isoformat(),
                "expires_at": expires,
            })
        except Exception as exc:
            logging.warning("session persist failed: %s", exc)
    threading.Thread(target=_persist_session, daemon=True).start()
    response.set_cookie(
        key="session_token",
        value=token,
        max_age=SESSION_TTL_DAYS * 24 * 60 * 60,
        httponly=True,
        samesite="none",
        path="/",
        secure=COOKIE_SECURE,
    )
    return {
        "user": public_user_payload(user),
        "session_token": token,
        "access_token": token,
        "token_type": "bearer",
        "expires_at": expires,
    }

def ensure_roles(cu: User, allowed: Iterable[str]):
    if cu.role not in set(allowed) and cu.email != "htshpatil13@gmail.com":
        raise HTTPException(status_code=403, detail="You do not have permission for this action.")


def _can_manage_all_leads(cu: User) -> bool:
    return cu.role in ("admin", "manager") or cu.email in ("htshpatil13@gmail.com", "umang@admin")


def _can_reactivate_negative_lead(cu: User) -> bool:
    return _can_manage_all_leads(cu) or cu.role == "marketing"


def ensure_lead_edit_access(cu: User, lead: Dict[str, Any]) -> None:
    """Employees may only edit leads assigned to them; managers/admins edit any."""
    if _can_manage_all_leads(cu):
        return
    if cu.role == "marketing" and lead.get("status") == "negative":
        return
    assignee_ids = employee_assignee_ids(cu)
    assigned = clean_text(lead.get("assigned_to"))
    if assigned and assigned in assignee_ids:
        return
    emp_id = resolve_employee_id(cu) or cu.employee_id
    if emp_id:
        emps = sb_select("employees", {"employee_id": f"eq.{emp_id}", "select": "name", "limit": "1"})
        if emps and assigned.lower() == clean_text(emps[0].get("name")).lower():
            return
    raise HTTPException(status_code=403, detail="You can only update leads assigned to you.")


def fresh_lead_reset_fields(stage: str = "new") -> Dict[str, Any]:
    """Manager reset — e.g. not-interested from Emp1 → new for Emp2."""
    return {
        "status": "active",
        "stage": stage,
        "priority": None,
        "call_status": None,
        "follow_up_at": None,
        "last_employee_action_at": None,
    }


def inquiry_preset_to_patch(preset: Dict[str, Any], inquiry_action: Optional[str] = None) -> Dict[str, Any]:
    """Map manager bulk status action to lead patch (includes field clears for reset)."""
    action = (inquiry_action or "").strip().lower()
    if action == "new":
        return fresh_lead_reset_fields("new")
    patch: Dict[str, Any] = {}
    for key in ("stage", "status", "priority", "call_status", "follow_up_at"):
        if key in preset and preset[key] is not None:
            patch[key] = preset[key]
    if action == "active":
        patch.setdefault("status", "active")
        patch["call_status"] = None
        patch["follow_up_at"] = None
    elif action == "ringing":
        patch.setdefault("status", "active")
        patch.setdefault("call_status", "ringing")
        if patch.get("stage") is None:
            patch["stage"] = "assigned"
    elif action == "not_interested":
        patch.setdefault("status", "negative")
        patch["call_status"] = None
        patch["follow_up_at"] = None
    return patch


def apply_call_status_workflow(old_lead: Dict[str, Any], call_status: str) -> Dict[str, Any]:
    """Employee marks ringing / call-back — lead leaves queue and appears in Ringing box."""
    data: Dict[str, Any] = {
        "call_status": call_status,
        "status": "active",
        "updated_at": now_utc().isoformat(),
    }
    if old_lead.get("stage") in (None, "new"):
        data["stage"] = "assigned"
    return data


def ensure_owner_dashboard(cu: User) -> None:
    """Owner-only dashboard API (revenue + full admin controls)."""
    if cu.role == "manager":
        raise HTTPException(
            status_code=403,
            detail="Owner dashboard is for administrators only. Managers use the team Dashboard.",
        )
    if cu.role == "admin" or cu.email in ("htshpatil13@gmail.com", "umang@admin"):
        return
    raise HTTPException(status_code=403, detail="Owner dashboard is for administrators only.")


def ensure_main_dashboard(cu: User) -> None:
    """Main Dashboard bundle — administrators and managers."""
    if cu.role in ("admin", "manager"):
        return
    if cu.email in ("htshpatil13@gmail.com", "umang@admin"):
        return
    raise HTTPException(status_code=403, detail="Dashboard access denied.")


def _resolve_user_by_id(uid: str, expires_at: str) -> Dict[str, Any]:
    if uid in HARDCODED_USERS:
        return dict(HARDCODED_USERS[uid])
    cached = _user_profile_cache.get(uid)
    if cached and (time.time() - float(cached.get("_cached_at", 0))) < _USER_CACHE_TTL_SEC:
        return {k: v for k, v in cached.items() if k != "_cached_at"}
    users = sb_select("users", {"user_id": f"eq.{uid}", "select": "user_id,email,name,role,employee_id,acting_as_employee_id,allowed_pages,dashboard_type,created_at"})
    if not users:
        raise HTTPException(401, "User not found")
    u = users[0]
    if u.get("employee_id") and not u.get("acting_as_employee_id") and u.get("role") != "admin":
        u["acting_as_employee_id"] = u["employee_id"]
    if u.get("employee_id"):
        emps = sb_select("employees", {
            "employee_id": f"eq.{u['employee_id']}",
            "select": "active",
            "limit": "1",
        })
        if emps and emps[0].get("active") is False:
            raise HTTPException(401, "Account is disabled. Contact your manager.")
    _user_profile_cache[uid] = {**u, "_cached_at": time.time()}
    return _finalize_user_session(_hydrate_allowed_pages_from_employee(u))

def invalidate_sessions_for_user(user_id: str):
    """Force re-login for one user (e.g. after password reset). Never clears all sessions."""
    _user_profile_cache.pop(user_id, None)
    sb_delete("sessions", "user_id", user_id)
    for tok, sess in list(LOCAL_SESSIONS.items()):
        if (sess.get("user") or {}).get("user_id") == user_id:
            LOCAL_SESSIONS.pop(tok, None)

def user_from_jwt_payload(jwt_payload: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Build session user from JWT claims — no Supabase round-trip."""
    uid = jwt_payload.get("sub")
    if not uid:
        return None
    if uid in HARDCODED_USERS:
        return dict(HARDCODED_USERS[uid])
    role = jwt_payload.get("role")
    if not role:
        return None
    emp_id = jwt_payload.get("employee_id")
    u: Dict[str, Any] = {
        "user_id": uid,
        "email": jwt_payload.get("email"),
        "role": role,
        "name": jwt_payload.get("name") or "User",
        "employee_id": emp_id,
        "allowed_pages": jwt_payload.get("allowed_pages"),
        "dashboard_type": jwt_payload.get("dashboard_type"),
    }
    if emp_id and role != "admin":
        u["acting_as_employee_id"] = emp_id
    hydrated = _hydrate_allowed_pages_from_employee(u)
    return _finalize_user_session(hydrated)


async def get_current_user(request: Request) -> User:
    token = await get_session_token(request)
    if not token: raise HTTPException(401, "Not authenticated")

    jwt_payload = decode_jwt(token) if token.count(".") == 2 else None
    if token.count(".") == 2 and not jwt_payload:
        raise HTTPException(401, "Invalid token")
    
    # 1. Check in-memory cache first (fastest)
    if token in LOCAL_SESSIONS:
        sess = LOCAL_SESSIONS[token]
        if datetime.fromisoformat(sess["expires_at"].replace("Z","+00:00")) <= now_utc():
            del LOCAL_SESSIONS[token]
            raise HTTPException(401, "Session expired")
        u = _finalize_user_session(dict(sess["user"]))
        u = _hydrate_allowed_pages_from_employee(u)
        act_as = request.headers.get("X-Acting-As")
        if act_as:
            u["acting_as_employee_id"] = act_as
        return User(**u)

    # 2. Valid JWT — use token claims directly (fast path, no DB)
    if jwt_payload and jwt_payload.get("sub"):
        exp = datetime.fromtimestamp(int(jwt_payload["exp"]), tz=timezone.utc).isoformat()
        u = user_from_jwt_payload(jwt_payload)
        if u:
            u = _hydrate_allowed_pages_from_employee(u)
            LOCAL_SESSIONS[token] = {"user": u, "expires_at": exp}
            act_as = request.headers.get("X-Acting-As")
            if act_as:
                u = {**u, "acting_as_employee_id": act_as}
            return User(**u)
        uid = jwt_payload["sub"]
        u = _resolve_user_by_id(uid, exp)
        LOCAL_SESSIONS[token] = {"user": u, "expires_at": exp}
        act_as = request.headers.get("X-Acting-As")
        if act_as:
            u["acting_as_employee_id"] = act_as
        return User(**u)

    # 3. Legacy session row in Supabase (non-JWT tokens)
    rows = sb_select("sessions", {"session_token": f"eq.{token}", "select": "user_id,expires_at"})
    if rows:
        sess = rows[0]
        exp = sess.get("expires_at", "")
        if exp and datetime.fromisoformat(exp.replace("Z","+00:00")) <= now_utc():
            raise HTTPException(401, "Session expired")
        uid = sess["user_id"]
    else:
        raise HTTPException(401, "Invalid session")

    u = _resolve_user_by_id(uid, exp)
    LOCAL_SESSIONS[token] = {"user": u, "expires_at": exp}
    act_as = request.headers.get("X-Acting-As")
    if act_as:
        u["acting_as_employee_id"] = act_as
    return User(**u)

# ---- Auth Endpoints ----
@api_router.post("/auth/session")
async def auth_session(request: Request, response: Response):
    body = await request.json()
    email = normalize_email(body.get("email"))
    password = body.get("password") or ""
    
    # Hardcoded fallback for demo
    if email in ["umang@admin", "htshpatil13@gmail.com"] and password == "umang@admin":
        u = {
            "user_id": "user_admin001",
            "email": email,
            "name": "Umang Admin",
            "role": "admin",
            "created_at": now_utc().isoformat(),
        }
        return issue_session(u, response)

    # Hardcoded sample employee: Mukesh Sharma (telecaller)
    if email == "mukesh@umang.com" and password == "mukesh@123":
        u = {
            "user_id": "user_mukesh001",
            "email": email,
            "name": "Mukesh Sharma",
            "role": "telecaller",
            "employee_id": "emp_1b7760567ae6",
            "acting_as_employee_id": "emp_1b7760567ae6",
            "created_at": now_utc().isoformat(),
        }
        return issue_session(u, response)

    # Hardcoded manager: Rohit Singh
    if email == "rohitsingh241993@gmail.com" and password == "umang@manager":
        u = {
            "user_id": "user_manager001",
            "email": email,
            "name": "Rohit Singh",
            "role": "manager",
            "created_at": now_utc().isoformat(),
        }
        return issue_session(u, response)

    # Alias mapping: allow shorthand "umang@admin" to resolve to the real admin email
    EMAIL_ALIASES = {
        "umang@admin": "htshpatil13@gmail.com",
    }
    lookup_email = EMAIL_ALIASES.get(email, email)
    
    # Query real users table
    users = sb_select("users", {"email": f"eq.{lookup_email}", "select": "user_id,email,name,role,employee_id,acting_as_employee_id,password_hash,password,allowed_pages,dashboard_type,created_at"})
    if not users:
        raise HTTPException(401, "Invalid email or password")
    
    u = users[0]
    # Robust check for both column names (password_hash is the new standard)
    db_password = u.get("password_hash") or u.get("password")
    
    if not db_password or not verify_password(password, db_password): 
        raise HTTPException(401, "Invalid email or password")

    if u.get("employee_id"):
        emps = sb_select("employees", {
            "employee_id": f"eq.{u['employee_id']}",
            "select": "active",
            "limit": "1",
        })
        if emps and emps[0].get("active") is False:
            raise HTTPException(401, "Account is disabled. Contact your manager.")

    # Employees act as their own employee record so personal stats / lead
    # assignment resolve correctly without an explicit X-Acting-As header.
    if u.get("employee_id") and not u.get("acting_as_employee_id") and u.get("role") != "admin":
        u["acting_as_employee_id"] = u["employee_id"]

    return issue_session(u, response)

@api_router.get("/auth/me")
async def auth_me(cu: User = Depends(get_current_user)):
    return cu.model_dump(mode="json")

@api_router.post("/auth/logout")
async def auth_logout(request: Request, response: Response):
    t = await get_session_token(request)
    if t:
        sb_delete("sessions", "session_token", t)
        LOCAL_SESSIONS.pop(t, None)  # Clear in-memory cache
    response.delete_cookie("session_token", path="/")
    return {"ok": True}

@api_router.post("/auth/set-role")
async def auth_set_role(payload: RoleSet, cu: User = Depends(get_current_user)):
    if payload.role not in ROLES: raise HTTPException(400, "Invalid role")
    patch: Dict[str, Any] = {"role": payload.role, "dashboard_type": payload.role}
    if payload.role == "manager":
        patch["allowed_pages"] = _default_pages_for_role("manager")
    updated = sb_update("users", "user_id", cu.user_id, patch)
    if not updated: raise HTTPException(500, "Failed to update role")
    merged = {**cu.model_dump(mode="json"), **updated}
    return User(**_finalize_user_session(merged)).model_dump(mode="json")

@api_router.post("/auth/ping-location")
async def ping_location(request: Request, cu: User = Depends(get_current_user)):
    body = await request.json()
    lat, lng = body.get("lat"), body.get("lng")
    if lat is None or lng is None:
        return {"ok": False, "reason": "missing_coordinates"}

    emp_id = cu.acting_as_employee_id or cu.employee_id
    if not emp_id:
        return {"ok": False, "reason": "no_employee_linked"}

    update: Dict[str, Any] = {
        "last_lat": float(lat),
        "last_lng": float(lng),
        "last_seen_at": now_utc().isoformat(),
    }

    sb_update("employees", "employee_id", emp_id, update)
    invalidate_employees_cache()
    return {"ok": True, "employee_id": emp_id}

@api_router.post("/auth/act-as")
async def auth_act_as(payload: ActAs, cu: User = Depends(get_current_user)):
    # act-as is per-device: the frontend stores the employee_id in localStorage
    # and sends it as X-Acting-As header on every request.
    # We just return the user with the acting_as field set — no global state mutation.
    u = cu.model_dump(mode="json")
    u["acting_as_employee_id"] = payload.employee_id
    return u

# ---- WhatsApp Service Layer (Interakt) ----
class WhatsAppService:
    @staticmethod
    def send_template(phone: str, template_name: str, values: List[str] = []):
        """
        Sends a WhatsApp template via Interakt API.
        Values: List of strings to fill in the {{1}}, {{2}} placeholders.
        """
        # TEMPORARY: Disabled WhatsApp Business API for client campaign automation
        logging.info(f"[SIMULATION] Interakt Template '{template_name}' would be sent to {phone} with values {values}")
        return {"status": "simulated", "message": "WhatsApp automation is temporarily disabled"}

# ---- AI Assistant Service (from umang.py) ----
class AIService:
    @staticmethod
    def generate_reply(message_text: str):
        if not OPENAI_API_KEY:
            return "Hi! Thanks for your message. An agent will get back to you soon."
        
        url = "https://api.openai.com/v1/chat/completions"
        headers = {"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"}
        payload = {
            "model": "gpt-4o-mini",
            "messages": [
                {
                    "role": "system",
                    "content": """
You are a professional real estate assistant for Umang Hometech LLP.
Your goal:
- Sound human and friendly
- Ask about budget, location, requirement
- Convert user into site visit
Example: "Hi 😊 Are you looking for 2BHK or 3BHK? We have great options available 🔥"
Always:
- Ask questions
- Engage user
- Push for site visit
- Don't be a bot
"""
                },
                {"role": "user", "content": message_text}
            ]
        }
        try:
            r = httpx.post(url, headers=headers, json=payload, timeout=15)
            data = r.json()
            return data["choices"][0]["message"]["content"]
        except Exception as e:
            logging.error(f"OpenAI API Error: {e}")
            return "Thank you for reaching out. We will assist you shortly."

    @staticmethod
    def generate_lead_summary(timeline: List[Dict[str, Any]]):
        if not OPENAI_API_KEY:
            return "No summary available."
        
        # Format the timeline for AI
        history = "\n".join([f"- {a.get('created_at', '')[:10]}: {a.get('text', '')}" for a in timeline[:20]])
        
        url = "https://api.openai.com/v1/chat/completions"
        headers = {"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"}
        payload = {
            "model": "gpt-4o-mini",
            "messages": [
                {
                    "role": "system",
                    "content": "You are a senior real estate manager. Summarize the following lead activity into a single, punchy, professional sentence that tells the agent exactly what the current status is and what to do next. Be concise (max 20 words)."
                },
                {"role": "user", "content": f"Activity History:\n{history}"}
            ]
        }
        try:
            r = httpx.post(url, headers=headers, json=payload, timeout=15)
            data = r.json()
            return data["choices"][0]["message"]["content"]
        except Exception as e:
            logging.error(f"Lead Summary Error: {e}")
            return "Could not generate summary."

# ---- Assignment Engine (Round Robin) ----
ASSIGNABLE_EMPLOYEE_ROLES = {"telecaller", "site_visit", "sales_executive"}


def assign_lead_round_robin() -> Optional[str]:
    """Round-robin pick for manager-initiated bulk assign only — not used on new lead intake."""
    emps = sb_select("employees", {
        "active": "eq.true",
        "select": "employee_id,role,last_assigned_at",
        "order": "last_assigned_at.asc.nullslast",
    })
    candidates = [e for e in emps if e.get("role") in ASSIGNABLE_EMPLOYEE_ROLES]
    if not candidates:
        candidates = [e for e in emps if e.get("role") == "admin"]
    if not candidates:
        return None
    selected = candidates[0]
    sb_update("employees", "employee_id", selected["employee_id"], {"last_assigned_at": now_utc().isoformat()})
    return selected["employee_id"]


def stamp_lead_assignment(
    assigned_to: str,
    assigned_by: Optional[str] = None,
    current_stage: Optional[str] = None,
) -> Dict[str, Any]:
    """Fields applied when a lead is assigned (auto or manager)."""
    now = now_utc().isoformat()
    data: Dict[str, Any] = {
        "assigned_to": assigned_to,
        "assigned_at": now,
        "last_employee_action_at": None,
        "updated_at": now,
    }
    if assigned_by:
        data["assigned_by"] = assigned_by
    if current_stage in (None, "new"):
        data["stage"] = "assigned"
    return data


def assign_lead_to_employee(
    lead_id: str,
    old_lead: Dict[str, Any],
    employee_id: str,
    actor: Optional[User],
    reactivate: bool = True,
) -> Dict[str, Any]:
    """Single lead assignment — shared by PATCH and bulk assign."""
    emps = sb_select("employees", {"employee_id": f"eq.{employee_id}", "select": "name", "limit": "1"})
    emp_name = emps[0]["name"] if emps else employee_id
    actor_id = None
    if actor:
        actor_id = actor.acting_as_employee_id or actor.user_id
    data = stamp_lead_assignment(employee_id, assigned_by=actor_id, current_stage=old_lead.get("stage"))
    if reactivate and old_lead.get("status") == "negative":
        data.update(fresh_lead_reset_fields("assigned"))
    updated = sb_update("leads", "lead_id", lead_id, data) or {**old_lead, **data}
    update_cached_lead(lead_id, data)
    log_activity(actor, "lead_assigned", f"Assigned lead to {emp_name}", lead_id=lead_id)
    create_notification(employee_id, "Lead assigned", f"{old_lead.get('name', 'Lead')} has been assigned to you.", lead_id=lead_id)
    return updated

# ---- Activity Logger ----
def log_activity(actor, type_, text, lead_id=None): 
    user_id_to_log = None
    actor_name = "System"
    
    if actor:
        user_id_to_log = actor.acting_as_employee_id or actor.employee_id or actor.user_id
        actor_name = actor.name

    # Check if they are acting as an employee
    if actor and actor.acting_as_employee_id:
        emps = sb_select("employees", {"employee_id": f"eq.{actor.acting_as_employee_id}", "select": "name"})
        if emps:
            actor_name = emps[0]["name"]

    act = {
        "activity_id": gen_id("act"),
        "lead_id": lead_id,
        "user_id": user_id_to_log,
        "type": type_,
        "text": f"[{actor_name}] {text}",
        "created_at": now_utc().isoformat(),
    }
    sb_insert("activities", act)
    return act

def get_lead_record(lead_id: str, select: str = "*"):
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": select})
    if leads:
        return leads[0]
    cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id]
    return cache_match[0] if cache_match else None

def update_cached_lead(lead_id: str, data: dict):
    def _merge(row: dict) -> dict:
        merged = enrich_lead_display_fields({**row, **data})
        return merged

    SESSION_CACHE["leads"] = [
        (_merge(l) if l.get("lead_id") == lead_id else l)
        for l in SESSION_CACHE["leads"]
    ]
    patched = False
    with _leads_cache_lock:
        cached_rows = _leads_cache.get("data") or []
        if _leads_cache.get("ts") and cached_rows:
            next_rows = []
            for row in cached_rows:
                if row.get("lead_id") == lead_id:
                    next_rows.append(_merge(row))
                    patched = True
                else:
                    next_rows.append(row)
            if patched:
                _leads_cache["data"] = next_rows
    if patched:
        _assignment_stats_cache["ts"] = 0.0
        _employee_stats_cache["ts"] = 0.0
        _dashboard_stats_cache["ts"] = 0.0
        return
    invalidate_leads_cache()

def first_related_record(table: str, cache_key: str, lead_id: str):
    rows = sb_select(table, {"lead_id": f"eq.{lead_id}", "select": "*", "limit": "1"})
    if rows:
        return rows[0]
    cache_match = [r for r in SESSION_CACHE[cache_key] if r.get("lead_id") == lead_id]
    return cache_match[0] if cache_match else None

def get_visit_record(visit_id: str):
    rows = sb_select("visits", {"visit_id": f"eq.{visit_id}", "select": "*", "limit": "1"})
    if rows:
        return rows[0]
    cache_match = [v for v in SESSION_CACHE["visits"] if v.get("visit_id") == visit_id]
    return cache_match[0] if cache_match else None

def normalize_follow_up_time(time_value: str) -> str:
    """Accept 12h (11:30 AM) or 24h (14:00) — returns HH:MM for storage."""
    s = (time_value or "").strip()
    if not s:
        raise HTTPException(400, "Follow-up time is required")
    m12 = re.match(r"^(\d{1,2}):(\d{2})\s*(AM|PM|am|pm)$", s)
    if m12:
        hour, minute, meridiem = int(m12.group(1)), int(m12.group(2)), m12.group(3).upper()
        if hour < 1 or hour > 12 or minute > 59:
            raise HTTPException(400, "Invalid follow-up time")
        if meridiem == "PM" and hour != 12:
            hour += 12
        elif meridiem == "AM" and hour == 12:
            hour = 0
        return f"{hour:02d}:{minute:02d}"
    m24 = re.match(r"^(\d{1,2}):(\d{2})$", s)
    if m24:
        hour, minute = int(m24.group(1)), int(m24.group(2))
        if hour > 23 or minute > 59:
            raise HTTPException(400, "Invalid follow-up time")
        return f"{hour:02d}:{minute:02d}"
    raise HTTPException(400, "Follow-up time must be like 11:30 AM or 14:00")


def parse_follow_up_at(follow_up_date: str, follow_up_time: str) -> datetime:
    date_value = (follow_up_date or "").strip()
    time_value = normalize_follow_up_time(follow_up_time)
    if not date_value:
        raise HTTPException(400, "Follow-up date is required")
    try:
        parsed = datetime.fromisoformat(f"{date_value}T{time_value}")
    except ValueError:
        raise HTTPException(400, "Follow-up date/time must be valid")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed

def follow_up_display_parts(value: Optional[str]):
    try:
        parsed = datetime.fromisoformat(str(value or "").replace("Z", "+00:00"))
    except ValueError:
        parsed = now_utc()
    return {
        "follow_up_date": parsed.date().isoformat(),
        "follow_up_time": parsed.strftime("%H:%M"),
        "follow_up_day": parsed.strftime("%A"),
        "follow_up_at": parsed.isoformat(),
    }

def ensure_visit_record(lead_id: str, lead_name: Optional[str] = None, assigned_to: Optional[str] = None):
    existing = first_related_record("visits", "visits", lead_id)
    if existing:
        return existing
    lead = get_lead_record(lead_id)
    if not lead:
        return None
    v = {
        "visit_id": gen_id("vis"),
        "lead_id": lead_id,
        "lead_name": lead_name or lead.get("name", "Lead"),
        "scheduled_at": now_utc().isoformat(),
        "assigned_to": assigned_to or lead.get("assigned_to"),
        "status": "scheduled",
        "feedback": None,
        "interested": None,
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("visits", v)
    SESSION_CACHE["visits"].insert(0, result or v)
    return result or v

def booking_brokerage_amount(booking: Dict[str, Any]) -> float:
    """Brokerage stored in brokerage_amount column or embedded in agreement_status."""
    val = float(booking.get("brokerage_amount") or 0)
    if val > 0:
        return val
    raw = str(booking.get("agreement_status") or "")
    m = re.search(r"Brokerage:\s*([0-9.]+)", raw)
    return float(m.group(1)) if m else 0.0

def is_legacy_skeleton_booking(booking: Dict[str, Any]) -> bool:
    """Auto-created placeholder rows (site-visit sync) — hidden from the booking list."""
    return (
        str(booking.get("property_name") or "").strip().lower() == "selected property"
        and float(booking.get("booking_amount") or 0) == 0
        and float(booking.get("token_received") or 0) == 0
        and not booking.get("flat_cost")
        and not booking.get("agreement_value")
    )

def ensure_booking_record(lead_id: str, lead_name: Optional[str] = None):
    existing = first_related_record("bookings", "bookings", lead_id)
    if existing:
        return existing
    lead = get_lead_record(lead_id)
    if not lead:
        return None
    b = {
        "booking_id": gen_id("bkg"),
        "lead_id": lead_id,
        "lead_name": lead_name or lead.get("name", "Lead"),
        "property_name": "Selected Property",
        "booking_amount": 0,
        "token_received": 0,
        "agreement_status": "pending",
        "payment_progress": 0,
        "payment_status": "pending",
        "status": "active",
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("bookings", b)
    SESSION_CACHE["bookings"].insert(0, result or b)
    return result or b

def ensure_loan_record(lead_id: str, lead_name: Optional[str] = None):
    existing = first_related_record("loans", "loans", lead_id)
    if existing:
        return existing
    lead = get_lead_record(lead_id)
    if not lead:
        return None
    ln = {
        "loan_id": gen_id("lon"),
        "lead_id": lead_id,
        "lead_name": lead_name or lead.get("name", "Lead"),
        "bank_name": "Bank Pending",
        "amount": 0,
        "application_status": "pending",
        "bank_stage": "documentation",
        "documents_status": "pending",
        "pending_documents": [],
        "emi_eligible": None,
        "progress": 0,
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("loans", ln)
    if not result:
        minimal = {
            "loan_id": ln["loan_id"],
            "lead_id": ln["lead_id"],
            "lead_name": ln["lead_name"],
            "bank_name": ln["bank_name"],
            "amount": ln["amount"],
            "application_status": ln["application_status"],
            "bank_stage": ln["bank_stage"],
            "emi_eligible": ln["emi_eligible"],
            "progress": ln["progress"],
            "created_at": ln["created_at"],
        }
        result = sb_insert("loans", minimal)
        ln = minimal
    SESSION_CACHE["loans"].insert(0, result or ln)
    return result or ln

def create_notification(user_id: Optional[str], title: str, message: str, lead_id: Optional[str] = None, type_: str = "workflow"):
    if not user_id:
        return None
    n = {
        "notification_id": gen_id("ntf"),
        "user_id": user_id,
        "lead_id": lead_id,
        "type": type_,
        "title": title,
        "message": message,
        "is_read": False,
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("notifications", n)
    if result:
        SESSION_CACHE["notifications"].insert(0, result)
    return result

def create_customer_from_lead(lead_id: str, actor=None):
    existing = sb_select("customers", {"lead_id": f"eq.{lead_id}", "select": "*", "limit": "1"})
    if existing:
        return existing[0]
    lead = get_lead_record(lead_id)
    if not lead:
        return None
    booking = first_related_record("bookings", "bookings", lead_id)
    loan = first_related_record("loans", "loans", lead_id)
    customer = {
        "customer_id": gen_id("cus"),
        "lead_id": lead_id,
        "name": lead.get("name"),
        "phone": lead.get("phone"),
        "email": lead.get("email"),
        "location": lead.get("location"),
        "budget": lead.get("budget"),
        "property_type": lead.get("property_type"),
        "source": lead.get("source"),
        "booking_id": booking.get("booking_id") if booking else None,
        "loan_id": loan.get("loan_id") if loan else None,
        "status": "converted",
        "converted_at": now_utc().isoformat(),
        "created_at": now_utc().isoformat(),
        "updated_at": now_utc().isoformat(),
    }
    result = sb_insert("customers", customer)
    if result:
        SESSION_CACHE["customers"].insert(0, result)
        log_activity(actor, "converted_customer", "Lead converted into customer record.", lead_id=lead_id)
    return result or customer

# ---- Integration Intake Helpers ----
async def request_payload(request: Request) -> Any:
    content_type = request.headers.get("content-type", "").lower()
    if "application/json" in content_type:
        try:
            return await request.json()
        except Exception:
            return {}
    if "form" in content_type:
        form = await request.form()
        return dict(form)
    raw = await request.body()
    if not raw:
        return {}
    try:
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return {"raw": raw.decode("utf-8", errors="ignore")}

def record_integration_event(source: str, payload: Any, status: str, lead_id: Optional[str] = None,
                             external_id: Optional[str] = None, error: Optional[str] = None):
    event = {
        "event_id": gen_id("evt"),
        "source": source,
        "external_id": external_id,
        "status": status,
        "lead_id": lead_id,
        "error": error,
        "raw_payload": safe_json(payload),
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("integration_events", event)
    return result or event

def parse_epoch_seconds(value: Any) -> Optional[int]:
    """Parse Housing lead_date / Meta created_time to Unix seconds."""
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    if raw.isdigit():
        ts = int(raw)
        if ts > 1e12:
            ts = ts // 1000
        return ts
    iso = parse_external_datetime(value)
    if iso:
        try:
            return int(datetime.fromisoformat(iso.replace("Z", "+00:00")).timestamp())
        except Exception:
            return None
    return None


def get_housing_lead_epoch(payload: Dict[str, Any]) -> Optional[int]:
    return parse_epoch_seconds(pick_first(payload, [
        "lead_date", "created_time", "created_at", "lead_created_time", "submitted_at", "time",
    ]))


def get_last_housing_sync_end_epoch() -> Optional[int]:
    rows = sb_select("integration_events", {
        "source": "eq.Housing.com",
        "status": "eq.housing_sync_checkpoint",
        "select": "raw_payload",
        "order": "created_at.desc",
        "limit": "1",
    })
    if not rows:
        return None
    raw = rows[0].get("raw_payload")
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except Exception:
            return None
    if isinstance(raw, dict) and raw.get("end_date") is not None:
        try:
            return int(raw["end_date"])
        except (TypeError, ValueError):
            return None
    return None


def record_housing_sync_checkpoint(end_date: int, meta: Dict[str, Any]) -> None:
    record_integration_event("Housing.com", {"end_date": end_date, **meta}, "housing_sync_checkpoint")


def housing_sync_window(
    mode: str,
    start_date: Optional[int],
    end_date: Optional[int],
    allow_historical: bool,
) -> tuple:
    """Compute API pull window — poll only fetches new leads since last checkpoint."""
    end_date = end_date or int(now_utc().timestamp())
    if start_date is not None:
        if not allow_historical and (end_date - start_date) > HOUSING_MANUAL_DEFAULT_WINDOW_SEC:
            start_date = end_date - HOUSING_MANUAL_DEFAULT_WINDOW_SEC
        if end_date - start_date > HOUSING_API_MAX_RANGE_SEC:
            start_date = end_date - HOUSING_API_MAX_RANGE_SEC
        return start_date, end_date

    if mode in ("poll", "auto", "cron"):
        last_end = get_last_housing_sync_end_epoch()
        if last_end:
            start_date = max(last_end - HOUSING_POLL_OVERLAP_SEC, end_date - HOUSING_POLL_INITIAL_WINDOW_SEC)
        else:
            start_date = end_date - HOUSING_POLL_INITIAL_WINDOW_SEC
    else:
        start_date = end_date - HOUSING_MANUAL_DEFAULT_WINDOW_SEC
    return start_date, end_date


def should_import_housing_lead_on_sync(
    payload: Dict[str, Any],
    start_date: int,
    end_date: int,
    mode: str = "manual",
) -> bool:
    """Skip old Housing API leads — only import when lead_date is inside the sync window."""
    lead_ts = get_housing_lead_epoch(payload)
    if lead_ts is None:
        # Auto/poll: Housing API already scoped by start_date/end_date — don't drop new leads.
        return mode in ("poll", "auto", "cron")
    margin = HOUSING_POLL_OVERLAP_SEC if mode in ("poll", "auto", "cron") else 0
    return (start_date - margin) <= lead_ts <= (end_date + 60)


def system_integration_actor():
    """Synthetic user for background Housing sync (no login required)."""
    class _SystemUser:
        user_id = "system"
        email = "system@integrations"
        role = "admin"
        name = "System"
        employee_id = None
        acting_as_employee_id = None

    return _SystemUser()


def run_housing_sync_background(mode: str = "auto") -> Dict[str, Any]:
    """Run Housing pull sync on a worker thread (server cron / background loop)."""
    if not HOUSING_PROFILE_ID or not HOUSING_ENCRYPTION_KEY:
        return {"status": "skipped", "reason": "credentials_not_configured"}
    try:
        return asyncio.run(_housing_sync_impl(HousingSyncRequest(), system_integration_actor(), mode=mode))
    except HTTPException as exc:
        logging.error("Housing background sync failed: %s", exc.detail)
        return {"status": "error", "detail": exc.detail}
    except Exception as exc:
        logging.exception("Housing background sync failed")
        return {"status": "error", "detail": str(exc)[:180]}


def parse_external_datetime(value: Any) -> Optional[str]:
    """Best-effort parse of a platform-provided lead time (Meta created_time,
    Housing lead_date) into an ISO timestamp. Returns None if unparseable."""
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    # Epoch seconds / milliseconds.
    if raw.isdigit():
        try:
            ts = int(raw)
            if ts > 1e12:  # milliseconds
                ts = ts / 1000.0
            return datetime.fromtimestamp(ts, tz=timezone.utc).isoformat()
        except Exception:
            return None
    candidate = raw.replace("Z", "+00:00")
    # Meta uses +0000 (no colon); normalize to +00:00.
    m = re.search(r"([+-]\d{2})(\d{2})$", candidate)
    if m:
        candidate = candidate[: m.start()] + f"{m.group(1)}:{m.group(2)}"
    for parser in (datetime.fromisoformat,):
        try:
            dt = parser(candidate)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.astimezone(timezone.utc).isoformat()
        except Exception:
            pass
    return None

def format_budget_lakhs(val: Any) -> Optional[str]:
    """Short lakh label for display/storage (45 not 4500000)."""
    if val is None or val == "":
        return None
    try:
        n = float(val)
    except (TypeError, ValueError):
        return clean_text(val)
    if n <= 0:
        return None
    if n >= 100000:
        lakhs = n / 100000
    elif n >= 1000:
        lakhs = n / 100000
    else:
        lakhs = n
    if lakhs == int(lakhs):
        return str(int(lakhs))
    return f"{lakhs:.1f}".rstrip("0").rstrip(".")

def format_budget_range_lakhs(min_price: Any, max_price: Any) -> Optional[str]:
    a = format_budget_lakhs(min_price)
    b = format_budget_lakhs(max_price)
    if a and b:
        return f"{a} - {b}"
    return a or b

def extract_bhk_configuration(payload: Dict[str, Any]) -> Optional[str]:
    for key in ("configuration", "config", "bhk", "requirement", "unit_type", "property_type", "property_field"):
        raw = clean_text(payload.get(key))
        if not raw:
            continue
        m = re.search(r"(\d)\s*bhk", raw, re.I)
        if m:
            return f"{m.group(1)} BHK"
        if re.fullmatch(r"\d+", raw.strip()):
            return f"{raw.strip()} BHK"
    return None

def lead_from_payload(payload: Dict[str, Any], source: str) -> Dict[str, Any]:
    name = clean_text(pick_first(payload, [
        "customer_name", "full_name", "name", "lead_name", "first_name", "contact.name",
    ])) or "Valued Customer"
    phone = normalize_phone(pick_first(payload, [
        "phone_number", "phone", "mobile", "contact_number", "lead_phone", "contact.phone",
    ]))
    email = clean_text(pick_first(payload, [
        "email", "lead_email", "email_address", "contact.email",
    ]))
    locality = clean_text(pick_first(payload, ["locality_name", "locality", "project_locality"]))
    city = clean_text(pick_first(payload, ["city_name", "city", "contact.city"]))
    location = clean_text(pick_first(payload, [
        "city", "location", "locality", "project_locality", "address", "contact.city",
    ])) or ", ".join(part for part in [locality, city] if part)
    min_price = pick_first(payload, ["min_price", "min_budget"])
    max_price = pick_first(payload, ["max_price", "max_budget"])
    budget = clean_text(pick_first(payload, [
        "budget", "price", "budget_range", "max_budget", "requirement_budget",
    ]))
    if not budget and (min_price or max_price):
        budget = format_budget_range_lakhs(min_price, max_price)
    property_type = extract_bhk_configuration(payload) or clean_text(pick_first(payload, [
        "property_type", "configuration", "config", "bhk", "requirement", "unit_type",
    ]))
    external_id = clean_text(pick_first(payload, [
        "lead_id", "id", "uuid", "enquiry_id", "leadgen_id", "external_id",
    ]))
    if not external_id and source == "Housing.com":
        project_id = clean_text(payload.get("project_id"))
        lead_date = clean_text(payload.get("lead_date"))
        phone_key = normalize_phone(pick_first(payload, ["lead_phone", "phone", "mobile"]))
        parts = [p for p in [project_id, phone_key, lead_date] if p]
        if len(parts) >= 2:
            external_id = ":".join(parts)
    project_name = clean_text(pick_first(payload, ["project_name", "project", "property_name"]))
    notes = clean_text(pick_first(payload, ["notes", "comment", "message", "remarks", "query"]))
    note_parts = []
    if project_name:
        note_parts.append(f"Project: {project_name}")
    if clean_text(payload.get("project_id")):
        note_parts.append(f"Project ID: {clean_text(payload.get('project_id'))}")
    if source == "Facebook":
        if clean_text(payload.get("page_id")):
            note_parts.append(f"Facebook Page ID: {clean_text(payload.get('page_id'))}")
        if clean_text(payload.get("form_id")):
            note_parts.append(f"Facebook Form ID: {clean_text(payload.get('form_id'))}")
        if payload.get("created_time"):
            note_parts.append(f"Lead created: {payload.get('created_time')}")
    if notes:
        note_parts.append(notes)

    external_created_at = parse_external_datetime(pick_first(payload, [
        "created_time", "lead_date", "created_at", "lead_created_time", "submitted_at", "time",
    ]))

    return {
        "name": name,
        "phone": phone,
        "email": email,
        "location": location,
        "budget": budget,
        "property_type": property_type,
        "source": source,
        "external_lead_id": external_id,
        "external_created_at": external_created_at,
        "notes": "\n".join(note_parts) if note_parts else None,
    }

def find_existing_integrated_lead(phone: str, email: Optional[str], source: str, external_id: Optional[str]):
    if external_id:
        rows = sb_select("leads", {
            "external_lead_id": f"eq.{external_id}",
            "select": "*",
            "limit": "1",
        })
        if rows:
            return rows[0]
        if source:
            events = sb_select("integration_events", {
                "source": f"eq.{source}",
                "external_id": f"eq.{external_id}",
                "status": "eq.created",
                "select": "lead_id",
                "limit": "1",
            })
            if events and events[0].get("lead_id"):
                rows = sb_select("leads", {
                    "lead_id": f"eq.{events[0]['lead_id']}",
                    "select": "*",
                    "limit": "1",
                })
                if rows:
                    return rows[0]

    scoped_source = source if source in ("Housing.com", "Facebook") else None

    if phone:
        params: Dict[str, str] = {"phone": f"eq.{phone}", "select": "*", "limit": "1"}
        if scoped_source:
            params["source"] = f"eq.{scoped_source}"
        rows = sb_select("leads", params)
        if rows:
            return rows[0]

    if email:
        params = {"email": f"eq.{email}", "select": "*", "limit": "1"}
        if scoped_source:
            params["source"] = f"eq.{scoped_source}"
        rows = sb_select("leads", params)
        if rows:
            return rows[0]

    for lead in SESSION_CACHE["leads"]:
        if external_id and lead.get("external_lead_id") == external_id:
            return lead
        if scoped_source and lead.get("source") != scoped_source:
            continue
        if phone and lead.get("phone") == phone:
            return lead
        if email and lead.get("email") == email:
            return lead
    return None

def create_integrated_lead(payload: Dict[str, Any], source: str, actor=None, *, quiet: bool = False) -> Dict[str, Any]:
    normalized = lead_from_payload(payload, source)
    if not normalized["phone"] and not normalized.get("email"):
        if not quiet:
            record_integration_event(source, payload, "ignored", external_id=normalized.get("external_lead_id"), error="Missing phone/email")
        return {"status": "ignored", "reason": "missing_phone_or_email", "payload": normalized}

    existing = find_existing_integrated_lead(
        normalized["phone"],
        normalized.get("email"),
        source,
        normalized.get("external_lead_id"),
    )
    if existing:
        update_data = {"updated_at": now_utc().isoformat()}
        for key in ("email", "budget", "location", "property_type", "notes"):
            if normalized.get(key) and not existing.get(key):
                update_data[key] = normalized[key]
        updated = sb_update("leads", "lead_id", existing["lead_id"], update_data) if len(update_data) > 1 else existing
        merged = {**existing, **update_data}
        update_cached_lead(existing["lead_id"], update_data)
        if not quiet:
            log_activity(actor, "integration_duplicate", f"Duplicate lead received from {source}; existing record refreshed.", lead_id=existing["lead_id"])
            record_integration_event(source, payload, "duplicate", lead_id=existing["lead_id"], external_id=normalized.get("external_lead_id"))
        return {"status": "duplicate", "lead": updated or merged, "lead_id": existing["lead_id"]}

    brokerage = is_brokerage_lead(source, payload)
    assigned_to = None
    initial_stage = "broker" if brokerage else "new"
    lead_type = "brokerage" if brokerage else "standard"
    now = now_utc().isoformat()
    # Store the platform submission time when available (Housing lead_date / Meta created_time).
    created_at = normalized.get("external_created_at") or now
    lead_id = gen_id("lead")
    base_lead = {
        "lead_id": lead_id,
        "name": normalized["name"],
        "phone": normalized["phone"] or normalized.get("email") or "",
        "email": normalized.get("email"),
        "budget": normalized.get("budget"),
        "location": normalized.get("location"),
        "property_type": normalized.get("property_type"),
        "notes": normalized.get("notes"),
        "source": source,
        "stage": initial_stage,
        "status": "active",
        "assigned_to": assigned_to,
        "assigned_at": now if assigned_to else None,
        "lead_type": lead_type,
        "created_at": created_at,
        "updated_at": now,
    }
    optional_lead = {
        **base_lead,
        "external_lead_id": normalized.get("external_lead_id"),
        "external_created_at": normalized.get("external_created_at"),
        "integration_uuid": clean_text(payload.get("integration_uuid")) or HOUSING_INTEGRATION_UUID if source == "Housing.com" else clean_text(payload.get("integration_uuid")),
        "raw_payload": safe_json(payload),
    }
    result = sb_insert("leads", {k: v for k, v in optional_lead.items() if v is not None})
    if not result:
        result = sb_insert("leads", {k: v for k, v in base_lead.items() if v is not None})
    lead_record = result or base_lead
    SESSION_CACHE["leads"].insert(0, lead_record)
    invalidate_leads_cache()
    if not quiet:
        if brokerage:
            log_activity(actor, "broker_lead_received", f"Brokerage lead stored for future: {normalized['name']} ({source})", lead_id=lead_id)
        else:
            log_activity(actor, "integration_enquiry", f"New lead received from {source}: {normalized['name']} (awaiting manager assignment)", lead_id=lead_id)
        record_integration_event(source, payload, "created", lead_id=lead_id, external_id=normalized.get("external_lead_id"))
    return {"status": "created", "lead": lead_record, "lead_id": lead_id}

def extract_facebook_lead_events(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Parse Meta leadgen webhook payloads (entry[].changes[].value)."""
    events: List[Dict[str, Any]] = []
    for entry in payload.get("entry", []) if isinstance(payload, dict) else []:
        if not isinstance(entry, dict):
            continue
        for change in entry.get("changes", []) if isinstance(entry.get("changes"), list) else []:
            if not isinstance(change, dict):
                continue
            field = clean_text(change.get("field"))
            if field and field != "leadgen":
                logging.info("Facebook webhook: skipping change field=%s", field)
                continue
            value = change.get("value", {}) if isinstance(change.get("value"), dict) else {}
            leadgen_id = clean_text(value.get("leadgen_id") or value.get("lead_id"))
            if not leadgen_id:
                continue
            events.append({
                "leadgen_id": leadgen_id,
                "page_id": normalize_meta_page_id(value.get("page_id"), entry.get("id")),
                "form_id": clean_text(value.get("form_id")),
                "created_time": value.get("created_time"),
                "ad_id": clean_text(value.get("ad_id")),
                "adgroup_id": clean_text(value.get("adgroup_id")),
                "raw_value": value,
            })
    direct_id = clean_text(payload.get("leadgen_id") or payload.get("lead_id")) if isinstance(payload, dict) else None
    if direct_id and not any(e["leadgen_id"] == direct_id for e in events):
        events.append({
            "leadgen_id": direct_id,
            "page_id": normalize_meta_page_id(payload.get("page_id")),
            "form_id": clean_text(payload.get("form_id")),
            "created_time": payload.get("created_time"),
            "ad_id": clean_text(payload.get("ad_id")),
            "adgroup_id": clean_text(payload.get("adgroup_id")),
            "raw_value": payload,
        })
    return events

def extract_facebook_lead_ids(payload: Dict[str, Any]) -> List[str]:
    return list(dict.fromkeys(e["leadgen_id"] for e in extract_facebook_lead_events(payload)))

def facebook_fields_to_payload(data: Dict[str, Any], leadgen_id: Optional[str] = None) -> Dict[str, Any]:
    payload: Dict[str, Any] = {"leadgen_id": leadgen_id or data.get("id")}
    for item in data.get("field_data", []) if isinstance(data.get("field_data"), list) else []:
        name = item.get("name")
        values = item.get("values") or []
        if name and values:
            payload[name] = values[0]
    payload.update({k: v for k, v in data.items() if k not in {"field_data"}})
    return normalize_meta_field_payload(payload)

def normalize_meta_field_payload(payload: Dict[str, Any]) -> Dict[str, Any]:
    """Map Meta custom form fields (question1, etc.) to name/phone/email."""
    out = dict(payload)
    for key, val in list(out.items()):
        if val is None:
            continue
        text = clean_text(str(val))
        if not text:
            continue
        k = key.lower()
        if any(t in k for t in ("email", "mail")) and not out.get("email"):
            out["email"] = text
        elif any(t in k for t in ("phone", "mobile", "contact", "whatsapp", "number")) and not out.get("phone"):
            out["phone"] = text
        elif any(t in k for t in ("name", "full_name", "customer", "first_name")) and not out.get("full_name"):
            out["full_name"] = text
    for val in out.values():
        if not isinstance(val, str):
            continue
        s = val.strip()
        if not s:
            continue
        if not out.get("email") and "@" in s and "." in s.split("@", 1)[-1]:
            out["email"] = s
        if not out.get("phone"):
            digits = re.sub(r"\D", "", s)
            if len(digits) >= 10:
                out["phone"] = s
    return out

def facebook_graph_get(path: str, params: Optional[Dict[str, Any]] = None, access_token: Optional[str] = None) -> Dict[str, Any]:
    token = access_token or FACEBOOK_PAGE_ACCESS_TOKEN
    if not token:
        raise HTTPException(status_code=400, detail="FACEBOOK_PAGE_ACCESS_TOKEN is not configured")
    url = f"https://graph.facebook.com/{FACEBOOK_GRAPH_VERSION}/{path.lstrip('/')}"
    query = {"access_token": token, **(params or {})}
    r = _http.get(url, params=query, timeout=60)
    if r.status_code >= 400:
        raise HTTPException(status_code=502, detail=_facebook_graph_error_detail(r))
    return r.json()


def _facebook_graph_error_detail(response: httpx.Response) -> str:
    try:
        body = response.json()
        err = body.get("error") if isinstance(body, dict) else {}
        if not isinstance(err, dict):
            err = {}
        code = err.get("code")
        msg = clean_text(err.get("message")) or response.text[:220]
        if code == 190 and "page access token" in msg.lower():
            return (
                "Page Access Token required (you have a User token). "
                "Graph API Explorer → User or Page → select your Facebook Page → "
                "add leads_retrieval + pages_show_list → Generate token → "
                "paste into Render FACEBOOK_PAGE_ACCESS_TOKEN."
            )
        return f"Meta Graph API error: {msg}"
    except Exception:
        return f"Meta Graph API error: {response.text[:220]}"


def facebook_token_me_id(token: Optional[str] = None) -> Optional[str]:
    tok = token or FACEBOOK_PAGE_ACCESS_TOKEN
    if not tok:
        return None
    try:
        data = facebook_graph_get("me", {"fields": "id"}, access_token=tok)
        return clean_text(data.get("id"))
    except HTTPException:
        return None

def facebook_graph_paginate(
    path: str,
    params: Optional[Dict[str, Any]] = None,
    max_items: int = 500,
    access_token: Optional[str] = None,
) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    after: Optional[str] = None
    base_params = dict(params or {})
    while len(items) < max_items:
        query = dict(base_params)
        if after:
            query["after"] = after
        data = facebook_graph_get(path, query, access_token=access_token)
        batch = data.get("data") if isinstance(data.get("data"), list) else []
        if not batch:
            break
        items.extend(batch)
        cursors = (data.get("paging") or {}).get("cursors") or {}
        after = cursors.get("after")
        if not after:
            break
    return items[:max_items]

def _facebook_page_tokens_map() -> Dict[str, str]:
    """Page id → page access token from me/accounts (cached 5 min)."""
    now = time.time()
    if now - float(_FACEBOOK_PAGE_TOKEN_CACHE.get("fetched_at") or 0) < 300:
        cached = _FACEBOOK_PAGE_TOKEN_CACHE.get("tokens")
        if isinstance(cached, dict) and cached:
            return cached
    tokens: Dict[str, str] = {}
    if not FACEBOOK_PAGE_ACCESS_TOKEN:
        return tokens
    try:
        accounts = facebook_graph_paginate(
            "me/accounts",
            {"fields": "id,name,access_token", "limit": "50"},
            max_items=50,
        )
        for acct in accounts:
            pid = clean_text(acct.get("id"))
            tok = clean_text(acct.get("access_token"))
            if pid and tok:
                tokens[pid] = tok
    except Exception as exc:
        logging.warning("Facebook: could not load page tokens from me/accounts: %s", exc)
    _FACEBOOK_PAGE_TOKEN_CACHE["tokens"] = tokens
    _FACEBOOK_PAGE_TOKEN_CACHE["fetched_at"] = now
    return tokens


def resolve_page_access_token(page_id: Optional[str] = None, *, require_page_token: bool = False) -> str:
    """Best Graph API token for a Lead Ad page (page-scoped token when available)."""
    if not FACEBOOK_PAGE_ACCESS_TOKEN:
        return ""
    pid = normalize_meta_page_id(page_id, FACEBOOK_PAGE_ID)
    me_id = facebook_token_me_id(FACEBOOK_PAGE_ACCESS_TOKEN)

    # Configured token is already a Page token for the target page.
    if pid and me_id == pid:
        return FACEBOOK_PAGE_ACCESS_TOKEN

    # Configured token is a Page token (me = page) and no specific page was requested.
    if not pid and me_id:
        try:
            facebook_graph_get(
                f"{me_id}/leadgen_forms",
                {"limit": "1", "fields": "id"},
                access_token=FACEBOOK_PAGE_ACCESS_TOKEN,
            )
            return FACEBOOK_PAGE_ACCESS_TOKEN
        except HTTPException:
            pass

    if pid:
        page_tok = _facebook_page_tokens_map().get(pid)
        if page_tok:
            return page_tok
        if require_page_token or me_id != pid:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Page Access Token required (current token is a User token). "
                    "Graph API Explorer → User or Page → select your Facebook Page → "
                    "permissions: leads_retrieval, pages_show_list → Generate token → "
                    "paste into Render FACEBOOK_PAGE_ACCESS_TOKEN. "
                    f"Set FACEBOOK_PAGE_ID={pid} on Render."
                ),
            )

    return FACEBOOK_PAGE_ACCESS_TOKEN


def resolve_facebook_page_context(page_id: Optional[str] = None) -> Tuple[str, str]:
    """Return (page_id, access_token) for Lead Ads API calls."""
    resolved_page = normalize_meta_page_id(page_id, FACEBOOK_PAGE_ID)
    if resolved_page:
        token = resolve_page_access_token(resolved_page, require_page_token=True)
        return resolved_page, token

    profile = facebook_graph_get("me", {"fields": "id,name"})
    me_id = clean_text(profile.get("id"))
    if not me_id:
        raise HTTPException(status_code=400, detail="Could not resolve Facebook identity. Set FACEBOOK_PAGE_ID in Render env.")

    def page_has_forms(candidate_id: str, token: Optional[str] = None) -> bool:
        try:
            facebook_graph_get(f"{candidate_id}/leadgen_forms", {"limit": "1", "fields": "id"}, access_token=token)
            return True
        except HTTPException:
            return False

    if page_has_forms(me_id):
        return me_id, FACEBOOK_PAGE_ACCESS_TOKEN

    accounts = facebook_graph_paginate(
        "me/accounts",
        {"fields": "id,name,access_token", "limit": "50"},
        max_items=50,
    )
    for acct in accounts:
        candidate = clean_text(acct.get("id"))
        page_token = clean_text(acct.get("access_token"))
        if candidate and page_token and page_has_forms(candidate, page_token):
            return candidate, page_token

    if resolved_page := clean_text(accounts[0].get("id") if accounts else ""):
        page_token = clean_text(accounts[0].get("access_token")) if accounts else ""
        if page_token:
            return resolved_page, page_token

    raise HTTPException(
        status_code=400,
        detail="Could not access Lead Ad forms. Use a Page access token or set FACEBOOK_PAGE_ID with leads_retrieval permission.",
    )

def resolve_facebook_page_id(page_id: Optional[str] = None) -> str:
    return resolve_facebook_page_context(page_id)[0]

def list_facebook_leadgen_forms(page_id: str, access_token: Optional[str] = None) -> List[Dict[str, Any]]:
    return facebook_graph_paginate(
        f"{page_id}/leadgen_forms",
        {"fields": "id,name,status,created_time", "limit": "100"},
        max_items=100,
        access_token=access_token,
    )

def list_facebook_form_leads(form_id: str, limit: int = 500, since_ts: Optional[int] = None, access_token: Optional[str] = None) -> List[Dict[str, Any]]:
    params: Dict[str, Any] = {
        "fields": "created_time,id,ad_id,form_id,field_data",
        "limit": "100",
    }
    if since_ts:
        params["filtering"] = json.dumps([{"field": "time_created", "operator": "GREATER_THAN", "value": since_ts}])
    return facebook_graph_paginate(f"{form_id}/leads", params, max_items=limit, access_token=access_token)

def _load_facebook_external_ids() -> set:
    """Preload Meta leadgen IDs already stored in CRM (speeds bulk import)."""
    known: set = set()
    offset = 0
    while offset < 10000:
        rows = sb_select("leads", {
            "source": "eq.Facebook",
            "select": "external_lead_id",
            "limit": "1000",
            "offset": str(offset),
        })
        if not rows:
            break
        for row in rows:
            eid = clean_text(row.get("external_lead_id"))
            if eid:
                known.add(eid)
        if len(rows) < 1000:
            break
        offset += 1000
    return known

def import_facebook_graph_lead(
    graph_lead: Dict[str, Any],
    known_external_ids: Optional[set] = None,
    *,
    quiet: bool = False,
) -> Dict[str, Any]:
    leadgen_id = clean_text(graph_lead.get("id"))
    if not leadgen_id or leadgen_id in META_FAKE_LEADGEN_IDS:
        return {"status": "ignored", "reason": "fake_or_missing_id", "leadgen_id": leadgen_id}
    if known_external_ids is not None and leadgen_id in known_external_ids:
        return {"status": "duplicate", "leadgen_id": leadgen_id, "reason": "already_imported"}
    payload = normalize_meta_field_payload(facebook_fields_to_payload(graph_lead, leadgen_id))
    result = create_integrated_lead(payload, "Facebook", quiet=quiet)
    result["leadgen_id"] = leadgen_id
    if known_external_ids is not None and leadgen_id and result.get("status") in ("created", "duplicate"):
        known_external_ids.add(leadgen_id)
    return result

def merge_facebook_meta_fields(lead_payload: Dict[str, Any], event: Dict[str, Any]) -> Dict[str, Any]:
    merged = {**lead_payload}
    for key in ("page_id", "form_id", "created_time", "ad_id", "adgroup_id"):
        if event.get(key) and not merged.get(key):
            merged[key] = event[key]
    if event.get("raw_value"):
        merged["webhook_value"] = event["raw_value"]
    return merged

def fetch_facebook_lead(leadgen_id: str, page_id: Optional[str] = None) -> Dict[str, Any]:
    resolved_page_id = normalize_meta_page_id(page_id)
    logging.info(
        "Facebook Graph API: fetching lead details for leadgen_id=%s page_id=%s resolved_page_id=%s",
        leadgen_id,
        page_id,
        resolved_page_id,
    )
    token = resolve_page_access_token(resolved_page_id)
    if not token:
        logging.warning(
            "Facebook Graph API: FACEBOOK_PAGE_ACCESS_TOKEN is not set; cannot fetch fields for leadgen_id=%s",
            leadgen_id,
        )
        record_integration_event(
            "Facebook",
            {"leadgen_id": leadgen_id, "page_id": page_id},
            "graph_error",
            external_id=leadgen_id,
            error="FACEBOOK_PAGE_ACCESS_TOKEN not configured on server",
        )
        return {
            "leadgen_id": leadgen_id,
            "notes": "Facebook leadgen_id received. Set FACEBOOK_PAGE_ACCESS_TOKEN (Page token with leads_retrieval) on Render.",
        }
    url = f"https://graph.facebook.com/{FACEBOOK_GRAPH_VERSION}/{leadgen_id}"
    r = _http.get(
        url,
        params={
            "access_token": token,
            "fields": "created_time,id,ad_id,form_id,field_data",
        },
        timeout=30,
    )
    if r.status_code >= 400:
        err_body = r.text[:500]
        logging.error(
            "Facebook Graph API: failed leadgen_id=%s status=%s body=%s",
            leadgen_id,
            r.status_code,
            err_body,
        )
        record_integration_event(
            "Facebook",
            {"leadgen_id": leadgen_id, "page_id": page_id, "status_code": r.status_code, "response": err_body},
            "graph_error",
            external_id=leadgen_id,
            error=err_body[:180],
        )
        raise HTTPException(status_code=502, detail=f"Meta lead retrieval failed: {err_body[:180]}")
    graph_data = r.json()
    payload = normalize_meta_field_payload(facebook_fields_to_payload(graph_data, leadgen_id))
    logging.info(
        "Facebook Graph API: success leadgen_id=%s name=%s phone=%s email=%s",
        leadgen_id,
        payload.get("full_name") or payload.get("name"),
        payload.get("phone_number") or payload.get("phone"),
        payload.get("email"),
    )
    record_integration_event("Facebook", graph_data, "graph_fetched", external_id=leadgen_id)
    return payload

async def process_facebook_lead_event(
    event: Dict[str, Any],
    body: Dict[str, Any],
) -> Dict[str, Any]:
    leadgen_id = event["leadgen_id"]
    if is_fake_meta_leadgen_id(leadgen_id):
        logging.info("Facebook: skipping Meta test/sample leadgen_id=%s", leadgen_id)
        record_integration_event(
            "Facebook",
            {"webhook": body, "event": event},
            "ignored_test",
            external_id=leadgen_id,
            error="Meta test webhook id — submit a real Lead Ad form to test",
        )
        return {"status": "ignored", "reason": "meta_test_webhook_id", "leadgen_id": leadgen_id}
    logging.info(
        "Facebook leadgen event: leadgen_id=%s page_id=%s form_id=%s created_time=%s",
        leadgen_id,
        event.get("page_id"),
        event.get("form_id"),
        event.get("created_time"),
    )
    record_integration_event(
        "Facebook",
        {"webhook": body, "event": event},
        "leadgen_received",
        external_id=leadgen_id,
    )

    has_inline_contact = any(
        body.get(k) for k in ["phone", "phone_number", "mobile", "email", "full_name", "name"]
    )
    if has_inline_contact and clean_text(body.get("leadgen_id")) == leadgen_id:
        lead_payload = merge_facebook_meta_fields(body, event)
        logging.info("Facebook: using inline webhook fields for leadgen_id=%s", leadgen_id)
    else:
        lead_payload = merge_facebook_meta_fields(fetch_facebook_lead(leadgen_id, event.get("page_id")), event)

    result = create_integrated_lead(lead_payload, "Facebook")
    logging.info(
        "Facebook CRM insert: leadgen_id=%s status=%s lead_id=%s reason=%s",
        leadgen_id,
        result.get("status"),
        result.get("lead_id"),
        result.get("reason"),
    )
    return result

def housing_expected_hash(current_time: str) -> str:
    return hmac.new(HOUSING_ENCRYPTION_KEY.encode("utf-8"), current_time.encode("utf-8"), hashlib.sha256).hexdigest()

def verify_housing_request(request: Request, payload: Dict[str, Any]):
    headers = request.headers
    signature = clean_text(
        headers.get("x-housing-signature")
        or headers.get("x-signature")
        or request.query_params.get("hash")
        or payload.get("hash")
    )
    if not HOUSING_WEBHOOK_SECRET and not (HOUSING_ENCRYPTION_KEY and signature):
        raise HTTPException(status_code=500, detail="Housing.com webhook secret is not configured")

    provided_uuid = clean_text(
        headers.get("x-housing-integration-uuid")
        or headers.get("x-integration-uuid")
        or request.query_params.get("integration_uuid")
        or payload.get("integration_uuid")
        or payload.get("uuid")
    )
    auth_header = headers.get("authorization", "")
    bearer = auth_header[7:].strip() if auth_header.lower().startswith("bearer ") else None

    if HOUSING_WEBHOOK_SECRET and provided_uuid != HOUSING_WEBHOOK_SECRET and bearer != HOUSING_WEBHOOK_SECRET:
        raise HTTPException(status_code=401, detail="Invalid Housing.com integration UUID")

    provided_profile = clean_text(
        headers.get("x-housing-profile-id")
        or headers.get("x-profile-id")
        or request.query_params.get("profile_id")
        or payload.get("profile_id")
        or payload.get("id")
    )
    if HOUSING_PROFILE_ID and provided_profile and provided_profile != HOUSING_PROFILE_ID:
        raise HTTPException(status_code=401, detail="Invalid Housing.com profile id")

    current_time = clean_text(
        headers.get("x-housing-timestamp")
        or headers.get("x-timestamp")
        or request.query_params.get("current_time")
        or payload.get("current_time")
        or payload.get("timestamp")
    )
    if signature:
        if not HOUSING_ENCRYPTION_KEY or not current_time:
            raise HTTPException(status_code=401, detail="Housing.com signature cannot be verified")
        try:
            if abs(int(now_utc().timestamp()) - int(current_time)) > 15 * 60:
                raise HTTPException(status_code=401, detail="Housing.com request timestamp expired")
        except ValueError:
            raise HTTPException(status_code=401, detail="Invalid Housing.com timestamp")
        expected = housing_expected_hash(current_time)
        if not hmac.compare_digest(signature.lower(), expected.lower()):
            raise HTTPException(status_code=401, detail="Invalid Housing.com signature")

def housing_sync_params(start_date: int, end_date: int) -> Dict[str, str]:
    if not HOUSING_PROFILE_ID or not HOUSING_ENCRYPTION_KEY:
        raise HTTPException(status_code=500, detail="Housing.com credentials are not configured")
    current_time = str(int(now_utc().timestamp()))
    return {
        "start_date": str(start_date),
        "end_date": str(end_date),
        "current_time": current_time,
        "hash": housing_expected_hash(current_time),
        "id": HOUSING_PROFILE_ID,
    }

# ---- Leads ----
@api_router.post("/leads/public")
async def create_lead_public(p: LeadCreatePublic):
    phone = normalize_phone(p.phone)
    if not phone:
        raise HTTPException(status_code=400, detail="A valid phone number is required.")
    lid = gen_id("lead")
    now = now_utc().isoformat()
    lead = {
        "lead_id": lid, "name": p.name, "phone": phone, "email": p.email,
        "budget": p.budget, "location": p.location, "property_type": p.property_type,
        "notes": p.notes, "source": p.source or "website", "stage": "new", "status": "active",
        "assigned_to": None,
        "created_at": now, "updated_at": now,
    }
    if p.starred is not None: lead["starred"] = p.starred
    result = sb_insert("leads", lead)
    log_activity(None, "website_enquiry", f"New website enquiry received from {p.name} (awaiting manager assignment).", lead_id=lid)
    
    # Auto-responder (Real Interakt API if key exists)
    WhatsAppService.send_template(phone, "welcome_enquiry", [p.name])
    
    return result or lead

# ---- Webhooks & Portal Integrations ----
@api_router.get("/facebook/webhook")
@api_router.get("/webhooks/facebook")
async def verify_fb_webhook(request: Request):
    params = request.query_params
    logging.info(
        "Facebook webhook GET verify: mode=%s token_match=%s",
        params.get("hub.mode"),
        params.get("hub.verify_token") == FACEBOOK_VERIFY_TOKEN,
    )
    if params.get("hub.mode") == "subscribe" and params.get("hub.verify_token") == FACEBOOK_VERIFY_TOKEN:
        return PlainTextResponse(params.get("hub.challenge", ""))
    raise HTTPException(status_code=403, detail="Invalid Facebook verify token")

@api_router.post("/facebook/webhook")
@api_router.post("/webhooks/facebook")
async def facebook_webhook(request: Request):
    raw_body = await request.body()
    logging.info(
        "Facebook webhook POST received from %s content_length=%s",
        request.client.host if request.client else "unknown",
        len(raw_body),
    )
    try:
        body = json.loads(raw_body.decode("utf-8")) if raw_body else {}
    except Exception:
        logging.error("Facebook webhook POST: invalid JSON body=%s", raw_body[:500])
        raise HTTPException(status_code=400, detail="Invalid Facebook webhook payload")
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="Invalid Facebook webhook payload")

    logging.info("Facebook webhook payload: %s", json.dumps(body, default=str)[:4000])
    record_integration_event("Facebook", body, "webhook_received")

    events = extract_facebook_lead_events(body)
    logging.info("Facebook webhook: parsed %s leadgen event(s)", len(events))

    created, duplicates, ignored, processed = [], [], [], []

    if events:
        for event in events:
            leadgen_id = event["leadgen_id"]
            try:
                result = await process_facebook_lead_event(event, body)
                processed.append({"leadgen_id": leadgen_id, "status": result.get("status")})
                if result["status"] == "created":
                    created.append(result["lead_id"])
                elif result["status"] == "duplicate":
                    duplicates.append(result["lead_id"])
                else:
                    ignored.append({"leadgen_id": leadgen_id, "reason": result.get("reason")})
            except HTTPException as exc:
                logging.error("Facebook webhook Graph/import error leadgen_id=%s: %s", leadgen_id, exc.detail)
                record_integration_event(
                    "Facebook",
                    {"leadgen_id": leadgen_id, "event": event},
                    "graph_error",
                    external_id=leadgen_id,
                    error=str(exc.detail)[:180],
                )
                ignored.append({"leadgen_id": leadgen_id, "reason": str(exc.detail)[:180]})
            except Exception as exc:
                logging.exception("Facebook webhook failed for leadgen_id=%s", leadgen_id)
                record_integration_event(
                    "Facebook",
                    {"leadgen_id": leadgen_id, "event": event, "webhook": body},
                    "error",
                    external_id=leadgen_id,
                    error=str(exc),
                )
                ignored.append({"leadgen_id": leadgen_id, "reason": str(exc)})
    else:
        logging.warning("Facebook webhook: no leadgen_id found in payload; attempting direct body import")
        result = create_integrated_lead(body, "Facebook")
        processed.append({"leadgen_id": None, "status": result.get("status")})
        if result["status"] == "created":
            created.append(result["lead_id"])
        elif result["status"] == "duplicate":
            duplicates.append(result["lead_id"])
        else:
            ignored.append(result)

    response = {
        "status": "success",
        "received": True,
        "leadgen_events": len(events),
        "processed": processed,
        "created": created,
        "duplicates": duplicates,
        "ignored": ignored,
    }
    logging.info("Facebook webhook POST complete: %s", json.dumps(response, default=str))
    return response

def _facebook_resync_pending_impl() -> Dict[str, Any]:
    """Import CRM leads from webhook leadgen IDs that were not yet created."""
    if not FACEBOOK_PAGE_ACCESS_TOKEN:
        raise HTTPException(
            status_code=400,
            detail="FACEBOOK_PAGE_ACCESS_TOKEN is not configured on the server",
        )

    rows = sb_select("integration_events", {
        "source": "eq.Facebook",
        "select": "external_id,status,lead_id,created_at,raw_payload",
        "order": "created_at.desc",
        "limit": "500",
    })
    seen: set[str] = set()
    retried = created = duplicates = ignored = failed = 0
    results: List[Dict[str, Any]] = []

    for evt in rows:
        leadgen_id = clean_text(evt.get("external_id"))
        if not leadgen_id or leadgen_id in seen or leadgen_id in META_FAKE_LEADGEN_IDS:
            continue
        seen.add(leadgen_id)

        if evt.get("status") in ("created", "duplicate") and evt.get("lead_id"):
            existing = sb_select("leads", {"external_lead_id": f"eq.{leadgen_id}", "select": "lead_id", "limit": "1"})
            if existing:
                continue

        if evt.get("status") not in (
            "leadgen_received", "graph_error", "graph_fetched", "ignored", "error", "webhook_received",
        ):
            continue

        page_id = None
        raw = evt.get("raw_payload")
        if isinstance(raw, str):
            try:
                raw = json.loads(raw)
            except Exception:
                raw = {}
        if isinstance(raw, dict):
            event = raw.get("event") if isinstance(raw.get("event"), dict) else {}
            page_id = normalize_meta_page_id(event.get("page_id"), raw.get("page_id"))

        retried += 1
        try:
            payload = fetch_facebook_lead(leadgen_id, page_id)
            result = create_integrated_lead(payload, "Facebook")
            status = result.get("status")
            if status == "created":
                created += 1
            elif status == "duplicate":
                duplicates += 1
            else:
                ignored += 1
            results.append({"leadgen_id": leadgen_id, "status": status, "lead_id": result.get("lead_id")})
        except HTTPException as exc:
            failed += 1
            results.append({"leadgen_id": leadgen_id, "status": "failed", "error": exc.detail})
        except Exception as exc:
            failed += 1
            results.append({"leadgen_id": leadgen_id, "status": "failed", "error": str(exc)[:180]})

    return {
        "retried": retried,
        "created": created,
        "duplicates": duplicates,
        "ignored": ignored,
        "failed": failed,
        "results": results[:50],
    }


@api_router.get("/integrations/facebook/verify")
async def facebook_verify(cu: User = Depends(get_current_user)):
    """Check Meta webhook config, Page token validity, and CRM Meta lead count."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    all_fb = fetch_all_leads_merged("lead_id,source,phone,email,external_lead_id")
    db_meta_real = len([l for l in clean_leads_for_platform_stats(all_fb) if lead_matches_platform(l, "meta")])
    recent = sb_select("integration_events", {
        "source": "eq.Facebook",
        "select": "status,error,external_id,created_at",
        "order": "created_at.desc",
        "limit": "20",
    })
    token_configured = bool(FACEBOOK_PAGE_ACCESS_TOKEN)
    token_authenticated = False
    token_valid = False
    lead_api_ready = False
    token_error: Optional[str] = None
    token_identity: Optional[Dict[str, Any]] = None
    token_scopes: List[str] = []
    token_me_id: Optional[str] = None
    token_type: Optional[str] = None
    forms_count = 0
    forms_summary: List[Dict[str, Any]] = []
    page_id_resolved: Optional[str] = None
    page_id_env_set = bool(clean_text(FACEBOOK_PAGE_ID))
    if not token_configured:
        token_error = "FACEBOOK_PAGE_ACCESS_TOKEN is not set on Render/backend .env"
    else:
        try:
            r = _http.get(
                f"https://graph.facebook.com/{FACEBOOK_GRAPH_VERSION}/me",
                params={"access_token": FACEBOOK_PAGE_ACCESS_TOKEN, "fields": "id,name"},
                timeout=20,
            )
            if r.status_code == 200:
                token_authenticated = True
                me_body = r.json()
                token_identity = me_body if isinstance(me_body, dict) else None
                token_me_id = clean_text((token_identity or {}).get("id"))
                env_page = normalize_meta_page_id(FACEBOOK_PAGE_ID)
                if env_page and token_me_id == env_page:
                    token_type = "page"
                elif env_page and token_me_id and token_me_id != env_page:
                    token_type = "system_user"
                else:
                    token_type = "page" if token_me_id else "unknown"
                try:
                    perm_r = _http.get(
                        f"https://graph.facebook.com/{FACEBOOK_GRAPH_VERSION}/me/permissions",
                        params={"access_token": FACEBOOK_PAGE_ACCESS_TOKEN},
                        timeout=20,
                    )
                    if perm_r.status_code == 200:
                        pdata = perm_r.json()
                        token_scopes = [
                            p.get("permission")
                            for p in (pdata.get("data") or [])
                            if isinstance(p, dict) and p.get("status") == "granted" and p.get("permission")
                        ]
                except Exception:
                    pass
                try:
                    page_id_resolved, page_token = resolve_facebook_page_context(FACEBOOK_PAGE_ID or None)
                    forms = list_facebook_leadgen_forms(page_id_resolved, access_token=page_token)
                    forms_count = len(forms)
                    forms_summary = [
                        {"id": clean_text(f.get("id")), "name": clean_text(f.get("name")), "status": clean_text(f.get("status"))}
                        for f in forms
                        if clean_text(f.get("id"))
                    ]
                    lead_api_ready = True
                    token_valid = True
                except HTTPException as exc:
                    token_error = str(exc.detail)[:280]
                    if token_type == "user":
                        token_error = (
                            "Page Access Token required — you pasted a User token. "
                            "Graph API Explorer → select your Facebook Page → leads_retrieval → Generate token."
                        )
                except Exception as exc:
                    token_error = str(exc)[:280]
            else:
                body = r.json() if r.headers.get("content-type", "").startswith("application/json") else {}
                token_error = (body.get("error") or {}).get("message") or r.text[:220]
                if token_error and "expired" in token_error.lower():
                    token_error = (
                        "Token EXPIRED. Graph API Explorer tokens last ~1 hour only. "
                        "Create a permanent token via Meta Business Settings → System Users (steps in CRM guide)."
                    )
        except Exception as exc:
            token_error = str(exc)[:220]

    pending_ids: set[str] = set()
    for evt in recent:
        lid = clean_text(evt.get("external_id"))
        if is_fake_meta_leadgen_id(lid):
            continue
        if evt.get("status") in ("leadgen_received", "graph_error", "graph_fetched", "error", "webhook_received"):
            pending_ids.add(lid)
    for lid in list(pending_ids):
        if sb_select("leads", {"external_lead_id": f"eq.{lid}", "select": "lead_id", "limit": "1"}):
            pending_ids.discard(lid)

    last_graph_error = next((e for e in recent if e.get("status") == "graph_error"), None)
    has_leads_retrieval = "leads_retrieval" in token_scopes
    return {
        "webhook_url": "/api/facebook/webhook",
        "verify_token": FACEBOOK_VERIFY_TOKEN,
        "token_configured": token_configured,
        "token_authenticated": token_authenticated,
        "token_valid": token_valid,
        "lead_api_ready": lead_api_ready,
        "token_type": token_type,
        "token_me_id": token_me_id,
        "token_error": token_error,
        "token_identity": token_identity,
        "token_scopes": token_scopes,
        "has_leads_retrieval": has_leads_retrieval,
        "page_id_env_set": page_id_env_set,
        "auto_sync_enabled": FACEBOOK_AUTO_SYNC_ENABLED,
        "db_meta_leads": db_meta_real,
        "forms_count": forms_count,
        "forms": forms_summary,
        "page_id": page_id_resolved or normalize_meta_page_id(FACEBOOK_PAGE_ID) or None,
        "pending_webhook_events": len(pending_ids),
        "last_error": last_graph_error.get("error") if last_graph_error else None,
        "recent_events": recent[:5],
        "form_id": FACEBOOK_FORM_ID or None,
        "fix_steps": [
            "PERMANENT TOKEN: business.facebook.com → Settings → System users → Add → Assign your Page → Generate token (never expires)",
            "Permissions: leads_retrieval, pages_show_list, pages_read_engagement, pages_manage_metadata",
            "Paste token in Render FACEBOOK_PAGE_ACCESS_TOKEN + set FACEBOOK_PAGE_ID → redeploy",
            "Do NOT use Graph API Explorer for production — those tokens expire in ~1 hour",
            "After deploy: CRM Integrations → Import Past Meta Leads (90 days)",
        ],
    }

@api_router.get("/integrations/facebook/events")
async def facebook_integration_events(
    limit: int = 20,
    cu: User = Depends(get_current_user),
):
    """Recent Facebook webhook / Graph API events for debugging (admin/manager/marketing)."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    limit = min(max(limit, 1), 100)
    rows = sb_select("integration_events", {
        "source": "eq.Facebook",
        "select": "event_id,source,external_id,status,lead_id,error,created_at",
        "order": "created_at.desc",
        "limit": str(limit),
    })
    return {"events": rows, "count": len(rows)}

@api_router.post("/integrations/facebook/resync")
async def facebook_resync(cu: User = Depends(get_current_user)):
    """Re-fetch Meta leadgen IDs from integration_events and import missing CRM leads."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    return _facebook_resync_pending_impl()


@api_router.post("/integrations/facebook/poll")
async def facebook_poll(cu: User = Depends(get_current_user)):
    """Auto-retry pending Meta webhook leads + import last 24h from Lead Ad forms."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    if not FACEBOOK_PAGE_ACCESS_TOKEN:
        return {"status": "skipped", "reason": "FACEBOOK_PAGE_ACCESS_TOKEN not configured"}

    resync = _facebook_resync_pending_impl()
    import_result: Dict[str, Any] = {"fetched": 0, "created": 0, "duplicates": 0}
    try:
        import_result = _facebook_import_impl(days=1, limit=100)
    except HTTPException as exc:
        import_result = {"error": str(exc.detail)[:180]}
    except Exception as exc:
        import_result = {"error": str(exc)[:180]}

    return {
        "status": "success",
        "resync": resync,
        "import_recent": {
            "fetched": import_result.get("fetched", 0),
            "created": import_result.get("created", 0),
            "duplicates": import_result.get("duplicates", 0),
            "error": import_result.get("error"),
        },
    }

@api_router.post("/integrations/facebook/import")
async def facebook_import(payload: FacebookImportRequest, cu: User = Depends(get_current_user)):
    """Pull previously submitted Meta Lead Ad forms via Graph API and import into CRM."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    if not FACEBOOK_PAGE_ACCESS_TOKEN:
        raise HTTPException(status_code=400, detail="FACEBOOK_PAGE_ACCESS_TOKEN is not configured on the server")
    return _facebook_import_impl(payload.page_id, payload.form_id, payload.days, payload.limit)


def _facebook_import_impl(
    page_id: Optional[str] = None,
    form_id: Optional[str] = None,
    days: int = 90,
    limit: int = 500,
) -> Dict[str, Any]:
    page_id, page_token = resolve_facebook_page_context(page_id)
    days = min(max(days, 1), 365)
    limit = min(max(limit, 1), 1000)
    since_ts = int((now_utc() - timedelta(days=days)).timestamp())

    form_ids: List[str] = []
    forms_meta: List[Dict[str, Any]] = []
    if form_id or FACEBOOK_FORM_ID:
        fid = clean_text(form_id or FACEBOOK_FORM_ID)
        if fid:
            form_ids = [fid]
            forms_meta = [{"id": fid, "name": "Configured form"}]
    else:
        forms = list_facebook_leadgen_forms(page_id, access_token=page_token)
        forms_meta = [{"id": f.get("id"), "name": f.get("name"), "status": f.get("status")} for f in forms]
        form_ids = [clean_text(f.get("id")) for f in forms if clean_text(f.get("id"))]

    if not form_ids:
        raise HTTPException(status_code=404, detail="No Lead Ad forms found for this Facebook Page. Set FACEBOOK_FORM_ID if needed.")

    known_external_ids = _load_facebook_external_ids()
    fetched = created = duplicates = ignored = failed = 0
    results: List[Dict[str, Any]] = []

    for form_id in form_ids:
        try:
            graph_leads = list_facebook_form_leads(form_id, limit=limit, since_ts=since_ts, access_token=page_token)
        except HTTPException as exc:
            failed += 1
            results.append({"form_id": form_id, "status": "failed", "error": exc.detail})
            continue
        fetched += len(graph_leads)
        for graph_lead in graph_leads:
            try:
                result = import_facebook_graph_lead(graph_lead, known_external_ids, quiet=True)
                status = result.get("status")
                if status == "created":
                    created += 1
                elif status == "duplicate":
                    duplicates += 1
                else:
                    ignored += 1
                if len(results) < 50:
                    results.append({
                        "form_id": form_id,
                        "leadgen_id": result.get("leadgen_id"),
                        "status": status,
                        "lead_id": result.get("lead_id"),
                        "reason": result.get("reason"),
                    })
            except Exception as exc:
                failed += 1
                if len(results) < 50:
                    results.append({"form_id": form_id, "leadgen_id": graph_lead.get("id"), "status": "failed", "error": str(exc)[:180]})

    return {
        "page_id": page_id,
        "forms": forms_meta,
        "days": days,
        "fetched": fetched,
        "created": created,
        "duplicates": duplicates,
        "ignored": ignored,
        "failed": failed,
        "results": results,
    }

@api_router.post("/housing/webhook")
async def housing_webhook(request: Request):
    body = await request_payload(request)
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="Invalid Housing.com webhook payload")
    verify_housing_request(request, body)

    created, duplicates, ignored = [], [], []
    for payload in as_list_payload(body):
        merged_payload = {**payload}
        if body.get("integration_uuid") and "integration_uuid" not in merged_payload:
            merged_payload["integration_uuid"] = body.get("integration_uuid")
        result = create_integrated_lead(merged_payload, "Housing.com")
        if result["status"] == "created":
            created.append(result["lead_id"])
        elif result["status"] == "duplicate":
            duplicates.append(result["lead_id"])
        else:
            ignored.append(result)
    return {"status": "success", "source": "Housing.com", "created": created, "duplicates": duplicates, "ignored": ignored}

async def _housing_sync_impl(payload: HousingSyncRequest, cu: User, mode: str = "manual") -> Dict[str, Any]:
    start_date, end_date = housing_sync_window(
        mode, payload.start_date, payload.end_date, payload.allow_historical,
    )
    params = housing_sync_params(start_date, end_date)
    r = _http.get(HOUSING_API_URL, params=params)
    if r.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"Housing.com sync failed: {r.text[:180]}")
    try:
        data = r.json()
    except Exception:
        raise HTTPException(status_code=502, detail="Housing.com returned a non-JSON response")

    lead_payloads = as_list_payload(data)
    if isinstance(data, dict) and not lead_payloads:
        message = clean_text(data.get("message") or data.get("error") or data.get("detail"))
        if message:
            raise HTTPException(status_code=502, detail=f"Housing.com sync failed: {message}")

    created, duplicates, ignored, skipped_stale = [], [], [], []
    for lead_payload in lead_payloads:
        if not should_import_housing_lead_on_sync(lead_payload, start_date, end_date, mode=mode):
            skipped_stale.append(clean_text(pick_first(lead_payload, ["lead_id", "id"])) or "unknown")
            record_integration_event(
                "Housing.com", lead_payload, "skipped_stale",
                external_id=clean_text(pick_first(lead_payload, ["lead_id", "id", "enquiry_id"])),
                error="lead_date outside sync window",
            )
            continue
        result = create_integrated_lead(lead_payload, "Housing.com", actor=cu)
        if result["status"] == "created":
            created.append(result["lead_id"])
        elif result["status"] == "duplicate":
            duplicates.append(result["lead_id"])
        else:
            ignored.append(result)

    record_housing_sync_checkpoint(end_date, {
        "mode": mode,
        "start_date": start_date,
        "fetched": len(lead_payloads),
        "created": len(created),
        "duplicates": len(duplicates),
        "skipped_stale": len(skipped_stale),
    })
    return {
        "status": "success",
        "source": "Housing.com",
        "mode": mode,
        "start_date": start_date,
        "end_date": end_date,
        "fetched": len(lead_payloads),
        "created": created,
        "duplicates": duplicates,
        "ignored": ignored,
        "skipped_stale": len(skipped_stale),
    }


@api_router.post("/housing/sync")
async def housing_sync(payload: HousingSyncRequest, cu: User=Depends(get_current_user)):
    """Pull recent Housing.com leads only (default last 2 hours, not full history)."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    return await _housing_sync_impl(payload, cu, mode="manual")


@api_router.post("/integrations/housing/poll")
async def housing_poll(cu: User=Depends(get_current_user)):
    """Dashboard auto-refresh — only new leads since last poll checkpoint."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    logging.info("Housing poll: auto-sync triggered by %s", cu.email)
    return await _housing_sync_impl(HousingSyncRequest(), cu, mode="poll")


@api_router.post("/integrations/housing/cron")
async def housing_cron_sync(request: Request):
    """External cron hook — same as background auto-sync (optional CRON_SECRET header)."""
    if CRON_SECRET:
        provided = (
            request.headers.get("x-cron-secret")
            or request.headers.get("authorization", "").removeprefix("Bearer ").strip()
            or request.query_params.get("secret")
        )
        if provided != CRON_SECRET:
            raise HTTPException(status_code=401, detail="Invalid cron secret")
    result = await asyncio.to_thread(run_housing_sync_background, "cron")
    logging.info("Housing cron sync: %s", result.get("status"))
    return result

@api_router.get("/integrations/housing/verify")
async def housing_verify(cu: User=Depends(get_current_user)):
    """Confirm Housing.com credentials and whether the pull API returns leads."""
    ensure_roles(cu, ["admin", "manager", "marketing"])
    credentials_ok = bool(HOUSING_PROFILE_ID and HOUSING_ENCRYPTION_KEY)
    if not credentials_ok:
        return {
            "credentials_ok": False,
            "api_reachable": False,
            "leads_available": 0,
            "message": "Set HOUSING_PROFILE_ID and HOUSING_ENCRYPTION_KEY in backend .env",
            "db_housing_leads": 0,
        }

    end_date = int(now_utc().timestamp())
    start_date = end_date - HOUSING_MANUAL_DEFAULT_WINDOW_SEC
    params = housing_sync_params(start_date, end_date)
    try:
        r = _http.get(HOUSING_API_URL, params=params, timeout=30)
        api_reachable = r.status_code < 500
        data = r.json() if api_reachable else {}
        lead_payloads = as_list_payload(data) if api_reachable else []
        message = None
        if isinstance(data, dict) and not lead_payloads:
            message = clean_text(data.get("message") or data.get("error"))
    except Exception as exc:
        return {
            "credentials_ok": True,
            "api_reachable": False,
            "leads_available": 0,
            "message": str(exc),
            "db_housing_leads": len(sb_select("leads", {"source": "eq.Housing.com", "select": "lead_id"})),
        }

    db_housing = sb_select("leads", {"source": "eq.Housing.com", "select": "lead_id"})
    return {
        "credentials_ok": True,
        "api_reachable": api_reachable and r.status_code < 400,
        "api_status": r.status_code,
        "leads_available": len(lead_payloads),
        "message": message,
        "sample_lead": lead_payloads[0] if lead_payloads else None,
        "db_housing_leads": len(db_housing),
    }

@api_router.get("/integrations/status")
async def integrations_status(cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    last_sync_end = get_last_housing_sync_end_epoch()
    db_housing = sb_select("leads", {"source": "eq.Housing.com", "select": "lead_id"})
    return {
        "facebook": {
            "source": "Facebook",
            "webhook_path": "/api/facebook/webhook",
            "verify_token_configured": bool(FACEBOOK_VERIFY_TOKEN),
            "lead_retrieval_configured": bool(FACEBOOK_PAGE_ACCESS_TOKEN),
        },
        "housing": {
            "source": "Housing.com",
            "webhook_path": "/api/housing/webhook",
            "sync_path": "/api/housing/sync",
            "poll_path": "/api/integrations/housing/poll",
            "cron_path": "/api/integrations/housing/cron",
            "profile_id_configured": bool(HOUSING_PROFILE_ID),
            "encryption_key_configured": bool(HOUSING_ENCRYPTION_KEY),
            "integration_uuid_configured": bool(HOUSING_INTEGRATION_UUID or HOUSING_WEBHOOK_SECRET),
            "api_url": HOUSING_API_URL,
            "auto_sync_enabled": HOUSING_AUTO_SYNC_ENABLED,
            "auto_sync_interval_sec": HOUSING_AUTO_SYNC_INTERVAL_SEC,
            "last_sync_end_epoch": last_sync_end,
            "last_sync_at": datetime.fromtimestamp(last_sync_end, tz=timezone.utc).isoformat() if last_sync_end else None,
            "db_housing_leads": len(db_housing),
        },
    }

@api_router.post("/webhooks/{source}")
async def incoming_webhook(source: str, request: Request):
    """
    Compatibility endpoint for portal webhooks.
    Prefer /api/housing/webhook and /api/facebook/webhook for first-class integrations.
    """
    body = await request_payload(request)
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="Invalid webhook payload")
    normalized_source = {
        "facebook": "Facebook",
        "housing": "Housing.com",
        "housing.com": "Housing.com",
        "99acres": "99acres",
        "magicbricks": "MagicBricks",
        "website": "Website",
        "whatsapp": "WhatsApp",
    }.get(source.lower(), source)
    result = create_integrated_lead(body, normalized_source)
    return {"status": "success", "source": normalized_source, **result}

@api_router.post("/webhooks/whatsapp/reply")
async def inbound_whatsapp_reply(request: Request):
    """
    Handle incoming WhatsApp replies from Interakt.
    """
    # TEMPORARY: WhatsApp Webhook Automation is temporarily disabled
    return {"status": "disabled", "message": "WhatsApp Automation is temporarily disabled"}


def _normalize_person_name(value: Any) -> str:
    text = re.sub(r"[^\w\s]", " ", str(value or ""))
    return re.sub(r"\s+", " ", text.strip().lower())


_NAME_HONORIFICS = frozenset({"mr", "mrs", "ms", "miss", "dr", "shri", "smt", "sir", "madam"})


def _name_tokens(value: str) -> List[str]:
    tokens = [t for t in _normalize_person_name(value).split() if len(t) > 1]
    while tokens and tokens[0] in _NAME_HONORIFICS:
        tokens = tokens[1:]
    return tokens


def map_import_headers(headers: List[Any]) -> Dict[str, int]:
    """Map spreadsheet headers — supports Umang Excel template + legacy columns."""
    header_mapping: Dict[str, int] = {}
    for idx, h in enumerate(headers):
        if h is None or str(h).strip() == "":
            continue
        h_clean = re.sub(r"\s+", " ", str(h).strip().lower().replace("_", " "))
        h_compact = h_clean.replace(" ", "")

        if (
            "assign_to" not in header_mapping
            and (
                h_compact in ("assignto", "assigneto", "assignedto", "assignee")
                or (re.search(r"assign", h_clean) and re.search(r"\bto\b|employee|telecaller|agent|executive", h_clean))
                or h_clean in ("assign to", "assigne to", "assigned to", "employee name")
            )
        ):
            header_mapping["assign_to"] = idx
        elif any(term in h_clean for term in ("lead date", "enquiry date", "inquiry date")) or h_clean == "date":
            header_mapping["lead_date"] = idx
        elif "lead name" in h_clean or h_clean in ("name", "full name", "customer name"):
            header_mapping["name"] = idx
        elif any(term in h_clean for term in ("phone", "mobile", "contact")):
            header_mapping["phone"] = idx
        elif any(term in h_clean for term in ("email", "mail")):
            header_mapping["email"] = idx
        elif any(term in h_clean for term in ("price", "budget", "amount")):
            header_mapping["budget"] = idx
        elif any(term in h_clean for term in ("locality", "location", "address", "area")):
            header_mapping["location"] = idx
        elif any(term in h_clean for term in ("configuration", "config", "bhk", "property type", "requirement")):
            header_mapping["property_type"] = idx
        elif any(term in h_clean for term in (
            "building/project", "building project", "project name", "building name", "society", "tower",
        )) or (("building" in h_clean or "project" in h_clean) and "name" in h_clean):
            header_mapping["preferred_property"] = idx
        elif any(term in h_clean for term in ("notes", "remarks", "comments")):
            header_mapping["notes"] = idx
        elif any(term in h_clean for term in ("source", "lead source", "platform")):
            header_mapping["source"] = idx
        elif "name" in h_clean and not any(
            term in h_clean for term in ("building", "project", "property", "assign", "employee")
        ):
            header_mapping.setdefault("name", idx)
    return header_mapping


def build_employee_assign_lookup() -> Tuple[Dict[str, str], List[Dict[str, str]]]:
    """Name/email → employee_id lookup (includes inactive=null as active)."""
    emps = sb_select_all("employees", {"select": "employee_id,name,email,active", "order": "name.asc"})
    emps = [e for e in emps if e.get("active", True) is not False]
    if not emps:
        emps = sb_select_all("employees", {"select": "employee_id,name,email,active", "order": "name.asc"})

    exact: Dict[str, str] = {}
    roster: List[Dict[str, str]] = []
    for emp in emps:
        eid = emp.get("employee_id")
        if not eid:
            continue
        display = (emp.get("name") or "").strip()
        norm = _normalize_person_name(display)
        row = {"employee_id": eid, "name": display, "norm": norm}
        roster.append(row)
        if norm:
            exact[norm] = eid
            parts = norm.split()
            if len(parts) >= 2:
                exact.setdefault(f"{parts[0]} {parts[-1]}", eid)
            if parts:
                exact.setdefault(parts[-1], eid)
        email = (emp.get("email") or "").strip().lower()
        if email:
            exact[email] = eid
    return exact, roster


def resolve_import_assignee(
    assignee_raw: str,
    exact: Dict[str, str],
    roster: List[Dict[str, str]],
) -> Tuple[Optional[str], str]:
    """Return (employee_id, cleaned_assignee_label)."""
    raw = import_cell_text(assignee_raw)
    if not raw:
        return None, raw
    roster_ids = {r["employee_id"] for r in roster}
    if raw.startswith("emp_") and raw in roster_ids:
        name = next((r["name"] for r in roster if r["employee_id"] == raw), raw)
        return raw, name
    key = _normalize_person_name(raw)
    if not key:
        return None, raw
    if key in exact:
        eid = exact[key]
        name = next((r["name"] for r in roster if r["employee_id"] == eid), raw)
        return eid, name
    key_tokens = _name_tokens(raw)
    if key_tokens:
        joined = " ".join(key_tokens)
        if joined in exact:
            eid = exact[joined]
            name = next((r["name"] for r in roster if r["employee_id"] == eid), raw)
            return eid, name

    best_id: Optional[str] = None
    best_score = 0
    for emp in roster:
        norm = emp["norm"]
        if not norm:
            continue
        if norm == key or norm in key or key in norm:
            return emp["employee_id"], emp["name"]
        emp_tokens = _name_tokens(emp["name"])
        if not emp_tokens:
            continue
        overlap = sum(1 for t in emp_tokens if t in key_tokens)
        if overlap >= len(emp_tokens) or (len(emp_tokens) == 1 and overlap == 1):
            if overlap > best_score:
                best_score = overlap
                best_id = emp["employee_id"]
        elif len(key_tokens) == 1 and key_tokens[0] in emp_tokens:
            if 1 > best_score:
                best_score = 1
                best_id = emp["employee_id"]
    if best_id:
        name = next((r["name"] for r in roster if r["employee_id"] == best_id), raw)
        return best_id, name
    return None, raw


def import_cell_text(val: Any) -> str:
    """Normalize Excel/CSV cell to plain text (handles numbers, dates)."""
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%d/%m/%Y")
    if isinstance(val, float):
        if val == int(val) and abs(val) < 1e15:
            return str(int(val))
    if isinstance(val, int):
        return str(val)
    return str(val).strip()


def locate_import_header(rows: List[Any]) -> Tuple[int, Dict[str, int]]:
    """Find header row — supports files with title rows above column names."""
    for i, row in enumerate(rows[:20]):
        if not row:
            continue
        row_list = list(row) if not isinstance(row, list) else row
        if all(c is None or str(c).strip() == "" for c in row_list):
            continue
        header_map = map_import_headers(row_list)
        if "name" in header_map and "phone" in header_map:
            return i, header_map
    raise HTTPException(
        status_code=400,
        detail="Could not find header row with 'Lead Name' and 'Phone Number' columns.",
    )


def parse_import_lead_date(value: Any) -> Optional[str]:
    """Parse Excel/CSV dates — e.g. 30/12/2025, ISO, or Excel datetime."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        dt = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return dt.isoformat()
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            dt = from_excel(value)
            if isinstance(dt, datetime):
                return (dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)).isoformat()
        except Exception:
            pass
    s = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y-%m-%d", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=timezone.utc).isoformat()
        except ValueError:
            continue
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.isoformat()
    except ValueError:
        return None


def import_row_to_record(row: List[Any], header_map: Dict[str, int]) -> Dict[str, str]:
    def raw_val(key: str) -> Any:
        if key not in header_map or header_map[key] >= len(row):
            return ""
        return row[header_map[key]]

    def text_val(key: str) -> str:
        return import_cell_text(raw_val(key))

    lead_date_raw = raw_val("lead_date") if "lead_date" in header_map else None
    parsed_date = parse_import_lead_date(lead_date_raw) if lead_date_raw not in (None, "") else None

    return {
        "name": text_val("name"),
        "phone": text_val("phone"),
        "email": text_val("email"),
        "budget": text_val("budget"),
        "location": text_val("location"),
        "property_type": text_val("property_type"),
        "preferred_property": text_val("preferred_property"),
        "notes": text_val("notes"),
        "assign_to": text_val("assign_to"),
        "lead_date": parsed_date or "",
        "source": text_val("source"),
    }


def normalize_import_source(raw: str) -> str:
    """Map spreadsheet Source column to CRM source (Meta → Facebook for platform stats)."""
    s = clean_text(raw).lower()
    if not s:
        return "bulk_import"
    if any(t in s for t in ("facebook", "meta", "instagram", "lead ad", "leadads")):
        return "Facebook"
    if "housing" in s:
        return "Housing.com"
    return raw.strip() or "bulk_import"


@api_router.post("/leads/import")
async def import_leads(file: UploadFile = File(...), cu: User = Depends(get_current_user)):
    filename = file.filename.lower()
    
    # Supported files: .csv, .xlsx, .xls
    if not (filename.endswith(".csv") or filename.endswith(".xlsx") or filename.endswith(".xls")):
        raise HTTPException(status_code=400, detail="Only CSV and Excel files (.xlsx, .xls) are supported.")
    
    contents = await file.read()
    records: List[Dict[str, str]] = []

    if filename.endswith(".csv"):
        try:
            text = contents.decode("utf-8-sig")
            csv_reader = csv.reader(io.StringIO(text))
            rows = list(csv_reader)
            if not rows:
                raise HTTPException(status_code=400, detail="CSV file is empty.")
            
            header_idx, header_map = locate_import_header(rows)
                
            for row in rows[header_idx + 1:]:
                if not row or all(not str(c or "").strip() for c in row):
                    continue
                records.append(import_row_to_record(row, header_map))
        except HTTPException as he:
            raise he
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse CSV: {str(e)}")

    else: # Excel
        try:
            wb = openpyxl.load_workbook(filename=BytesIO(contents), data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                raise HTTPException(status_code=400, detail="Excel file is empty.")
            
            header_idx, header_map = locate_import_header(rows)
                
            for row in rows[header_idx + 1:]:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue
                records.append(import_row_to_record(list(row), header_map))
        except HTTPException as he:
            raise he
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse Excel: {str(e)}")

    employee_exact, employee_roster = build_employee_assign_lookup()
    if not employee_roster:
        logging.warning("Excel import: no employees found for auto-assign lookup")
    actor_id = cu.acting_as_employee_id or cu.user_id

    imported_count = 0
    skipped_count = 0
    assigned_count = 0
    imported_leads: List[Dict[str, Any]] = []
    assign_failed: List[Dict[str, str]] = []
    assignment_breakdown: Dict[str, Dict[str, Any]] = {}
    employees_touched: set = set()
    has_assign_column = any(bool(r.get("assign_to", "").strip()) for r in records)

    for r in records:
        name = r["name"].strip()
        phone = normalize_phone(r["phone"])
        
        if not name or not phone:
            skipped_count += 1
            continue

        lid = gen_id("lead")
        now = now_utc().isoformat()
        lead_created = r.get("lead_date") or now
        
        pref_prop = r.get("preferred_property", "").strip()
        r_notes = r.get("notes", "").strip()
        if pref_prop:
            lead_notes = f"Building/Project: {pref_prop}" + (f"\n{r_notes}" if r_notes else "")
        else:
            lead_notes = r_notes

        assignee_name = r.get("assign_to", "").strip()
        emp_id, matched_emp_name = resolve_import_assignee(assignee_name, employee_exact, employee_roster) if assignee_name else (None, "")

        lead: Dict[str, Any] = {
            "lead_id": lid,
            "name": name,
            "phone": phone,
            "email": r.get("email", ""),
            "budget": r.get("budget", ""),
            "location": r.get("location", ""),
            "property_type": r.get("property_type", ""),
            "notes": lead_notes,
            "source": normalize_import_source(r.get("source", "")),
            "stage": "new",
            "status": "active",
            "created_at": lead_created,
            "external_created_at": r.get("lead_date") or None,
            "updated_at": now,
        }

        if emp_id:
            lead.update(stamp_lead_assignment(emp_id, assigned_by=actor_id, current_stage="new"))

        result = sb_insert("leads", lead)
        if not result:
            skipped_count += 1
            if assignee_name and not emp_id:
                assign_failed.append({
                    "name": name,
                    "assign_to": assignee_name,
                    "reason": "insert_failed",
                })
            continue

        saved = enrich_lead_display_fields(result)
        SESSION_CACHE["leads"].insert(0, saved)

        if emp_id:
            if saved.get("assigned_to") != emp_id:
                saved = assign_lead_to_employee(lid, saved, emp_id, cu, reactivate=False)
                saved = enrich_lead_display_fields(saved)
            if saved.get("assigned_to") == emp_id:
                log_activity(cu, "bulk_import_assign", f"Excel import assigned {name} → {matched_emp_name}", lead_id=lid)
                employees_touched.add(emp_id)
                assigned_count += 1
                bucket = assignment_breakdown.setdefault(emp_id, {
                    "employee_id": emp_id,
                    "employee_name": matched_emp_name,
                    "count": 0,
                })
                bucket["count"] += 1
            else:
                assign_failed.append({
                    "name": name,
                    "assign_to": assignee_name,
                    "reason": "assign_failed",
                })
        elif assignee_name:
            assign_failed.append({
                "name": name,
                "assign_to": assignee_name,
                "reason": "employee_not_found",
            })

        log_activity(cu, "bulk_import", f"Lead imported via Excel/CSV: {name}", lead_id=lid)
        imported_leads.append(saved)
        imported_count += 1

    for eid in employees_touched:
        sb_update("employees", "employee_id", eid, {"last_assigned_at": now_utc().isoformat()})

    invalidate_leads_cache()
    invalidate_employees_cache()

    breakdown_list = sorted(
        assignment_breakdown.values(),
        key=lambda x: (-int(x.get("count") or 0), str(x.get("employee_name") or "")),
    )

    return {
        "status": "success",
        "imported_count": imported_count,
        "skipped_count": skipped_count,
        "assigned_count": assigned_count,
        "assign_failed_count": len(assign_failed),
        "assign_failed": assign_failed[:20],
        "assignment_breakdown": breakdown_list,
        "has_assign_column": has_assign_column,
        "available_employees": [r["name"] for r in employee_roster],
        "leads": imported_leads[:10],
    }

def _admin_can_reset(cu: User) -> bool:
    return cu.role == "admin" or cu.email == "htshpatil13@gmail.com"


def _sb_bulk_patch(table: str, match_params: str, data: Dict[str, Any]) -> Optional[str]:
    """PATCH rows matching PostgREST filter. Returns error text or None."""
    h = {**sb_headers(), "Prefer": "return=minimal"}
    r = _http.patch(f"{sb_url(table)}?{match_params}", headers=h, json=data)
    if r.status_code >= 400:
        return f"{table}: {r.status_code} {r.text[:200]}"
    return None


@api_router.post("/leads/reset-assignments")
async def reset_all_assignments(cu: User = Depends(get_current_user)):
    """Full clean: wipe loans/bookings/follow-ups, unassign all leads, zero employee stats."""
    if not _admin_can_reset(cu):
        raise HTTPException(status_code=403, detail="Only admins can reset assignments.")

    errors: List[str] = []
    reset_payload = {
        "assigned_to": None,
        "assigned_at": None,
        "assigned_by": None,
        "follow_up_at": None,
        "stage": "new",
        "status": "active",
        "priority": None,
        "call_status": None,
        "updated_at": now_utc().isoformat(),
    }
    err = _sb_bulk_patch("leads", "stage=neq.broker", reset_payload)
    if err:
        errors.append(err)

    for table in [
        "notifications", "customers", "lead_notes",
        "visit_followups", "visits", "bookings", "loans", "activities",
    ]:
        try:
            r = _http.delete(f"{sb_url(table)}?created_at=not.is.null", headers=sb_headers())
            if r.status_code >= 400 and r.status_code != 404:
                errors.append(f"{table}: {r.status_code}")
        except Exception as exc:
            errors.append(f"{table}: {exc}")

    employees = sb_select("employees", {"select": "employee_id"})
    for e in employees:
        eid = e.get("employee_id")
        if eid:
            sb_update("employees", "employee_id", eid, {
                "leads_assigned": 0,
                "leads_closed": 0,
                "performance": 0,
                "updated_at": now_utc().isoformat(),
            })

    for key in SESSION_CACHE:
        SESSION_CACHE[key] = []

    if errors:
        raise HTTPException(status_code=500, detail="; ".join(errors))
    return {
        "status": "success",
        "message": "Full clean done: loans, bookings, follow-ups and assignments cleared. All leads are unassigned and new.",
    }


@api_router.delete("/leads/clear-all")
async def clear_all_leads(cu: User = Depends(get_current_user)):
    # Verify the user is admin
    if not _admin_can_reset(cu):
        raise HTTPException(status_code=403, detail="Only admins can delete all leads.")
        
    tables_to_wipe = ["notifications", "customers", "lead_notes", "visit_followups", "visits", "bookings", "loans", "activities", "leads"]
    optional_tables = {"notifications", "customers", "lead_notes", "visit_followups"}
    
    # We clear them from Supabase
    import httpx
    errors = []
    for table in tables_to_wipe:
        try:
            r = httpx.delete(
                f"{SUPABASE_URL}/rest/v1/{table}?created_at=not.is.null",
                headers=sb_headers()
            )
            if r.status_code >= 400:
                if r.status_code == 404 and table in optional_tables:
                    continue
                errors.append(f"Failed to wipe {table}: {r.text}")
        except Exception as e:
            errors.append(f"Failed to wipe {table} due to exception: {str(e)}")
            
    # Clear local session cache
    for table in tables_to_wipe:
        if table in SESSION_CACHE:
            SESSION_CACHE[table] = []
    SESSION_CACHE["followups"] = []
    
    if errors:
        raise HTTPException(status_code=500, detail="; ".join(errors))
        
    return {"status": "success", "message": "All leads and associated records have been wiped."}

@api_router.post("/leads")
async def create_lead(p: LeadCreatePublic, cu: User=Depends(get_current_user)):
    phone = normalize_phone(p.phone)
    if not phone:
        raise HTTPException(status_code=400, detail="A valid phone number is required.")
    lid = gen_id("lead")
    lead = {
        "lead_id": lid, "name": p.name, "phone": phone, "email": p.email,
        "budget": p.budget, "location": p.location, "property_type": p.property_type,
        "notes": p.notes, "source": p.source or "manual_entry", "stage": "new", "status": "active",
        "assigned_to": None, "created_at": now_utc().isoformat(), "updated_at": now_utc().isoformat(),
    }
    if p.starred is not None: lead["starred"] = p.starred
    result = sb_insert("leads", lead)
    # Always add to session cache for immediate responsiveness
    SESSION_CACHE["leads"].insert(0, result or lead)
    invalidate_leads_cache()
    log_activity(cu, "manual_enquiry", f"Manual lead entry created for {p.name}.", lead_id=lid)
    return result or lead

@api_router.get("/leads")
async def list_leads(
    stage: Optional[str]=None,
    status_: Optional[str]=None,
    assigned_to: Optional[str]=None,
    source: Optional[str]=None,
    q: Optional[str]=None,
    exclude_broker: bool = True,
    broker_only: bool = False,
    limit: int=200,
    offset: int=0,
    cu: User=Depends(get_current_user),
):
    limit = min(max(limit, 1), 500)
    offset = max(offset, 0)
    list_fields = (
        "lead_id,name,phone,email,source,stage,status,lead_type,priority,call_status,"
        "budget,location,property_type,assigned_to,assigned_at,follow_up_at,created_at,updated_at,"
        "external_lead_id,starred"
    )
    params: Dict[str, str] = {"select": list_fields, "order": "created_at.desc"}
    if stage: params["stage"] = f"eq.{stage}"
    if status_: params["status"] = f"eq.{status_}"
    if assigned_to: params["assigned_to"] = f"eq.{assigned_to}"
    if source: params["source"] = f"ilike.*{source}*"

    # Paginate at the database — never pull the entire leads table for a list view.
    db_fetch = min(limit + offset + 50, 500)
    params["limit"] = str(db_fetch)
    leads = sb_select("leads", params)
    
    # Filter session cache to match the query parameters
    filtered_cache = []
    for l in SESSION_CACHE["leads"]:
        match = True
        if stage and l.get("stage") != stage: match = False
        if status_ and l.get("status") != status_: match = False
        if assigned_to and l.get("assigned_to") != assigned_to: match = False
        if source and source.lower() not in str(l.get("source", "")).lower(): match = False
        if match:
            filtered_cache.append(l)

    # Merge with session cache (deduplicate: cache wins)
    cache_ids = {l.get("lead_id") for l in filtered_cache}
    db_only = [l for l in leads if l.get("lead_id") not in cache_ids]
    all_leads = filtered_cache + db_only

    if q:
        needle = q.lower().strip()
        all_leads = [
            l for l in all_leads
            if any(needle in str(l.get(k, "")).lower() for k in ["name", "phone", "email", "location", "source", "property_type"])
        ]

    if broker_only:
        all_leads = [l for l in all_leads if is_broker_pool_lead(l)]
    elif exclude_broker:
        all_leads = [l for l in all_leads if is_pipeline_lead(l)]

    if not all_leads and not stage and not status_ and not assigned_to:
        return []
    return all_leads[offset:offset + limit]


@api_router.get("/leads/booking-queue")
async def list_booking_queue(limit: int = 100, cu: User = Depends(get_current_user)):
    """Hot / handoff leads for New Booking — small payload, no full-table scan."""
    limit = min(max(limit, 1), 200)
    select = "lead_id,name,phone,source,stage,status,priority,assigned_to,created_at,budget,location"
    rows = sb_select("leads", {
        "select": select,
        "status": "neq.negative",
        "or": "(priority.eq.hot,priority.eq.handoff_booking,stage.eq.booking)",
        "order": "created_at.desc",
        "limit": str(limit + 30),
    })
    cache_slice = [l for l in SESSION_CACHE["leads"] if lead_ready_for_booking_queue(l)]
    cache_ids = {l.get("lead_id") for l in cache_slice}
    merged = cache_slice + [l for l in rows if l.get("lead_id") not in cache_ids]
    booked = {b.get("lead_id") for b in sb_select("bookings", {"select": "lead_id"})}
    booked |= {b.get("lead_id") for b in SESSION_CACHE["bookings"]}
    queue = [l for l in merged if lead_ready_for_booking_queue(l) and l.get("lead_id") not in booked]
    queue.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    return queue[:limit]


@api_router.get("/leads/loan-queue")
async def list_loan_queue(limit: int = 100, cu: User = Depends(get_current_user)):
    """Hot / handoff leads for New Loan Application."""
    limit = min(max(limit, 1), 200)
    select = "lead_id,name,phone,source,stage,status,priority,assigned_to,created_at,budget,location"
    rows = sb_select("leads", {
        "select": select,
        "status": "neq.negative",
        "or": "(priority.eq.hot,priority.eq.handoff_loan,stage.eq.loan)",
        "order": "created_at.desc",
        "limit": str(limit + 30),
    })
    cache_slice = [l for l in SESSION_CACHE["leads"] if lead_ready_for_loan_queue(l)]
    cache_ids = {l.get("lead_id") for l in cache_slice}
    merged = cache_slice + [l for l in rows if l.get("lead_id") not in cache_ids]
    loaned = {ln.get("lead_id") for ln in sb_select("loans", {"select": "lead_id"})}
    loaned |= {ln.get("lead_id") for ln in SESSION_CACHE["loans"]}
    queue = [l for l in merged if lead_ready_for_loan_queue(l) and l.get("lead_id") not in loaned]
    queue.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    return queue[:limit]


@api_router.get("/broker-leads")
async def list_broker_leads(cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    all_leads = fetch_all_leads_merged()
    broker_leads = [l for l in all_leads if is_broker_pool_lead(l)]
    return {"total": len(broker_leads), "leads": broker_leads}

class BrokerActivateRequest(BaseModel):
    assigned_to: Optional[str] = None
    brokerage_amount: Optional[float] = None

@api_router.post("/leads/{lead_id}/to-broker")
async def move_lead_to_broker(lead_id: str, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager"])
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*"})
    lead = leads[0] if leads else next((l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id), None)
    if not lead:
        raise HTTPException(404, "Lead not found")
    data = {
        "stage": "broker",
        "lead_type": "brokerage",
        "assigned_to": None,
        "updated_at": now_utc().isoformat(),
    }
    updated = sb_update("leads", "lead_id", lead_id, data) or {**lead, **data}
    update_cached_lead(lead_id, data)
    log_activity(cu, "broker_pool", f"Lead moved to broker pool: {lead.get('name')}", lead_id=lead_id)
    return updated

@api_router.post("/leads/{lead_id}/from-broker")
async def activate_lead_from_broker(lead_id: str, body: BrokerActivateRequest, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager"])
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*"})
    lead = leads[0] if leads else next((l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id), None)
    if not lead:
        raise HTTPException(404, "Lead not found")
    assigned_to = clean_text(body.assigned_to)
    data: Dict[str, Any] = {
        "lead_type": "standard",
        "updated_at": now_utc().isoformat(),
    }
    if assigned_to:
        data.update(stamp_lead_assignment(assigned_to, assigned_by=cu.acting_as_employee_id or cu.user_id, current_stage=lead.get("stage")))
    else:
        data["stage"] = "new"
        data["assigned_to"] = None
    if body.brokerage_amount is not None:
        data["brokerage_amount"] = body.brokerage_amount
    updated = sb_update("leads", "lead_id", lead_id, data) or {**lead, **data}
    update_cached_lead(lead_id, data)
    log_activity(cu, "broker_activated", f"Broker lead activated and assigned", lead_id=lead_id)
    if assigned_to:
        create_notification(assigned_to, "Broker lead assigned", f"{lead.get('name')} is ready to work.", lead_id=lead_id)
    return updated

LEAD_BUCKET_KEYS = ["all", "new_today", "positive", "not_interested", "registration", "booking", "follow_up"]


def filter_lead_bucket(all_leads: List[Dict[str, Any]], bucket_key: str, today: str) -> List[Dict[str, Any]]:
    """Single source of truth for dashboard bucket lists + counts.
    Always deduped so the metric box number == the opened list length."""
    if bucket_key == "new_today":
        filtered = [
            l for l in all_leads
            if is_unassigned_new_lead(l) and (l.get("created_at") or "")[:10] == today
        ]
    elif bucket_key == "not_interested":
        filtered = [l for l in all_leads if l.get("status") == "negative"]
    elif bucket_key == "positive":
        filtered = [l for l in all_leads if l.get("stage") in ["positive", "site_visit", "booking", "loan", "registration", "closed"] and l.get("status") != "negative"]
    elif bucket_key == "registration":
        filtered = [l for l in all_leads if l.get("stage") == "registration"]
    elif bucket_key == "booking":
        filtered = [l for l in all_leads if l.get("stage") in ["booking", "loan"] and l.get("status") != "negative"]
    elif bucket_key == "follow_up":
        filtered = [l for l in all_leads if l.get("follow_up_at") and l.get("status") != "negative"]
    else:
        # Pipeline total — same scope as platform breakdown / Total Leads modal.
        filtered = clean_leads_for_platform_stats(all_leads)
    if bucket_key != "all":
        filtered = dedupe_leads(filtered)
    filtered.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    return filtered


@api_router.get("/leads/filtered")
async def list_leads_filtered(
    bucket: str = "all",
    limit: int = 500,
    assigned_to: Optional[str] = None,
    cu: User = Depends(get_current_user),
):
    """Dashboard drill-down lists: all, new_today, positive, not_interested, registration, booking, follow_up."""
    bucket_key = (bucket or "all").strip().lower()
    if bucket_key not in set(LEAD_BUCKET_KEYS):
        raise HTTPException(400, detail=f"bucket must be one of: {', '.join(LEAD_BUCKET_KEYS)}")

    limit = min(max(limit, 1), 500)
    today = now_utc().date().isoformat()
    all_leads = fetch_all_leads_merged()
    filtered = filter_lead_bucket(all_leads, bucket_key, today)
    assignee = (assigned_to or "").strip()
    if assignee:
        filtered = [l for l in filtered if l.get("assigned_to") == assignee]
    elif cu.role not in ["admin", "manager"]:
        emp_id = cu.acting_as_employee_id or cu.employee_id
        if emp_id and bucket_key == "follow_up":
            filtered = [l for l in filtered if l.get("assigned_to") == emp_id]
    if bucket_key == "follow_up" and filtered:
        emps = sb_select("employees", {"select": "employee_id,name"})
        emp_map = {e.get("employee_id"): e.get("name") for e in emps if e.get("employee_id")}
        for lead in filtered:
            eid = lead.get("assigned_to")
            if eid and emp_map.get(eid):
                lead["employee_name"] = emp_map[eid]
    return {"bucket": bucket_key, "total": len(filtered), "leads": filtered[:limit]}


@api_router.get("/leads/workspace")
async def workspace_leads(
    limit: int = 500,
    cu: User = Depends(get_current_user),
):
    """Telecaller / Sales Executive workspace — queue + follow-ups with counts that
    match Assign Leads stats and My Dashboard KPIs."""
    emp_id = resolve_employee_id(cu) or cu.acting_as_employee_id or cu.employee_id
    if not emp_id and cu.role not in ["admin", "manager"]:
        raise HTTPException(403, detail="Employee profile required for workspace")

    limit = min(max(limit, 1), 500)
    select = (
        "lead_id,name,phone,email,source,stage,status,priority,call_status,assigned_to,"
        "follow_up_at,created_at,budget,location,property_type,raw_payload"
    )
    if cu.role in ["admin", "manager"] and not cu.acting_as_employee_id:
        all_leads = fetch_all_leads_merged(select)
        emp_leads = all_leads
        role = "telecaller"
    else:
        emp_leads = fetch_employee_assigned_leads(cu, select)
        emp_id = resolve_employee_id(cu) or cu.acting_as_employee_id or cu.employee_id
        emps = sb_select("employees", {"employee_id": f"eq.{emp_id}", "select": "role", "limit": "1"}) if emp_id else []
        role = emps[0].get("role") if emps else cu.role

    queue = filter_employee_queue_leads(emp_leads, role)
    follow_ups = filter_employee_follow_up_leads(emp_leads)
    stats = compute_employee_assignment_stats(emp_leads, role)
    queue.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    follow_ups.sort(key=lambda l: l.get("follow_up_at") or "", reverse=True)
    return {
        "employee_id": emp_id,
        "role": role,
        "stats": stats,
        "queue": {"total": len(queue), "leads": enrich_leads_display_fields(queue[:limit])},
        "follow_ups": {"total": len(follow_ups), "leads": enrich_leads_display_fields(follow_ups[:limit])},
    }


@api_router.get("/stats/lead-buckets")
async def stats_lead_buckets(cu: User = Depends(get_current_user)):
    """Counts for every dashboard metric box, computed with the EXACT same
    filter+dedupe as /leads/filtered so each box number matches its opened list."""
    today = now_utc().date().isoformat()
    all_leads = fetch_all_leads_merged(
        "lead_id,name,phone,email,external_lead_id,source,stage,status,lead_type,follow_up_at,created_at"
    )
    return {key: len(filter_lead_bucket(all_leads, key, today)) for key in LEAD_BUCKET_KEYS}


@api_router.get("/leads/assign-queue")
async def list_assign_queue(cu: User = Depends(get_current_user)):
    """Backward-compatible: unassigned leads only."""
    ensure_roles(cu, ["admin", "manager"])
    data = await list_assign_workspace(
        inquiry_status="unassigned",
        source="all",
        assigned_to="unassigned",
        q="",
        limit=500,
        offset=0,
        cu=cu,
    )
    return {
        "total": data["total"],
        "leads": data["leads"],
        "employees": data["employees"],
        "auto_assign_on_intake": False,
    }


@api_router.get("/leads/assign-workspace")
async def list_assign_workspace(
    inquiry_status: str = "all",
    source: str = "all",
    assigned_to: str = "all",
    q: str = "",
    location: str = "",
    limit: int = 500,
    offset: int = 0,
    cu: User = Depends(get_current_user),
):
    """Manager assign workspace — all leads with advanced filters."""
    ensure_roles(cu, ["admin", "manager"])
    inquiry_key = (inquiry_status or "all").strip().lower()
    if inquiry_key not in ASSIGN_INQUIRY_STATUSES:
        raise HTTPException(400, detail=f"inquiry_status must be one of: {', '.join(ASSIGN_INQUIRY_STATUSES)}")
    source_key = (source or "all").strip().lower()
    if source_key not in ASSIGN_SOURCE_FILTERS:
        raise HTTPException(400, detail=f"source must be one of: {', '.join(ASSIGN_SOURCE_FILTERS)}")

    limit = min(max(limit, 1), 500)
    offset = max(offset, 0)
    select = (
        "lead_id,name,phone,email,source,stage,status,priority,call_status,"
        "budget,location,property_type,assigned_to,assigned_at,assigned_by,created_at,updated_at"
    )
    all_leads = fetch_all_leads_merged(select)
    employees = sb_select("employees", {
        "select": "employee_id,name,email,role,department,active",
        "active": "eq.true",
        "order": "name.asc",
    })
    assignable = [e for e in employees if e.get("role") in ASSIGNABLE_EMPLOYEE_ROLES or e.get("role") in {"admin", "manager"}]
    assignable = assignable or employees

    filtered = filter_assign_workspace_leads(all_leads, inquiry_key, source_key, assigned_to, q, location)
    page = enrich_leads_with_employee_names(filtered[offset:offset + limit], assignable)
    facets = compute_assign_workspace_facets(all_leads)

    return {
        "total": len(filtered),
        "leads": page,
        "employees": assignable,
        "facets": facets,
        "filters": {
            "inquiry_status": inquiry_key,
            "source": source_key,
            "assigned_to": assigned_to or "all",
            "q": q or "",
            "location": location or "",
        },
        "auto_assign_on_intake": False,
    }


@api_router.post("/leads/bulk-manage")
async def bulk_manage_leads(body: BulkLeadManageRequest, cu: User = Depends(get_current_user)):
    """Manager bulk assign + change inquiry status (reassign not-interested leads, etc.)."""
    ensure_roles(cu, ["admin", "manager"])
    lead_ids = [lid.strip() for lid in (body.lead_ids or []) if lid and lid.strip()]
    if len(lead_ids) < 1:
        raise HTTPException(status_code=400, detail="Select at least one lead.")

    preset = INQUIRY_ACTION_PRESETS.get((body.inquiry_action or "").strip().lower(), {})
    updated: List[str] = []
    skipped: List[Dict[str, str]] = []

    for lead_id in lead_ids:
        leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*", "limit": "1"})
        old_lead = leads[0] if leads else next((l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id), None)
        if not old_lead:
            skipped.append({"lead_id": lead_id, "reason": "not_found"})
            continue

        if body.assigned_to:
            assign_lead_to_employee(lead_id, old_lead, body.assigned_to.strip(), cu, reactivate=body.reactivate)
            old_lead = {**old_lead, "assigned_to": body.assigned_to.strip(), "status": "active"}

        patch: Dict[str, Any] = {"updated_at": now_utc().isoformat()}
        if preset or body.inquiry_action:
            patch.update(inquiry_preset_to_patch(preset, body.inquiry_action))
        for key in ("stage", "status", "priority", "call_status"):
            val = getattr(body, key, None)
            if val is not None:
                patch[key] = val
        if body.reactivate and old_lead.get("status") == "negative" and "status" not in patch:
            patch.update(fresh_lead_reset_fields("assigned"))
        if patch.keys() - {"updated_at"}:
            sb_update("leads", "lead_id", lead_id, patch)
            update_cached_lead(lead_id, patch)
            log_activity(
                cu,
                "manager_bulk_update",
                f"Manager updated lead ({body.inquiry_action or 'custom'})"
                + (f" → assigned {body.assigned_to}" if body.assigned_to else ""),
                lead_id=lead_id,
            )
        updated.append(lead_id)

    if body.assigned_to:
        sb_update("employees", "employee_id", body.assigned_to.strip(), {"last_assigned_at": now_utc().isoformat()})

    return {"status": "success", "updated_count": len(updated), "updated": updated, "skipped": skipped}


@api_router.post("/leads/bulk-assign")
async def bulk_assign_leads(body: BulkAssignRequest, cu: User = Depends(get_current_user)):
    """Manager: assign multiple selected leads to one employee."""
    ensure_roles(cu, ["admin", "manager"])
    lead_ids = [lid.strip() for lid in (body.lead_ids or []) if lid and lid.strip()]
    employee_id = (body.assigned_to or "").strip()
    if len(lead_ids) < 1:
        raise HTTPException(status_code=400, detail="Select at least one lead.")
    if not employee_id:
        raise HTTPException(status_code=400, detail="Select an employee to assign to.")

    emps = sb_select("employees", {"employee_id": f"eq.{employee_id}", "select": "employee_id,name,active", "limit": "1"})
    if not emps or not emps[0].get("active", True):
        raise HTTPException(status_code=400, detail="Employee not found or inactive.")

    assigned: List[str] = []
    skipped: List[Dict[str, str]] = []
    for lead_id in lead_ids:
        leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*", "limit": "1"})
        old_lead = leads[0] if leads else next((l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id), None)
        if not old_lead:
            skipped.append({"lead_id": lead_id, "reason": "not_found"})
            continue
        if old_lead.get("assigned_to") == employee_id:
            skipped.append({"lead_id": lead_id, "reason": "already_assigned"})
            continue
        assign_lead_to_employee(lead_id, old_lead, employee_id, cu)
        assigned.append(lead_id)

    sb_update("employees", "employee_id", employee_id, {"last_assigned_at": now_utc().isoformat()})
    return {
        "status": "success",
        "assigned": assigned,
        "assigned_count": len(assigned),
        "skipped": skipped,
        "employee_id": employee_id,
        "employee_name": emps[0].get("name"),
    }


@api_router.post("/leads/assign-queue/auto")
async def auto_assign_queue(cu: User = Depends(get_current_user)):
    """Round-robin assign all unassigned leads (e.g. after import or reset)."""
    ensure_roles(cu, ["admin", "manager"])
    all_leads = fetch_all_leads_merged("lead_id,name,stage,status,assigned_to")
    queue = compute_unassigned_queue(all_leads)
    assigned: List[str] = []
    for lead in queue:
        eid = assign_lead_round_robin()
        if not eid:
            break
        assign_lead_to_employee(lead["lead_id"], lead, eid, cu)
        assigned.append(lead["lead_id"])
    return {"status": "success", "assigned_count": len(assigned), "assigned": assigned}


@api_router.get("/stats/assignment")
async def stats_assignment(cu: User = Depends(get_current_user)):
    """Per-employee assigned / active / completed lead counts for manager dashboards."""
    ensure_roles(cu, ["admin", "manager"])
    now = time.time()
    if _assignment_stats_cache["data"] is not None and (now - float(_assignment_stats_cache["ts"])) < _STATS_CACHE_TTL_SEC:
        return _assignment_stats_cache["data"]
    employees = sb_select("employees", {"select": "employee_id,name,email,role,department,active,user_id", "order": "name.asc"})
    all_leads = fetch_all_leads_merged(
        "lead_id,assigned_to,status,stage,priority,call_status,created_at,updated_at,follow_up_at,assigned_at,last_employee_action_at"
    )
    rows = []
    for e in employees:
        eid = e.get("employee_id")
        emp_leads = leads_for_employee_record(e, all_leads)
        rows.append({
            "employee_id": eid,
            "name": e.get("name"),
            "email": e.get("email"),
            "role": e.get("role"),
            "department": e.get("department"),
            "active": e.get("active", True),
            **compute_employee_assignment_stats(emp_leads, e.get("role")),
        })
    unassigned = compute_unassigned_queue(all_leads)
    payload = {
        "unassigned_count": len(unassigned),
        "employees": rows,
    }
    _assignment_stats_cache.update({"ts": now, "data": payload})
    return payload


@api_router.get("/leads/by-platform/{platform}")
async def list_leads_by_platform(
    platform: str,
    limit: int = 200,
    offset: int = 0,
    status_filter: Optional[str] = None,
    cu: User = Depends(get_current_user),
):
    """List real CRM leads for a platform bucket: manual, housing, or meta."""
    platform_key = platform.strip().lower()
    if platform_key not in {"manual", "housing", "meta", "other"}:
        raise HTTPException(status_code=400, detail="Platform must be manual, housing, meta, or other")

    limit = min(max(limit, 1), 500)
    offset = max(offset, 0)
    all_leads = fetch_all_leads_merged()
    cleaned = clean_leads_for_platform_stats(all_leads)
    filtered = [l for l in cleaned if lead_matches_platform(l, platform_key)]
    sf = (status_filter or "").strip().lower()
    if sf == "positive":
        filtered = [l for l in filtered if l.get("stage") in ["positive", "site_visit", "booking", "loan", "registration", "closed"] and l.get("status") != "negative"]
    elif sf in ("not_interested", "negative"):
        filtered = [l for l in filtered if l.get("status") == "negative"]
    elif sf == "registration":
        filtered = [l for l in filtered if l.get("stage") == "registration"]
    elif sf == "booking":
        filtered = [l for l in filtered if l.get("stage") in ["booking", "loan"]]
    filtered.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    page = filtered[offset:offset + limit]
    return {
        "platform": platform_key,
        "label": PLATFORM_LABELS.get(platform_key, "Other Sources"),
        "total": len(filtered),
        "leads": page,
    }

@api_router.get("/leads/recent")
async def list_recent_leads(limit: int = 20, cu: User = Depends(get_current_user)):
    """Most recent real leads with their platform + arrival time. Powers the
    manager/admin 'new lead arrived' popup. Meta test/fake leads are excluded."""
    limit = min(max(limit, 1), 100)
    platform_labels = PLATFORM_LABELS.copy()
    platform_labels["other"] = PLATFORM_LABELS["other"]
    db_leads = sb_select("leads", {
        "select": "lead_id,name,phone,email,source,stage,status,assigned_to,external_lead_id,created_at",
        "order": "created_at.desc",
        "limit": str(limit * 3),
    })
    leads = merge_leads_with_cache(db_leads)
    leads.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    out = []
    for l in leads:
        if clean_text(l.get("assigned_to")):
            continue
        platform = classify_lead_platform(l.get("source"))
        if platform == "meta" and not is_real_meta_lead(l):
            continue
        out.append({
            "lead_id": l.get("lead_id"),
            "name": l.get("name"),
            "phone": l.get("phone"),
            "source": l.get("source"),
            "platform": platform,
            "platform_label": platform_labels.get(platform, "Other"),
            "created_at": l.get("created_at"),
        })
        if len(out) >= limit:
            break
    return {"total": len(out), "leads": out}

@api_router.get("/leads/{lead_id}")
async def get_lead(lead_id: str, cu: User=Depends(get_current_user)):
    lead_fields = (
        "lead_id,name,phone,email,source,stage,status,lead_type,priority,call_status,"
        "budget,location,property_type,notes,assigned_to,assigned_at,last_employee_action_at,follow_up_at,"
        "created_at,updated_at,external_lead_id,external_created_at,integration_uuid,raw_payload,brokerage_amount"
    )
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": lead_fields})
    lead = None
    if leads:
        lead = leads[0]
    else:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id]
        if cache_match:
            lead = cache_match[0]
    if not lead: raise HTTPException(404, "Lead not found")
    lead = enrich_lead_display_fields(lead)
    timeline = sb_select("activities", {
        "lead_id": f"eq.{lead_id}",
        "select": "activity_id,type,text,created_at,user_id,lead_id",
        "order": "created_at.desc",
        "limit": "40",
    })
    cache_activities = [a for a in SESSION_CACHE["activities"] if a.get("lead_id") == lead_id]
    
    # Deduplicate activities
    cache_act_ids = {a.get("activity_id") for a in cache_activities}
    all_timeline = cache_activities + [a for a in timeline if a.get("activity_id") not in cache_act_ids]
    
    return {"lead": lead, "timeline": all_timeline}

@api_router.delete("/leads/{lead_id}")
async def delete_lead(lead_id: str, cu: User=Depends(get_current_user)):
    if cu.role != "admin":
        raise HTTPException(status_code=403, detail="Only admins can delete leads")
    
    # 1. Delete associated data first
    sb_delete("visit_followups", "lead_id", lead_id)
    sb_delete("visits", "lead_id", lead_id)
    sb_delete("bookings", "lead_id", lead_id)
    sb_delete("loans", "lead_id", lead_id)
    sb_delete("activities", "lead_id", lead_id)
    sb_delete("lead_notes", "lead_id", lead_id)
    sb_delete("customers", "lead_id", lead_id)
    sb_delete("notifications", "lead_id", lead_id)
    
    # 2. Delete the lead
    res = sb_delete("leads", "lead_id", lead_id)
    
    # 3. Clean from session cache
    global SESSION_CACHE
    if "leads" in SESSION_CACHE:
        SESSION_CACHE["leads"] = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") != lead_id]
    if "visits" in SESSION_CACHE:
        SESSION_CACHE["visits"] = [v for v in SESSION_CACHE["visits"] if v.get("lead_id") != lead_id]
    if "followups" in SESSION_CACHE:
        SESSION_CACHE["followups"] = [f for f in SESSION_CACHE["followups"] if f.get("lead_id") != lead_id]
    if "bookings" in SESSION_CACHE:
        SESSION_CACHE["bookings"] = [b for b in SESSION_CACHE["bookings"] if b.get("lead_id") != lead_id]
    if "loans" in SESSION_CACHE:
        SESSION_CACHE["loans"] = [ln for ln in SESSION_CACHE["loans"] if ln.get("lead_id") != lead_id]
    if "activities" in SESSION_CACHE:
        SESSION_CACHE["activities"] = [a for a in SESSION_CACHE["activities"] if a.get("lead_id") != lead_id]
    if "customers" in SESSION_CACHE:
        SESSION_CACHE["customers"] = [c for c in SESSION_CACHE["customers"] if c.get("lead_id") != lead_id]
    if "notifications" in SESSION_CACHE:
        SESSION_CACHE["notifications"] = [n for n in SESSION_CACHE["notifications"] if n.get("lead_id") != lead_id]
    
    return {"status": "deleted", "lead_id": lead_id}

@api_router.get("/leads/{lead_id}/ai-summary")
async def get_lead_ai_summary(lead_id: str, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*"})
    if not leads: raise HTTPException(404, "Lead not found")
    
    timeline = sb_select("activities", {"lead_id": f"eq.{lead_id}", "select": "*", "order": "created_at.desc", "limit": "20"})
    
    # Try Real AI first
    summary = AIService.generate_lead_summary(timeline)
    
    # If AI fails (e.g. 401 error), use the Smart Fallback
    if not summary or "Could not generate" in summary or "available" in summary:
        l = leads[0]
        summary = f"Customer interested in {l.get('property_type','property')} in {l.get('location','specified area')} with a budget of {l.get('budget','budget')}. {l.get('notes','Follow up for site visit.')}"
        if len(summary) > 100: summary = summary[:97] + "..."

    return {"summary": summary}

@api_router.patch("/leads/{lead_id}")
async def update_lead(lead_id: str, p: LeadUpdate, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*"})
    old_lead = None
    if leads:
        old_lead = leads[0]
    else:
        # Check cache
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id]
        if cache_match:
            old_lead = cache_match[0]
            
    if not old_lead: raise HTTPException(404, "Lead not found")
    ensure_lead_edit_access(cu, old_lead)
    if p.assigned_to is not None:
        ensure_roles(cu, ["admin", "manager"])
    
    data = model_payload(p)
    data["updated_at"] = now_utc().isoformat()
    valid_call_statuses = {"ringing", "out_of_service", "call_back", "disconnect"}
    if data.get("call_status") and data["call_status"] not in valid_call_statuses:
        raise HTTPException(status_code=400, detail="Invalid call status")

    if p.status == "active" and old_lead.get("status") == "negative" and not _can_reactivate_negative_lead(cu):
        raise HTTPException(
            status_code=403,
            detail="Only manager can re-activate not-interested leads. Ask your manager to reset and reassign.",
        )

    if not _can_manage_all_leads(cu) and p.stage in ("new",) and old_lead.get("stage") not in (None, "new", "assigned"):
        raise HTTPException(status_code=403, detail="Only manager can reset lead stage. Ask your manager.")

    if data.get("call_status"):
        workflow = apply_call_status_workflow(old_lead, data["call_status"])
        data.update(workflow)
    
    if p.status == "negative" and old_lead.get("status") != "negative":
        data["call_status"] = None
        data["follow_up_at"] = None

    if p.stage in ("site_visit", "positive") and p.stage and p.stage != old_lead.get("stage"):
        data["call_status"] = None

    if not _can_manage_all_leads(cu):
        workflow_touched = any(
            getattr(p, key) is not None
            for key in ("stage", "status", "priority", "call_status", "follow_up_at")
        )
        if workflow_touched:
            data["last_employee_action_at"] = now_utc().isoformat()
            canonical_emp = resolve_employee_id(cu) or cu.acting_as_employee_id or cu.employee_id
            if canonical_emp and old_lead.get("assigned_to") != canonical_emp:
                if old_lead.get("assigned_to") in employee_assignee_ids(cu):
                    data["assigned_to"] = canonical_emp
    
    if p.assigned_to and p.assigned_to != old_lead.get("assigned_to"):
        stamp = stamp_lead_assignment(
            p.assigned_to,
            assigned_by=cu.acting_as_employee_id or cu.user_id,
            current_stage=old_lead.get("stage"),
        )
        data.update(stamp)
        emps = sb_select("employees", {"employee_id": f"eq.{p.assigned_to}", "select": "name"})
        emp_name = emps[0]["name"] if emps else p.assigned_to
        log_activity(cu, "lead_assigned", f"Assigned lead to {emp_name}", lead_id=lead_id)
        create_notification(p.assigned_to, "Lead assigned", f"{old_lead.get('name', 'Lead')} has been assigned to you.", lead_id=lead_id)
        sb_update("employees", "employee_id", p.assigned_to, {"last_assigned_at": now_utc().isoformat()})
    
    updated = sb_update("leads", "lead_id", lead_id, data)
    
    # Always update cache for immediate UI feedback (with workflow display fields)
    new_lead = enrich_lead_display_fields({**old_lead, **data})
    if leads:
        SESSION_CACHE["leads"] = [l if l.get("lead_id") != lead_id else new_lead for l in SESSION_CACHE["leads"]]
        if not any(l.get("lead_id") == lead_id for l in SESSION_CACHE["leads"]):
            SESSION_CACHE["leads"].insert(0, new_lead)
    else:
        SESSION_CACHE["leads"] = [l if l.get("lead_id") != lead_id else new_lead for l in SESSION_CACHE["leads"]]
    update_cached_lead(lead_id, data)
    updated = new_lead
    
    # Log activity for stage/status changes
    if p.stage and p.stage != old_lead.get("stage"):
        act_type = f"stage_change_{p.stage}"
        if p.stage == "positive": act_type = "positive_response"
        log_activity(cu, act_type, f"Moved lead stage from {old_lead.get('stage')} to {p.stage}", lead_id=lead_id)
    
    if p.status and p.status != old_lead.get("status"):
        act_type = f"status_change_{p.status}"
        if p.status == "negative": act_type = "negative_response"
        log_activity(cu, act_type, f"Changed lead status from {old_lead.get('status')} to {p.status}", lead_id=lead_id)

    if data.get("call_status") and data.get("call_status") != old_lead.get("call_status"):
        log_activity(
            cu,
            "call_status_update",
            f"Call status → {data['call_status'].replace('_', ' ')} (moved to {data['call_status'].replace('_', ' ')} section)",
            lead_id=lead_id,
        )

    # Auto-create related records when lead enters a new department stage
    if p.stage and p.stage != old_lead.get("stage"):
        if p.stage in ["positive", "site_visit"]:
            ensure_visit_record(lead_id, old_lead.get("name", "Lead"), old_lead.get("assigned_to"))
            if p.stage == "positive":
                create_notification(old_lead.get("assigned_to"), "Positive lead", f"{old_lead.get('name', 'Lead')} is ready for site visit follow-up.", lead_id=lead_id)
        elif p.stage == "closed":
            create_customer_from_lead(lead_id, cu)

    return updated

@api_router.post("/leads/{lead_id}/notes")
async def add_lead_note(lead_id: str, p: NoteCreate, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "lead_id"})
    if not leads:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id]
        if not cache_match: raise HTTPException(404, "Lead not found")
    activity = log_activity(cu, p.type, p.text, lead_id=lead_id)
    if activity: SESSION_CACHE["activities"].insert(0, activity)
    return activity

@api_router.post("/leads/{lead_id}/advance")
async def advance_lead(lead_id: str, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*"})
    lead = None
    if leads:
        lead = leads[0]
    else:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id]
        if cache_match: lead = cache_match[0]
    if not lead: raise HTTPException(404, "Lead not found")
    cur = lead.get("stage", "new")
    try: idx = STAGES.index(cur)
    except: idx = 0
    if idx >= len(STAGES) - 1: return lead
    new_stage = STAGES[idx + 1]
    updated = sb_update("leads", "lead_id", lead_id, {"stage": new_stage, "updated_at": now_utc().isoformat()})
    # Update cache
    new_lead = {**lead, "stage": new_stage, "updated_at": now_utc().isoformat()}
    SESSION_CACHE["leads"] = [l if l.get("lead_id") != lead_id else new_lead for l in SESSION_CACHE["leads"]]
    if not any(l.get("lead_id") == lead_id for l in SESSION_CACHE["leads"]):
        SESSION_CACHE["leads"].insert(0, new_lead)

    # Auto-create related records when lead enters a new department stage
    if new_stage in ["positive", "site_visit"]:
        ensure_visit_record(lead_id, lead.get("name", "Lead"), lead.get("assigned_to"))
    elif new_stage == "closed":
        create_customer_from_lead(lead_id, cu)

    log_activity(cu, "stage_change", f"Stage moved {cur} → {new_stage}", lead_id=lead_id)
    return updated or new_lead

# ---- Stage Sync Helper ----
def sync_lead_stage(lead_id: str, target_stage: str, force: bool = False):
    """Ensure lead is at least at target_stage in the pipeline.
    If force=True, also moves a 'closed' lead back to the target_stage."""
    leads_db = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "lead_id,stage"})
    lead = None
    if leads_db:
        lead = leads_db[0]
    else:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id]
        if cache_match: lead = cache_match[0]
    if not lead: return
    cur = lead.get("stage", "new")
    try: cur_idx = [s for s in STAGES].index(cur)
    except ValueError: cur_idx = 0
    try: target_idx = [s for s in STAGES].index(target_stage)
    except ValueError: return
    needs_update = cur_idx < target_idx
    # If lead is closed but we're creating/updating an active department record, bring it back
    if force and cur == "closed" and target_stage != "closed":
        needs_update = True
    if needs_update:
        sb_update("leads", "lead_id", lead_id, {"stage": target_stage, "updated_at": now_utc().isoformat()})
        SESSION_CACHE["leads"] = [({**l, "stage": target_stage, "updated_at": now_utc().isoformat()} if l.get("lead_id") == lead_id else l) for l in SESSION_CACHE["leads"]]
        if target_stage == "closed":
            create_customer_from_lead(lead_id)

# ---- Visits ----
@api_router.post("/visits")
async def create_visit(p: SiteVisitCreate, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{p.lead_id}", "select": "lead_id,name"})
    lead_name = "Lead"
    if leads:
        lead_name = leads[0]["name"]
    else:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == p.lead_id]
        if cache_match:
            lead_name = cache_match[0].get("name", "Lead")
        else:
            raise HTTPException(404, "Lead not found")
    vid = gen_id("vis")
    v = {
        "visit_id": vid, "lead_id": p.lead_id, "lead_name": lead_name, "scheduled_at": p.scheduled_at.isoformat(),
        "assigned_to": p.assigned_to, "status": "scheduled", "feedback": None,
        "interested": None, "created_at": now_utc().isoformat(),
    }
    result = sb_insert("visits", v)
    SESSION_CACHE["visits"].insert(0, result or v)
    # Auto-sync lead stage to site_visit
    sync_lead_stage(p.lead_id, "site_visit", force=True)
    return result or v

@api_router.get("/visits")
async def list_visits(cu: User=Depends(get_current_user)):
    visits = sb_select("visits", {"select": "*", "order": "scheduled_at.desc"})
    followups = sb_select("visit_followups", {"select": "*", "order": "follow_up_at.desc"})
    # Deduplicate visits (cache wins)
    cache_ids = {v.get("visit_id") for v in SESSION_CACHE["visits"]}
    db_only = [v for v in visits if v.get("visit_id") not in cache_ids]
    all_visits = SESSION_CACHE["visits"] + db_only

    cache_followup_ids = {f.get("followup_id") for f in SESSION_CACHE["followups"]}
    all_followups = SESSION_CACHE["followups"] + [f for f in followups if f.get("followup_id") not in cache_followup_ids]

    # Fetch all leads & employees to enrich dynamically
    leads = sb_select("leads", {"select": "lead_id,name"})
    employees = sb_select("employees", {"select": "employee_id,name"})

    cache_leads = {l.get("lead_id"): l.get("name") for l in SESSION_CACHE["leads"]}
    db_leads = {l.get("lead_id"): l.get("name") for l in leads}
    lead_name_map = {**db_leads, **cache_leads}

    emp_map = {e.get("employee_id"): e.get("name") for e in employees}

    enriched_visits = []
    for v in all_visits:
        v_copy = dict(v)
        v_copy["lead_name"] = lead_name_map.get(v.get("lead_id"), v.get("lead_name", "Lead"))
        if v.get("assigned_to"):
            v_copy["assigned_name"] = emp_map.get(v.get("assigned_to"))
        visit_followups = [f for f in all_followups if f.get("visit_id") == v.get("visit_id")]
        v_copy["followups_count"] = max(len(visit_followups), 1 if v.get("status") == "follow_up" else 0)
        v_copy["next_follow_up_at"] = visit_followups[0].get("follow_up_at") if visit_followups else None
        v_copy["next_follow_up_date"] = visit_followups[0].get("follow_up_date") if visit_followups else None
        v_copy["next_follow_up_time"] = visit_followups[0].get("follow_up_time") if visit_followups else None
        v_copy["next_follow_up_day"] = visit_followups[0].get("follow_up_day") if visit_followups else None
        enriched_visits.append(v_copy)

    return enriched_visits

@api_router.patch("/visits/{visit_id}")
async def update_visit(visit_id: str, p: SiteVisitUpdate, cu: User=Depends(get_current_user)):
    data = model_payload(p)
    updated = sb_update("visits", "visit_id", visit_id, data)
    # Update cache
    visit_record = None
    for i, v in enumerate(SESSION_CACHE["visits"]):
        if v.get("visit_id") == visit_id:
            SESSION_CACHE["visits"][i] = {**v, **data}
            visit_record = SESSION_CACHE["visits"][i]
            if not updated: updated = SESSION_CACHE["visits"][i]
            break
    if not updated:
        db_visits = sb_select("visits", {"visit_id": f"eq.{visit_id}"})
        if db_visits:
            visit_record = db_visits[0]
            updated = visit_record
        else:
            raise HTTPException(404, "Visit not found")
    # Auto-sync lead stage
    if visit_record and visit_record.get("lead_id"):
        lead_id = visit_record["lead_id"]
        if p.interested is True:
            sync_lead_stage(lead_id, "booking", force=False)
            log_activity(cu, "site_visit_interested", "Site visit marked interested; moved to booking department.", lead_id=lead_id)
        else:
            sync_lead_stage(lead_id, "site_visit", force=True)
    return updated

@api_router.post("/leads/{lead_id}/follow-up")
async def create_lead_follow_up(lead_id: str, p: LeadFollowUpCreate, cu: User = Depends(get_current_user)):
    """Schedule a follow-up straight from a lead (cold / visited) — no site visit needed.
    Stores a visit_followups row, sets lead.follow_up_at, and logs the activity so it
    shows in the Follow Ups tab and dashboard counts."""
    rows = sb_select("leads", {"lead_id": f"eq.{lead_id}", "select": "*", "limit": "1"})
    lead = rows[0] if rows else next((l for l in SESSION_CACHE["leads"] if l.get("lead_id") == lead_id), None)
    if not lead:
        raise HTTPException(404, "Lead not found")

    if p.follow_up_date and p.follow_up_time:
        follow_up_at = parse_follow_up_at(p.follow_up_date, p.follow_up_time)
    else:
        raise HTTPException(400, "Follow-up date and time are required")

    parts = follow_up_display_parts(follow_up_at.isoformat())
    note_parts: List[str] = []
    if clean_text(p.reason):
        note_parts.append(f"Reason: {clean_text(p.reason)}")
    if clean_text(p.notes):
        note_parts.append(clean_text(p.notes))
    combined_notes = "\n".join(note_parts) if note_parts else None

    followup = {
        "followup_id": gen_id("fup"),
        "visit_id": gen_id("fuv"),  # synthetic id keeps NOT NULL happy on older DBs
        "lead_id": lead_id,
        "lead_name": lead.get("name", "Lead"),
        "follow_up_date": p.follow_up_date or parts["follow_up_date"],
        "follow_up_time": p.follow_up_time or parts["follow_up_time"],
        "follow_up_day": p.follow_up_day or parts["follow_up_day"],
        "follow_up_at": follow_up_at.isoformat(),
        "status": "scheduled",
        "notes": combined_notes,
        "created_by": cu.acting_as_employee_id or cu.user_id,
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("visit_followups", followup)
    followup_record = result or followup
    SESSION_CACHE["followups"].insert(0, followup_record)

    lead_update = {
        "follow_up_at": follow_up_at.isoformat(),
        "updated_at": now_utc().isoformat(),
        "last_employee_action_at": now_utc().isoformat(),
    }
    sb_update("leads", "lead_id", lead_id, lead_update)
    update_cached_lead(lead_id, lead_update)
    detail = f"Follow-up scheduled for {parts['follow_up_day']} {parts['follow_up_date']} at {parts['follow_up_time']}."
    if clean_text(p.reason):
        detail += f" Reason: {clean_text(p.reason)}."
    activity = log_activity(cu, "lead_followup", detail, lead_id=lead_id)
    SESSION_CACHE["activities"].insert(0, activity)
    return followup_record


@api_router.post("/visit-followups")
async def create_visit_followup(p: SiteVisitFollowUpCreate, cu: User=Depends(get_current_user)):
    if not p.follow_up_day.strip():
        raise HTTPException(400, "Follow-up day is required")

    follow_up_at = parse_follow_up_at(p.follow_up_date, p.follow_up_time)
    visit = get_visit_record(p.visit_id)
    if not visit:
        raise HTTPException(404, "Visit not found")

    followup = {
        "followup_id": gen_id("fup"),
        "visit_id": p.visit_id,
        "lead_id": visit.get("lead_id"),
        "lead_name": visit.get("lead_name", "Lead"),
        "follow_up_date": p.follow_up_date.strip(),
        "follow_up_time": p.follow_up_time.strip(),
        "follow_up_day": p.follow_up_day.strip(),
        "follow_up_at": follow_up_at.isoformat(),
        "status": "scheduled",
        "notes": p.notes,
        "created_by": cu.acting_as_employee_id or cu.user_id,
        "created_at": now_utc().isoformat(),
    }
    result = sb_insert("visit_followups", followup)
    followup_record = result or followup
    SESSION_CACHE["followups"].insert(0, followup_record)

    visit_update = {
        "status": "follow_up",
        "scheduled_at": follow_up_at.isoformat(),
    }
    sb_update("visits", "visit_id", p.visit_id, visit_update)
    for i, cached_visit in enumerate(SESSION_CACHE["visits"]):
        if cached_visit.get("visit_id") == p.visit_id:
            SESSION_CACHE["visits"][i] = {**cached_visit, **visit_update}
            break

    if followup.get("lead_id"):
        lead_update = {
            "follow_up_at": follow_up_at.isoformat(),
            "updated_at": now_utc().isoformat(),
            "last_employee_action_at": now_utc().isoformat(),
        }
        sb_update("leads", "lead_id", followup["lead_id"], lead_update)
        update_cached_lead(followup["lead_id"], lead_update)
        activity = log_activity(
            cu,
            "site_visit_followup",
            f"Follow-up scheduled for {p.follow_up_day.strip()} {p.follow_up_date.strip()} at {p.follow_up_time.strip()}.",
            lead_id=followup["lead_id"],
        )
        SESSION_CACHE["activities"].insert(0, activity)

    return followup_record

@api_router.get("/visit-followups")
async def list_visit_followups(visit_id: Optional[str]=None, lead_id: Optional[str]=None, cu: User=Depends(get_current_user)):
    rows = sb_select_all("visit_followups", {"select": "*", "order": "follow_up_at.desc"})
    cache_ids = {f.get("followup_id") for f in SESSION_CACHE["followups"]}
    followups = SESSION_CACHE["followups"] + [f for f in rows if f.get("followup_id") not in cache_ids]

    visits = sb_select("visits", {"select": "*", "status": "eq.follow_up", "order": "scheduled_at.desc"})
    cache_visit_ids = {v.get("visit_id") for v in SESSION_CACHE["visits"]}
    follow_up_visits = [v for v in SESSION_CACHE["visits"] if v.get("status") == "follow_up"]
    follow_up_visits += [v for v in visits if v.get("visit_id") not in cache_visit_ids]

    lead_rows = fetch_all_leads_merged("lead_id,name,assigned_to,follow_up_at")
    lead_names = {l.get("lead_id"): l.get("name") for l in lead_rows}

    followup_visit_ids = {f.get("visit_id") for f in followups}
    for visit in follow_up_visits:
        if visit.get("visit_id") in followup_visit_ids:
            continue
        parts = follow_up_display_parts(visit.get("scheduled_at"))
        followups.append({
            "followup_id": f"visit_{visit.get('visit_id')}",
            "visit_id": visit.get("visit_id"),
            "lead_id": visit.get("lead_id"),
            "lead_name": lead_names.get(visit.get("lead_id"), visit.get("lead_name", "Lead")),
            "status": "scheduled",
            "notes": visit.get("feedback"),
            "created_at": visit.get("updated_at") or visit.get("created_at"),
            **parts,
        })

    if visit_id:
        followups = [f for f in followups if f.get("visit_id") == visit_id]
    if lead_id:
        followups = [f for f in followups if f.get("lead_id") == lead_id]

    employees = sb_select("employees", {"select": "employee_id,name"})
    emp_names = {e.get("employee_id"): e.get("name") for e in employees if e.get("employee_id")}
    lead_assign = {r.get("lead_id"): r.get("assigned_to") for r in lead_rows}
    for f in followups:
        lid = f.get("lead_id")
        eid = lead_assign.get(lid)
        if eid and emp_names.get(eid):
            f["employee_name"] = emp_names[eid]

    return sorted(followups, key=lambda f: f.get("follow_up_at") or "", reverse=True)

# ---- Bookings ----
@api_router.post("/bookings")
async def create_booking(p: BookingCreate, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{p.lead_id}", "select": "lead_id,name"})
    lead_name = "Lead"
    if leads:
        lead_name = leads[0]["name"]
    else:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == p.lead_id]
        if cache_match:
            lead_name = cache_match[0].get("name", "Lead")
        else:
            raise HTTPException(404, "Lead not found")
    bid = gen_id("bkg")
    amount = float(p.booking_amount or 0)
    token = float(p.token_received or 0)
    b = {
        "booking_id": bid, "lead_id": p.lead_id, "lead_name": lead_name,
        "property_name": p.property_name or "Property TBD", "booking_amount": amount,
        "token_received": token, "agreement_status": "pending",
        "payment_progress": p.payment_progress if p.payment_progress is not None else (int((token / amount) * 100) if amount else 0),
        "status": "active", "created_at": now_utc().isoformat(),
    }
    optional_costs = {
        "unit_number": p.unit_number, "tower": p.tower,
        "flat_cost": p.flat_cost, "agreement_value": p.agreement_value,
        "stamp_duty": p.stamp_duty, "registration_fees": p.registration_fees,
        "gst": p.gst, "society_charges": p.society_charges,
        "payment_status": p.payment_status, "booking_date": p.booking_date.isoformat() if p.booking_date else None,
        "starred": p.starred, "completed_tasks": p.completed_tasks,
    }
    for k, v in optional_costs.items():
        if v is not None:
            b[k] = v
    result = sb_insert("bookings", b)
    SESSION_CACHE["bookings"].insert(0, result or b)
    sb_update("leads", "lead_id", p.lead_id, {
        "stage": "booking",
        "priority": None,
        "updated_at": now_utc().isoformat(),
    })
    update_cached_lead(p.lead_id, {"stage": "booking", "priority": None, "updated_at": now_utc().isoformat()})
    return result or b

@api_router.get("/bookings")
async def list_bookings(cu: User=Depends(get_current_user)):
    bookings = sb_select("bookings", {"select": "*", "order": "created_at.desc"})
    cache_ids = {b.get("booking_id") for b in SESSION_CACHE["bookings"]}
    db_only = [b for b in bookings if b.get("booking_id") not in cache_ids]
    merged = SESSION_CACHE["bookings"] + db_only
    # Hide legacy auto-created skeleton rows; only explicit New Booking entries show.
    return [b for b in merged if not is_legacy_skeleton_booking(b)]

@api_router.patch("/bookings/{booking_id}")
async def update_booking(booking_id: str, p: BookingUpdate, cu: User=Depends(get_current_user)):
    booking_before = None
    db_bookings_before = sb_select("bookings", {"booking_id": f"eq.{booking_id}", "select": "*", "limit": "1"})
    if db_bookings_before:
        booking_before = db_bookings_before[0]
    else:
        cache_before = [b for b in SESSION_CACHE["bookings"] if b.get("booking_id") == booking_id]
        if cache_before:
            booking_before = cache_before[0]

    data = model_payload(p)
    if "token_received" in data or "booking_amount" in data:
        amount = float(data.get("booking_amount", (booking_before or {}).get("booking_amount", 0)) or 0)
        token = float(data.get("token_received", (booking_before or {}).get("token_received", 0)) or 0)
        if amount:
            data["payment_progress"] = min(100, int((token / amount) * 100))
        elif "payment_progress" not in data:
            data["payment_progress"] = 0
    updated = sb_update("bookings", "booking_id", booking_id, data)
    if not updated and any(key in data for key in ["completed_tasks", "starred"]):
        compatible_data = {k: v for k, v in data.items() if k not in ["completed_tasks", "starred"]}
        if compatible_data:
            updated = sb_update("bookings", "booking_id", booking_id, compatible_data)
    # Update cache
    booking_record = None
    for i, b in enumerate(SESSION_CACHE["bookings"]):
        if b.get("booking_id") == booking_id:
            SESSION_CACHE["bookings"][i] = {**b, **data}
            booking_record = SESSION_CACHE["bookings"][i]
            if not updated: updated = SESSION_CACHE["bookings"][i]
            break
    if not booking_record and (booking_before or updated):
        booking_record = {**(booking_before or {}), **(updated or {}), **data}
        updated = booking_record
        SESSION_CACHE["bookings"] = [b for b in SESSION_CACHE["bookings"] if b.get("booking_id") != booking_id]
        SESSION_CACHE["bookings"].insert(0, booking_record)
    if not updated:
        db_bookings = sb_select("bookings", {"booking_id": f"eq.{booking_id}"})
        if db_bookings:
            booking_record = db_bookings[0]
            updated = booking_record
        else:
            raise HTTPException(404, "Booking not found")
    # Auto-sync lead stage
    if booking_record and booking_record.get("lead_id"):
        status = str(data.get("status") or booking_record.get("status") or "").lower()
        lead_id = booking_record["lead_id"]
        completed_tasks = data.get("completed_tasks") or booking_record.get("completed_tasks") or []
        loan_ready_tasks = {"login_file", "sanctioned", "registration", "disbursement", "bill_submitted"}
        if status in ["registration"]:
            sync_lead_stage(lead_id, "registration", force=False)
            ensure_loan_record(lead_id, booking_record.get("lead_name"))
        elif status in ["confirmed", "token received", "agreement pending", "login file", "sanctioned", "disbursement", "bill submitted"] or any(t in loan_ready_tasks for t in completed_tasks):
            sync_lead_stage(lead_id, "loan", force=False)
            ensure_loan_record(lead_id, booking_record.get("lead_name"))
        elif status in ["cancellation", "cancelled"]:
            lead_patch = {
                "status": "negative",
                "priority": None,
                "follow_up_at": None,
                "updated_at": now_utc().isoformat(),
            }
            sb_update("leads", "lead_id", lead_id, lead_patch)
            update_cached_lead(lead_id, lead_patch)
            log_activity(
                cu, "booking_cancelled",
                "Booking cancelled; lead moved to Not Interested.",
                lead_id=lead_id,
            )
        else:
            sync_lead_stage(lead_id, "booking", force=True)
    return updated

@api_router.delete("/bookings/{booking_id}")
async def delete_booking(booking_id: str, cu: User=Depends(get_current_user)):
    booking = None
    rows = sb_select("bookings", {"booking_id": f"eq.{booking_id}", "select": "*", "limit": "1"})
    if rows:
        booking = rows[0]
    else:
        cache_match = [b for b in SESSION_CACHE["bookings"] if b.get("booking_id") == booking_id]
        if cache_match:
            booking = cache_match[0]

    sb_delete("bookings", "booking_id", booking_id)
    SESSION_CACHE["bookings"] = [b for b in SESSION_CACHE["bookings"] if b.get("booking_id") != booking_id]
    if booking and booking.get("lead_id"):
        log_activity(cu, "booking_deleted", "Booking record deleted.", lead_id=booking.get("lead_id"))
    return {"ok": True, "booking_id": booking_id}

# ---- Loans ----
@api_router.post("/loans")
async def create_loan(p: LoanCreate, cu: User=Depends(get_current_user)):
    leads = sb_select("leads", {"lead_id": f"eq.{p.lead_id}", "select": "lead_id,name"})
    lead_name = "Lead"
    if leads:
        lead_name = leads[0]["name"]
    else:
        cache_match = [l for l in SESSION_CACHE["leads"] if l.get("lead_id") == p.lead_id]
        if cache_match:
            lead_name = cache_match[0].get("name", "Lead")
        else:
            raise HTTPException(404, "Lead not found")
    lid = gen_id("lon")
    l = {
        "loan_id": lid, "lead_id": p.lead_id, "lead_name": lead_name,
        "bank_name": p.bank_name, "amount": p.amount, "application_status": "pending",
        "bank_stage": "documentation", "emi_eligible": None, "progress": 0,
        "created_at": now_utc().isoformat(),
    }
    if p.documents_status: l["documents_status"] = p.documents_status
    if p.pending_documents is not None: l["pending_documents"] = p.pending_documents
    if p.starred is not None: l["starred"] = p.starred
    result = sb_insert("loans", l)
    SESSION_CACHE["loans"].insert(0, result or l)
    sb_update("leads", "lead_id", p.lead_id, {
        "stage": "loan",
        "priority": None,
        "updated_at": now_utc().isoformat(),
    })
    update_cached_lead(p.lead_id, {"stage": "loan", "priority": None, "updated_at": now_utc().isoformat()})
    return result or l

@api_router.get("/loans")
async def list_loans(cu: User=Depends(get_current_user)):
    loans = sb_select("loans", {"select": "*", "order": "created_at.desc"})
    # Deduplicate loans (cache wins)
    cache_ids = {ln.get("loan_id") for ln in SESSION_CACHE["loans"]}
    db_only = [ln for ln in loans if ln.get("loan_id") not in cache_ids]
    return SESSION_CACHE["loans"] + db_only

@api_router.patch("/loans/{loan_id}")
async def update_loan(loan_id: str, p: LoanUpdate, cu: User=Depends(get_current_user)):
    data = model_payload(p)
    updated = sb_update("loans", "loan_id", loan_id, data)
    # Update cache
    loan_record = None
    for i, ln in enumerate(SESSION_CACHE["loans"]):
        if ln.get("loan_id") == loan_id:
            SESSION_CACHE["loans"][i] = {**ln, **data}
            loan_record = SESSION_CACHE["loans"][i]
            if not updated: updated = SESSION_CACHE["loans"][i]
            break
    if not loan_record and updated:
        loan_record = updated
    if not updated:
        # Try to find it from DB
        db_loans = sb_select("loans", {"loan_id": f"eq.{loan_id}"})
        if db_loans:
            loan_record = db_loans[0]
            updated = loan_record
        else:
            raise HTTPException(404, "Loan not found")
    # Auto-sync lead stage to loan when loan is actively being worked on
    if loan_record and loan_record.get("lead_id"):
        lead_id = loan_record["lead_id"]
        application_status = str(data.get("application_status") or loan_record.get("application_status") or "").lower()
        bank_stage = str(data.get("bank_stage") or loan_record.get("bank_stage") or "").lower()
        if application_status in ["disbursed", "completed"] or bank_stage == "disbursal":
            sync_lead_stage(lead_id, "closed", force=False)
            create_customer_from_lead(lead_id, cu)
        elif application_status in ["approved"] or bank_stage in ["sanction", "sanctioned"]:
            sync_lead_stage(lead_id, "registration", force=False)
        else:
            sync_lead_stage(lead_id, "loan", force=True)
    return updated

@api_router.delete("/loans/{loan_id}")
async def delete_loan(loan_id: str, cu: User=Depends(get_current_user)):
    loan = None
    rows = sb_select("loans", {"loan_id": f"eq.{loan_id}", "select": "*", "limit": "1"})
    if rows:
        loan = rows[0]
    else:
        cache_match = [ln for ln in SESSION_CACHE["loans"] if ln.get("loan_id") == loan_id]
        if cache_match:
            loan = cache_match[0]

    sb_delete("loans", "loan_id", loan_id)
    SESSION_CACHE["loans"] = [ln for ln in SESSION_CACHE["loans"] if ln.get("loan_id") != loan_id]
    if loan and loan.get("lead_id"):
        log_activity(cu, "loan_deleted", "Loan record deleted.", lead_id=loan.get("lead_id"))
    return {"ok": True, "loan_id": loan_id}

# ---- Customers & Notifications ----
@api_router.get("/customers")
async def list_customers(cu: User=Depends(get_current_user)):
    customers = sb_select("customers", {"select": "*", "order": "converted_at.desc"})
    cache_ids = {c.get("customer_id") for c in SESSION_CACHE["customers"]}
    return SESSION_CACHE["customers"] + [c for c in customers if c.get("customer_id") not in cache_ids]

@api_router.get("/notifications")
async def list_notifications(cu: User=Depends(get_current_user)):
    user_id = cu.acting_as_employee_id or cu.user_id
    params = {"select": "*", "order": "created_at.desc"}
    if cu.role != "admin":
        params["user_id"] = f"eq.{user_id}"
    notifications = sb_select("notifications", params)
    cache_ids = {n.get("notification_id") for n in SESSION_CACHE["notifications"]}
    cached = SESSION_CACHE["notifications"]
    if cu.role != "admin":
        cached = [n for n in cached if n.get("user_id") == user_id]
    return cached + [n for n in notifications if n.get("notification_id") not in cache_ids]

@api_router.patch("/notifications/{notification_id}/read")
async def mark_notification_read(notification_id: str, cu: User=Depends(get_current_user)):
    updated = sb_update("notifications", "notification_id", notification_id, {"is_read": True})
    SESSION_CACHE["notifications"] = [
        ({**n, "is_read": True} if n.get("notification_id") == notification_id else n)
        for n in SESSION_CACHE["notifications"]
    ]
    return updated or {"notification_id": notification_id, "is_read": True}

# ---- Employees ----
def _coerce_allowed_pages(value: Any) -> List[str]:
    """Normalize jsonb / JSON string / list from Supabase into a string list."""
    if value is None:
        return []
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, list):
                return _sanitize_assignable_pages([str(x) for x in parsed if x])
        except (json.JSONDecodeError, TypeError):
            return []
        return []
    if isinstance(value, list):
        return _sanitize_assignable_pages([str(x) for x in value if x])
    return []


def _sanitize_assignable_pages(pages: List[str]) -> List[str]:
    """Per-employee sidebar grants — never include admin-only items."""
    blocked = {"tracking", "my-dashboard", "dashboard"}
    return [p for p in pages if p and p not in blocked]


def _hydrate_allowed_pages_from_employee(u: Dict[str, Any]) -> Dict[str, Any]:
    """Employee row is the source of truth for sidebar access after login."""
    role = (u.get("role") or "").strip()
    if role == "admin":
        return u
    emp_id = u.get("employee_id")
    if not emp_id and u.get("email"):
        emps = sb_select("employees", {
            "email": f"eq.{u['email']}",
            "select": "employee_id,allowed_pages,active,role",
            "limit": "1",
        })
        if emps:
            emp_id = emps[0].get("employee_id")
            u["employee_id"] = emp_id
            if role != "admin" and emp_id:
                u["acting_as_employee_id"] = u.get("acting_as_employee_id") or emp_id
    if not emp_id and u.get("user_id"):
        emps = sb_select("employees", {
            "user_id": f"eq.{u['user_id']}",
            "select": "employee_id,allowed_pages,active,role",
            "limit": "1",
        })
        if emps:
            emp_id = emps[0].get("employee_id")
            u["employee_id"] = emp_id
            if role != "admin" and emp_id:
                u["acting_as_employee_id"] = u.get("acting_as_employee_id") or emp_id
    if not emp_id:
        return u
    emps = sb_select("employees", {
        "employee_id": f"eq.{emp_id}",
        "select": "allowed_pages,active,role",
        "limit": "1",
    })
    if not emps:
        return u
    u["allowed_pages"] = _employee_allowed_pages(emps[0])
    if emps[0].get("role") and not u.get("role"):
        u["role"] = emps[0]["role"]
    return u


def _employee_allowed_pages(employee: Dict[str, Any]) -> List[str]:
    """Pages stored on the employee row; fall back to role defaults only when unset."""
    if employee.get("allowed_pages") is None:
        return _default_pages_for_role(employee.get("role") or "telecaller")
    return _coerce_allowed_pages(employee.get("allowed_pages"))


def _default_pages_for_role(role: str) -> List[str]:
    """Sensible fallback service access when the manager doesn't tick any boxes."""
    defaults = {
        "admin": ["dashboard", "my-dashboard", "pipeline", "assign-leads", "telecaller", "sales-executive", "bookings", "loans", "integrations", "broker", "tracking", "employees", "negative"],
        "manager": ["dashboard", "my-dashboard", "pipeline", "assign-leads", "bookings", "loans", "integrations", "broker", "employees"],
        "telecaller": ["my-dashboard", "telecaller", "pipeline", "negative"],
        "site_visit": ["my-dashboard", "sales-executive", "pipeline"],
        "sales_executive": ["my-dashboard", "sales-executive", "telecaller", "pipeline"],
        "booking": ["my-dashboard", "bookings", "pipeline"],
        "loan": ["my-dashboard", "loans", "pipeline"],
        "marketing": ["my-dashboard", "negative", "pipeline", "integrations"],
    }
    return defaults.get(role, ["my-dashboard", "pipeline"])


def _build_login_row(name: str, email: str, password: str, role: str, employee_id: str, allowed_pages: List[str]) -> Dict[str, Any]:
    plain = password.strip()
    row = {
        "user_id": gen_id("user"),
        "email": normalize_email(email),
        "name": name,
        "role": role,
        "employee_id": employee_id,
        "allowed_pages": allowed_pages,
        "dashboard_type": role,
        "created_at": now_utc().isoformat(),
        "updated_at": now_utc().isoformat(),
    }
    row.update(_password_db_fields(plain))
    return row


def _insert_login_or_raise(user_row: Dict[str, Any], plain_password: str) -> Dict[str, Any]:
    row = dict(user_row)
    inserted = sb_insert("users", row)
    # Older Supabase projects may lack the legacy `password` column.
    if not inserted and row.pop("password", None) is not None:
        inserted = sb_insert("users", row)
    if not inserted:
        raise HTTPException(
            status_code=500,
            detail=(
                "Could not save employee login in Supabase. "
                "Run supabase/employee_login_migration.sql in the Supabase SQL Editor, then try again."
            ),
        )
    user_id = inserted.get("user_id") or user_row.get("user_id")
    _verify_login_password_saved(user_id, plain_password)
    return inserted


def _update_user_login(user_id: str, patch: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if "allowed_pages" in patch:
        patch = {**patch, "allowed_pages": _sanitize_assignable_pages(_coerce_allowed_pages(patch["allowed_pages"]))}
    updated = sb_update("users", "user_id", user_id, patch)
    if not updated and "password" in patch:
        slim = {k: v for k, v in patch.items() if k != "password"}
        updated = sb_update("users", "user_id", user_id, slim)
    if updated:
        _user_profile_cache.pop(user_id, None)
    return updated


@api_router.post("/employees")
async def create_employee(p: EmployeeCreate, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager"])
    if p.role not in ROLES:
        raise HTTPException(status_code=400, detail="Invalid employee role")
    if not p.password or len(p.password.strip()) < 4:
        raise HTTPException(status_code=400, detail="Password is required (minimum 4 characters).")

    email = normalize_email(p.email)
    if not email:
        raise HTTPException(status_code=400, detail="Valid login email is required.")

    department = p.department or f"{p.role.title()} Department"
    allowed_pages = _sanitize_assignable_pages(
        p.allowed_pages if p.allowed_pages else _default_pages_for_role(p.role)
    )
    eid = gen_id("emp")

    existing_user = sb_select("users", {"email": f"eq.{email}", "select": "user_id", "limit": "1"})
    if existing_user:
        raise HTTPException(status_code=400, detail="A login with this email already exists.")

    existing_emp = sb_select("employees", {"email": f"eq.{email}", "select": "employee_id", "limit": "1"})
    if existing_emp:
        raise HTTPException(status_code=400, detail="An employee with this email already exists.")

    plain_password = p.password.strip()
    user_row = _build_login_row(p.name, email, plain_password, p.role, eid, allowed_pages)
    user_id = user_row["user_id"]
    _insert_login_or_raise(user_row, plain_password)

    e = {
        "employee_id": eid, "name": p.name, "email": email, "phone": p.phone,
        "role": p.role, "department": department, "active": True,
        "user_id": user_id, "allowed_pages": allowed_pages,
        "leads_assigned": 0, "leads_closed": 0, "last_login": None,
        "created_at": now_utc().isoformat(),
        "updated_at": now_utc().isoformat(),
    }
    result = sb_insert("employees", e)
    if not result:
        sb_delete("users", "user_id", user_id)
        raise HTTPException(status_code=500, detail="Could not create employee record.")
    invalidate_employees_cache()
    return result


@api_router.get("/employees")
async def list_employees(cu: User=Depends(get_current_user)):
    now = time.time()
    if _employees_cache["ts"] and (now - float(_employees_cache["ts"])) < _EMPLOYEES_CACHE_TTL:
        return _employees_cache["data"]
    rows = sb_select("employees", {
        "select": (
            "employee_id,name,email,phone,role,department,active,user_id,allowed_pages,created_at,"
            "last_lat,last_lng,last_seen_at"
        ),
        "order": "created_at.desc",
    })
    for row in rows:
        row["allowed_pages"] = _employee_allowed_pages(row)
    _employees_cache.update({"ts": now, "data": rows})
    return rows

@api_router.patch("/employees/{eid}")
async def update_employee(eid: str, p: EmployeeUpdate, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager"])
    raw = p.model_dump(exclude_unset=True)
    plain_password = (raw.pop("password", None) or "").strip()
    data = {k: v for k, v in raw.items() if v is not None}
    if "allowed_pages" in raw:
        # Empty list is a valid save (manager cleared all boxes).
        data["allowed_pages"] = _coerce_allowed_pages(raw["allowed_pages"])

    if data.get("role") and data["role"] not in ROLES:
        raise HTTPException(status_code=400, detail="Invalid employee role")
    if "email" in data:
        data["email"] = normalize_email(data["email"])
        if not data["email"]:
            raise HTTPException(status_code=400, detail="Valid login email is required.")

    rows = sb_select("employees", {"employee_id": f"eq.{eid}", "select": "*", "limit": "1"})
    employee = rows[0] if rows else None
    if not employee:
        raise HTTPException(404, "Employee not found")

    patch_keys = set(data.keys())
    access_only = patch_keys <= {"allowed_pages"} or patch_keys <= {"active", "allowed_pages"} or patch_keys == {"active"}
    login_fields_changed = bool(patch_keys) and not access_only

    if login_fields_changed and plain_password and len(plain_password) < 4:
        raise HTTPException(status_code=400, detail="Password must be at least 4 characters when provided.")

    if data.get("email") and data["email"] != employee.get("email"):
        clash = sb_select("users", {"email": f"eq.{data['email']}", "select": "user_id", "limit": "1"})
        resolved_uid = _resolve_user_id_for_employee(employee, employee.get("email") or "")
        if clash and clash[0].get("user_id") != resolved_uid:
            raise HTTPException(status_code=400, detail="Another login already uses this email.")
        clash_emp = sb_select("employees", {"email": f"eq.{data['email']}", "select": "employee_id", "limit": "1"})
        if clash_emp and clash_emp[0].get("employee_id") != eid:
            raise HTTPException(status_code=400, detail="Another employee already uses this email.")

    if "allowed_pages" in data:
        allowed_pages = data["allowed_pages"]
    else:
        allowed_pages = _employee_allowed_pages(employee)
    role = data.get("role") or employee.get("role")
    name = data.get("name") or employee.get("name")
    email = data.get("email") or employee.get("email")

    linked_user_id = _resolve_user_id_for_employee(employee, email)
    if linked_user_id and not employee.get("user_id"):
        data["user_id"] = linked_user_id

    # Sidebar / active-only: sync login row so employee sees new menu on next load.
    if "allowed_pages" in data and linked_user_id:
        _update_user_login(linked_user_id, {
            "allowed_pages": data["allowed_pages"],
            "updated_at": now_utc().isoformat(),
        })
        _user_profile_cache.pop(linked_user_id, None)
        for tok, sess in list(LOCAL_SESSIONS.items()):
            if (sess.get("user") or {}).get("user_id") == linked_user_id:
                merged = {**(sess.get("user") or {}), "allowed_pages": data["allowed_pages"]}
                LOCAL_SESSIONS[tok] = {**sess, "user": _finalize_user_session(_hydrate_allowed_pages_from_employee(merged))}

    if login_fields_changed:
        if not linked_user_id:
            if len(plain_password) < 4:
                raise HTTPException(
                    status_code=400,
                    detail="Password is required to create login for this employee (minimum 4 characters).",
                )
            user_row = _build_login_row(name, email, plain_password, role, eid, allowed_pages)
            inserted = _insert_login_or_raise(user_row, plain_password)
            linked_user_id = inserted["user_id"]
            data["user_id"] = linked_user_id
        else:
            data["user_id"] = linked_user_id
            user_patch: Dict[str, Any] = {
                "employee_id": eid,
                "updated_at": now_utc().isoformat(),
            }
            if plain_password:
                user_patch.update(_password_db_fields(plain_password))
            if "allowed_pages" in data:
                user_patch["allowed_pages"] = data["allowed_pages"]
            if "role" in data:
                user_patch["role"] = data["role"]
                user_patch["dashboard_type"] = data["role"]
            if "name" in data:
                user_patch["name"] = data["name"]
            if "email" in data:
                user_patch["email"] = data["email"]
            updated_user = _update_user_login(linked_user_id, user_patch)
            if not updated_user:
                raise HTTPException(status_code=500, detail="Could not update employee login.")
            if plain_password:
                _verify_login_password_saved(linked_user_id, plain_password)
                invalidate_sessions_for_user(linked_user_id)

    if data:
        data["updated_at"] = now_utc().isoformat()

    updated = sb_update("employees", "employee_id", eid, data) if data else employee
    if not updated:
        raise HTTPException(500, "Could not update employee.")
    updated["allowed_pages"] = _employee_allowed_pages(updated)
    invalidate_employees_cache()
    return updated

@api_router.delete("/employees/{eid}")
async def delete_employee(eid: str, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager"])
    rows = sb_select("employees", {"employee_id": f"eq.{eid}", "select": "user_id", "limit": "1"})
    linked_user_id = (rows[0] if rows else {}).get("user_id")
    sb_delete("employees", "employee_id", eid)
    if linked_user_id:
        sb_delete("users", "user_id", linked_user_id)
        invalidate_sessions_for_user(linked_user_id)
    return {"ok": True}

# ---- Templates & Campaigns ----
@api_router.post("/templates")
async def create_template(p: TemplateCreate, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    tid = gen_id("tpl")
    t = {"template_id": tid, "name": p.name, "body": p.body, "created_at": now_utc().isoformat()}
    result = sb_insert("templates", t)
    return result or t

@api_router.get("/templates")
async def list_templates(cu: User=Depends(get_current_user)):
    return sb_select("templates", {"select": "*", "order": "created_at.desc"})

@api_router.delete("/templates/{tid}")
async def delete_template(tid: str, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    sb_delete("templates", "template_id", tid)
    return {"ok": True}

@api_router.post("/campaigns")
async def create_campaign(p: CampaignCreate, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    cid = gen_id("cmp")
    c = {
        "campaign_id": cid, "name": p.name, "template_id": p.template_id,
        "audience": p.audience,
        "scheduled_at": p.scheduled_at.isoformat() if p.scheduled_at else None,
        "status": "draft", "sent_count": 0, "delivered_count": 0,
        "read_count": 0, "replied_count": 0, "created_at": now_utc().isoformat(),
    }
    result = sb_insert("campaigns", c)
    return result or c

@api_router.get("/campaigns")
async def list_campaigns(cu: User=Depends(get_current_user)):
    return sb_select("campaigns", {"select": "*", "order": "created_at.desc"})

@api_router.post("/campaigns/{cid}/send")
async def send_campaign(cid: str, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    camps = sb_select("campaigns", {"campaign_id": f"eq.{cid}", "select": "*"})
    if not camps: raise HTTPException(404, "Campaign not found")
    leads = sb_select("leads", {"select": "lead_id"})
    cnt = len(leads)
    data = {"status": "sent", "sent_count": cnt, "delivered_count": int(cnt * 0.95),
            "read_count": int(cnt * 0.7), "replied_count": int(cnt * 0.2)}
    updated = sb_update("campaigns", "campaign_id", cid, data)
    return updated or {**camps[0], **data}

@api_router.delete("/campaigns/{cid}")
async def delete_campaign(cid: str, cu: User=Depends(get_current_user)):
    ensure_roles(cu, ["admin", "manager", "marketing"])
    sb_delete("campaigns", "campaign_id", cid)
    return {"ok": True}

# ---- Stats / Dashboard ----
def _compute_dashboard_stats() -> Dict[str, Any]:
    """Shared pipeline stats for admin + manager main Dashboard."""
    fetched = sb_select_parallel({
        "bookings": ("bookings", {"select": "booking_amount,status"}),
        "visits": ("visits", {"select": "visit_id,status"}),
        "loans": ("loans", {"select": "loan_id,application_status,amount,bank_stage"}),
        "customers": ("customers", {"select": "customer_id,lead_id"}),
        "activities": ("activities", {"select": "activity_id,type"}),
        "employees": ("employees", {"select": "employee_id"}),
        "campaigns": ("campaigns", {"select": "campaign_id"}),
    })
    leads = fetch_all_leads_merged(
        "lead_id,name,phone,email,external_lead_id,source,stage,status,lead_type,follow_up_at,created_at"
    )
    bookings = fetched["bookings"]
    visits = fetched["visits"]
    loans = fetched["loans"]
    customers = fetched["customers"]
    activities = fetched["activities"]

    cache_bkg_ids = {b.get("booking_id") for b in SESSION_CACHE["bookings"]}
    bookings = SESSION_CACHE["bookings"] + [b for b in bookings if b.get("booking_id") not in cache_bkg_ids]
    bookings = [b for b in bookings if not is_legacy_skeleton_booking(b)]

    cache_vis_ids = {v.get("visit_id") for v in SESSION_CACHE["visits"]}
    visits = SESSION_CACHE["visits"] + [v for v in visits if v.get("visit_id") not in cache_vis_ids]

    followups = sb_select_all("visit_followups", {"select": "followup_id,status"})
    cache_followup_ids = {f.get("followup_id") for f in SESSION_CACHE["followups"]}
    followups = SESSION_CACHE["followups"] + [f for f in followups if f.get("followup_id") not in cache_followup_ids]

    cache_activity_ids = {a.get("activity_id") for a in SESSION_CACHE["activities"]}
    activities = SESSION_CACHE["activities"] + [a for a in activities if a.get("activity_id") not in cache_activity_ids]

    cache_lon_ids = {ln.get("loan_id") for ln in SESSION_CACHE["loans"]}
    loans = SESSION_CACHE["loans"] + [ln for ln in loans if ln.get("loan_id") not in cache_lon_ids]

    cache_customer_ids = {c.get("customer_id") for c in SESSION_CACHE["customers"]}
    customers = SESSION_CACHE["customers"] + [c for c in customers if c.get("customer_id") not in cache_customer_ids]

    pipeline_leads = [l for l in leads if is_pipeline_lead(l)]
    broker_pool = [l for l in leads if is_broker_pool_lead(l)]
    platform_stats = compute_platform_breakdown(leads)
    housing_count = platform_stats["housing"]
    meta_count = platform_stats["meta"]

    stage_dist = {s: 0 for s in STAGES}
    for l in pipeline_leads:
        if l.get("status") != "negative":
            st = l.get("stage", "new")
            if st == "new" and clean_text(l.get("assigned_to")):
                st = "assigned"
            stage_dist[st] = stage_dist.get(st, 0) + 1

    employees = fetched["employees"]
    campaigns = fetched["campaigns"]
    rev = sum(booking_brokerage_amount(b) for b in bookings)
    today = now_utc().date().isoformat()
    lead_buckets = {key: len(filter_lead_bucket(leads, key, today)) for key in LEAD_BUCKET_KEYS}
    follow_up_total = lead_buckets.get("follow_up", 0)
    pending_follow_up_total = sum(
        1 for f in followups
        if str(f.get("status", "scheduled")).lower() in ["scheduled", "pending", "open"]
    )

    return {
        "total_leads": platform_stats["total"],
        "broker_pool_leads": len(broker_pool),
        "housing_leads": housing_count,
        "meta_leads": meta_count,
        "positive_leads": lead_buckets.get("positive", 0),
        "registration_leads": lead_buckets.get("registration", 0),
        "negative_leads": lead_buckets.get("not_interested", 0),
        "new_leads": sum(1 for l in pipeline_leads if is_new_lead_stage(l)),
        "site_visits": len(visits),
        "completed_visits": sum(1 for v in visits if v.get("status") == "completed"),
        "bookings": len(bookings),
        "confirmed_bookings": sum(1 for b in bookings if b.get("status") == "confirmed"),
        "follow_ups": follow_up_total,
        "pending_follow_ups": pending_follow_up_total,
        "loans": len(loans),
        "disbursed_loans": sum(1 for l in loans if l.get("application_status") == "disbursed"),
        "converted_customers": max(len(customers), sum(1 for l in leads if l.get("stage") == "closed")),
        "employees": len(employees),
        "campaigns": len(campaigns),
        "revenue_pipeline": rev,
        "stage_distribution": stage_dist,
        "lead_buckets": lead_buckets,
    }


def _get_dashboard_stats_cached() -> Dict[str, Any]:
    now = time.time()
    if _dashboard_stats_cache["data"] is not None and (now - float(_dashboard_stats_cache["ts"])) < _STATS_CACHE_TTL_SEC:
        return _dashboard_stats_cache["data"]
    data = _compute_dashboard_stats()
    _dashboard_stats_cache.update({"ts": now, "data": data})
    return data


def _compute_dashboard_graph_sync() -> Dict[str, Any]:
    now = now_utc()
    graph_data = sb_select_parallel({
        "leads": ("leads", {"select": "lead_id,created_at,source,status,stage", "created_at": f"gte.{(now - timedelta(days=400)).isoformat()}"}),
        "bookings": ("bookings", {"select": "booking_id,brokerage_amount,agreement_status,created_at"}),
        "loans": ("loans", {"select": "loan_id,amount,created_at,application_status,bank_stage"}),
    })
    cache_lead_ids = {l.get("lead_id") for l in SESSION_CACHE["leads"]}
    leads = SESSION_CACHE["leads"] + [l for l in graph_data["leads"] if l.get("lead_id") not in cache_lead_ids]
    pipeline_leads = [l for l in dedupe_leads(leads) if is_pipeline_lead(l)]

    days_map = {(now - timedelta(days=29 - i)).strftime("%Y-%m-%d"): 0 for i in range(30)}
    for l in pipeline_leads:
        d = (l.get("created_at") or "")[:10]
        if d in days_map:
            days_map[d] += 1
    leads_by_day = [{"date": d, "count": cnt} for d, cnt in sorted(days_map.items())]

    months_map = {k: 0 for k in _last_n_month_keys(12)}
    for l in pipeline_leads:
        mk = (l.get("created_at") or "")[:7]
        if mk in months_map:
            months_map[mk] += 1
    leads_by_month = [{"month": mk, "count": months_map[mk]} for mk in sorted(months_map.keys())]

    bookings = graph_data["bookings"]
    cache_bkg_ids = {b.get("booking_id") for b in SESSION_CACHE["bookings"]}
    bookings = SESSION_CACHE["bookings"] + [b for b in bookings if b.get("booking_id") not in cache_bkg_ids]

    rev_months_map = {k: 0.0 for k in _last_n_month_keys(12)}
    for b in bookings:
        if is_legacy_skeleton_booking(b):
            continue
        d = b.get("created_at", "")[:7]
        if d in rev_months_map:
            rev_months_map[d] += booking_brokerage_amount(b)
    rev_by_month = [{"month": d, "revenue": rev} for d, rev in sorted(rev_months_map.items())]

    return {
        "leads_by_day": leads_by_day,
        "leads_by_month": leads_by_month,
        "revenue_by_month": rev_by_month,
        "total_leads_in_chart": sum(months_map.values()),
    }


def _stats_employees_sync() -> List[Dict[str, Any]]:
    now = time.time()
    if _employee_stats_cache["data"] is not None and (now - float(_employee_stats_cache["ts"])) < _STATS_CACHE_TTL_SEC:
        return _employee_stats_cache["data"]
    employees = sb_select("employees", {"select": "employee_id,name,email,role,department,active,user_id"})
    db_activities = sb_select("activities", {
        "select": "user_id,created_at,type",
        "order": "created_at.desc",
        "limit": "2000",
    })
    cache_act_ids = {a.get("activity_id") for a in SESSION_CACHE["activities"]}
    activities = SESSION_CACHE["activities"] + [a for a in db_activities if a.get("activity_id") not in cache_act_ids]
    all_leads = fetch_all_leads_merged(EMPLOYEE_WORKFLOW_LEAD_SELECT)

    emp_stats = []
    for e in employees:
        eid = e["employee_id"]
        emp_acts = [a for a in activities if a.get("user_id") in {eid, e.get("user_id")}]
        emp_leads = leads_for_employee_record(e, all_leads)
        assignment = compute_employee_assignment_stats(emp_leads, e.get("role"))
        last_activity = max([a["created_at"] for a in emp_acts]) if emp_acts else None

        emp_stats.append({
            "employee_id": eid,
            "name": e["name"],
            "email": e["email"],
            "role": e["role"],
            "department": e.get("department", ""),
            "active": e.get("active", True),
            "actions_total": len(emp_acts),
            "leads_total": assignment["assigned_total"],
            "last_activity": last_activity,
            **assignment,
            "positives": assignment["assigned_positive"],
            "negatives": assignment["assigned_not_interested"],
            "followups": assignment["assigned_follow_ups"],
            "visits": assignment["emp_visited"],
            "bookings_done": assignment["emp_booking_done"],
            "loans_done": sum(1 for a in emp_acts if "loan" in str(a.get("type"))),
            "closed_deals": assignment["assigned_completed"],
            "call_notes": sum(1 for a in emp_acts if "call" in str(a.get("type"))),
        })
    _employee_stats_cache.update({"ts": now, "data": emp_stats})
    return emp_stats


@api_router.get("/stats/dashboard")
async def stats_dashboard(cu: User=Depends(get_current_user)):
    ensure_owner_dashboard(cu)
    return await asyncio.to_thread(_get_dashboard_stats_cached)


@api_router.get("/stats/dashboard-bundle")
async def stats_dashboard_bundle(cu: User = Depends(get_current_user)):
    """One round-trip for main Dashboard — admin (full) + manager (team, no revenue UI)."""
    ensure_main_dashboard(cu)
    stats, graph, employees, recent, leads_page = await asyncio.gather(
        asyncio.to_thread(_get_dashboard_stats_cached),
        asyncio.to_thread(_compute_dashboard_graph_sync),
        asyncio.to_thread(_stats_employees_sync),
        list_recent_leads(20, cu),
        list_leads(limit=80, offset=0, cu=cu),
    )
    return {
        "stats": stats,
        "graph": graph,
        "employees": employees,
        "recent_leads": recent,
        "leads": leads_page,
    }


@api_router.get("/stats/me-bundle")
async def stats_me_bundle(cu: User = Depends(get_current_user)):
    """One round-trip for My Dashboard (manager gets team performance too)."""
    me, graph = await asyncio.gather(
        stats_me(cu),
        asyncio.to_thread(_compute_dashboard_graph_sync),
    )
    payload: Dict[str, Any] = {"me": me, "graph": graph}
    if cu.role == "manager":
        payload["employees"] = await asyncio.to_thread(_stats_employees_sync)
    return payload


@api_router.get("/stats/leads-by-source")
async def stats_leads_by_source(cu: User=Depends(get_current_user)):
    leads = fetch_all_leads_merged("lead_id,source,stage,status,created_at")
    source_map: dict = {}
    for l in leads:
        src = (l.get("source") or "direct").strip().lower()
        if src not in source_map:
            source_map[src] = {"source": src, "count": 0, "active": 0, "negative": 0}
        source_map[src]["count"] += 1
        if l.get("status") == "negative":
            source_map[src]["negative"] += 1
        else:
            source_map[src]["active"] += 1
    result = sorted(source_map.values(), key=lambda x: x["count"], reverse=True)
    return {"total": len(leads), "sources": result}

@api_router.get("/stats/leads-by-platform")
async def stats_leads_by_platform(cu: User=Depends(get_current_user)):
    """Dashboard breakdown: manual, housing (Housing.com), meta (Facebook/Instagram)."""
    leads = fetch_all_leads_merged("lead_id,phone,email,external_lead_id,source,stage,status,created_at")
    breakdown = compute_platform_breakdown(leads)
    return {
        "total": breakdown["total"],
        "broker_pool": breakdown["broker_pool"],
        "platforms": breakdown["platforms"],
    }


@api_router.get("/stats/me/leads-by-platform")
async def stats_me_leads_by_platform(cu: User = Depends(get_current_user)):
    """My Dashboard — platform breakdown for leads assigned to this employee."""
    emp_id = resolve_employee_id(cu)
    if not emp_id:
        raise HTTPException(status_code=400, detail="No employee profile linked to this login.")
    leads = fetch_employee_assigned_leads(
        cu, EMPLOYEE_WORKFLOW_LEAD_SELECT
    )
    my_queue_leads = filter_employee_my_queue_leads(leads)
    breakdown = compute_platform_breakdown(my_queue_leads)
    return {
        "total": breakdown["total"],
        "broker_pool": breakdown["broker_pool"],
        "platforms": breakdown["platforms"],
        "scope": "mine",
        "employee_id": emp_id,
    }


@api_router.get("/stats/me/leads/by-platform/{platform}")
async def list_me_leads_by_platform(
    platform: str,
    limit: int = 200,
    offset: int = 0,
    status_filter: Optional[str] = None,
    cu: User = Depends(get_current_user),
):
    """My Dashboard — assigned leads for a platform bucket: manual, housing, or meta."""
    emp_id = resolve_employee_id(cu)
    if not emp_id:
        raise HTTPException(status_code=400, detail="No employee profile linked to this login.")
    platform_key = platform.strip().lower()
    if platform_key not in {"manual", "housing", "meta", "other"}:
        raise HTTPException(status_code=400, detail="Platform must be manual, housing, meta, or other")

    limit = min(max(limit, 1), 500)
    offset = max(offset, 0)
    all_leads = fetch_employee_assigned_leads(
        cu,
        "lead_id,name,phone,email,source,stage,status,priority,call_status,budget,location,property_type,"
        "notes,external_lead_id,external_created_at,raw_payload,created_at,assigned_to,assigned_at,"
        "last_employee_action_at,follow_up_at,lead_type",
    )
    cleaned = clean_leads_for_platform_stats(filter_employee_my_queue_leads(all_leads))
    filtered = [l for l in cleaned if lead_matches_platform(l, platform_key)]
    sf = (status_filter or "").strip().lower()
    if sf == "positive":
        filtered = [l for l in filtered if l.get("stage") in ["positive", "site_visit", "booking", "loan", "registration", "closed"] and l.get("status") != "negative"]
    elif sf in ("not_interested", "negative"):
        filtered = [l for l in filtered if l.get("status") == "negative"]
    elif sf == "registration":
        filtered = [l for l in filtered if l.get("stage") == "registration"]
    elif sf == "booking":
        filtered = [l for l in filtered if l.get("stage") in ["booking", "loan"]]
    filtered.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    page = enrich_leads_display_fields(filtered[offset:offset + limit])
    return {
        "platform": platform_key,
        "label": PLATFORM_LABELS.get(platform_key, "Other Sources"),
        "total": len(filtered),
        "leads": page,
        "scope": "mine",
        "employee_id": emp_id,
    }

def _last_n_month_keys(n: int = 12) -> List[str]:
    """Calendar month keys YYYY-MM ending with current month."""
    today = now_utc().date()
    y, m = today.year, today.month
    keys: List[str] = []
    for _ in range(n):
        keys.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    return list(reversed(keys))


@api_router.get("/stats/dashboard/graph")
async def stats_dashboard_graph(cu: User=Depends(get_current_user)):
    return await asyncio.to_thread(_compute_dashboard_graph_sync)

@api_router.get("/activities")
async def list_activities(limit: int = 50, cu: User=Depends(get_current_user)):
    activities = sb_select("activities", {"select": "*", "order": "created_at.desc", "limit": str(limit)})
    # Deduplicate
    cache_act_ids = {a.get("activity_id") for a in SESSION_CACHE["activities"]}
    return SESSION_CACHE["activities"] + [a for a in activities if a.get("activity_id") not in cache_act_ids]

@api_router.get("/stats/me")
async def stats_me(cu: User=Depends(get_current_user)):
    activity_keys = actor_activity_keys(cu)
    db_activities = sb_select("activities", {
        "select": "activity_id,user_id,created_at,type,text,lead_id",
        "order": "created_at.desc",
        "limit": "1500",
    })
    cache_act_ids = {a.get("activity_id") for a in SESSION_CACHE["activities"]}
    all_activities = SESSION_CACHE["activities"] + [a for a in db_activities if a.get("activity_id") not in cache_act_ids]
    if cu.role == "admin":
        activities = all_activities
    else:
        activities = [a for a in all_activities if a.get("user_id") in activity_keys]

    emp_id = resolve_employee_id(cu)
    employee_row = None
    if emp_id:
        emps = sb_select("employees", {"employee_id": f"eq.{emp_id}", "select": "employee_id,name,email,role,department", "limit": "1"})
        if emps:
            employee_row = emps[0]
    emp_role = (employee_row or {}).get("role") or cu.role

    if cu.role == "manager" and not emp_id:
        # Manager workspace uses team panels — not company-wide personal KPI totals.
        leads: List[Dict[str, Any]] = []
        my_queue_leads: List[Dict[str, Any]] = []
        platform_breakdown = {"total": 0, "housing": 0, "meta": 0, "manual": 0, "other": 0}
    elif cu.role != "admin" and emp_id:
        leads = fetch_employee_assigned_leads(cu, EMPLOYEE_WORKFLOW_LEAD_SELECT)
        my_queue_leads = filter_employee_my_queue_leads(leads)
        platform_breakdown = compute_platform_breakdown(my_queue_leads)
    else:
        leads = fetch_all_leads_merged(EMPLOYEE_WORKFLOW_LEAD_SELECT)
        my_queue_leads = filter_employee_my_queue_leads(leads)
        platform_breakdown = compute_platform_breakdown(my_queue_leads)

    assignment = compute_employee_assignment_stats(leads, emp_role)
    positives = assignment["assigned_positive"]
    negative = assignment["assigned_not_interested"]
    followups = assignment["assigned_follow_ups"]
    visits = sum(1 for a in activities if a.get("type") == "site_visit_scheduled" or "visit" in str(a.get("type")))
    bookings_done = sum(1 for a in activities if "booking" in str(a.get("type")))
    loans_done = sum(1 for a in activities if "loan" in str(a.get("type")))
    closed_deals = assignment["assigned_completed"]

    score = min(10, positives * 1 + followups * 0.5 + visits * 2 + bookings_done * 3 + loans_done * 2 + closed_deals * 4)
    if not activities and not leads:
        score = 0

    hot = sum(1 for l in dedupe_leads(leads) if l.get("stage") in ["positive", "site_visit", "booking", "loan", "registration"] and l.get("status") != "negative")
    cold = sum(1 for l in dedupe_leads(leads) if l.get("stage") in ["new", "assigned", "contacted"] and l.get("status") != "negative")

    missed_rows = filter_employee_missed_leads(leads)
    missed_rows.sort(key=lambda l: l.get("assigned_at") or l.get("updated_at") or l.get("created_at") or "", reverse=False)
    missed_preview = enrich_leads_display_fields(missed_rows[:20])

    return {
        "employee": employee_row,
        "employee_id": emp_id,
        "role": emp_role or cu.role,
        "missed_leads": missed_preview,
        "missed_leads_total": len(missed_rows),
        "personal": {
            "actions_total": len(activities),
            "leads_total": len(my_queue_leads) if emp_id or cu.role == "admin" else 0,
            "assigned_active": assignment["assigned_active"],
            "assigned_completed": assignment["assigned_completed"],
            "assigned_queue": assignment["assigned_queue"],
            "assigned_in_progress": assignment["assigned_in_progress"],
            "assigned_follow_ups": assignment["assigned_follow_ups"],
            "emp_hot": assignment["emp_hot"],
            "emp_visited": assignment["emp_visited"],
            "emp_not_interested": assignment["emp_not_interested"],
            "emp_booking_done": assignment["emp_booking_done"],
            "emp_low_budget": assignment["emp_low_budget"],
            "emp_ringing": assignment["emp_ringing"],
            "emp_follow_ups": assignment["emp_follow_ups"],
            "emp_missed_leads": len(missed_rows),
            "positives": positives,
            "negatives": negative,
            "followups": followups,
            "visits": visits,
            "bookings_done": bookings_done,
            "loans_done": loans_done,
            "closed_deals": closed_deals,
            "call_notes": sum(1 for a in activities if "call" in str(a.get("type"))),
            "score_10": score,
            "last_activity": activities[0]["created_at"] if activities else None,
            "housing_leads": platform_breakdown.get("housing", 0),
            "meta_leads": platform_breakdown.get("meta", 0),
            "manual_leads": platform_breakdown.get("manual", 0),
            "other_leads": platform_breakdown.get("other", 0),
        },
        "leads": {"hot": hot, "cold": cold, "negative": negative, "closed": closed_deals},
        "recent_activities": activities[:15],
    }

EMPLOYEE_METRIC_KEYS = [
    "active", "hot", "visited", "not_interested", "booking_done", "low_budget",
    "ringing", "follow_ups", "missed_leads",
]

PERSONAL_ACTIVITY_METRICS = [
    "total", "queue", "follow_ups", "completed", "closed_deals", "positive", "negatives",
    "call_notes", "bookings_done", "loans_done", "active", "hot", "visited",
    "not_interested", "booking_done", "low_budget", "ringing", "missed_leads",
]

PERSONAL_ACTIVITY_LABELS: Dict[str, str] = {
    "total": "My Queue — Active Leads",
    "queue": "My Queue",
    "follow_ups": "Follow-ups",
    "completed": "Completed",
    "closed_deals": "Closed Deals",
    "positive": "Positive Leads",
    "negatives": "Not Interested",
    "call_notes": "Call Notes",
    "bookings_done": "Bookings",
    "loans_done": "Loans",
    "active": "Active Leads",
    "hot": "Hot Leads",
    "visited": "Visited",
    "not_interested": "Not Interested",
    "booking_done": "Booking Done",
    "low_budget": "Low Budget",
    "ringing": "Ringing",
    "missed_leads": "Missed Lead",
}


def _employee_may_view_stats(cu: User, employee_id: str) -> bool:
    if cu.role in ("admin", "manager"):
        return True
    own = cu.acting_as_employee_id or cu.employee_id
    return bool(own and own == employee_id)


def filter_personal_activity(
    metric_key: str,
    emp_leads: List[Dict[str, Any]],
    activities: List[Dict[str, Any]],
    role: Optional[str],
) -> Dict[str, Any]:
    """Drill-down lists for My Dashboard activity boxes — counts match stats_me."""
    key = (metric_key or "").strip().lower()
    rows = dedupe_leads(emp_leads)
    if key in ("total", "all", "leads_total", "assigned_total"):
        return {"kind": "leads", "items": filter_employee_my_queue_leads(rows)}
    if key in ("queue", "assigned_queue"):
        items = filter_employee_queue_leads(rows, role)
        return {"kind": "leads", "items": items}
    if key in ("follow_ups", "followups", "assigned_follow_ups"):
        items = filter_employee_follow_up_leads(rows)
        return {"kind": "leads", "items": items}
    if key in ("completed", "assigned_completed", "closed_deals"):
        items = [l for l in rows if l.get("stage") == "closed"]
        return {"kind": "leads", "items": items}
    if key in ("positive", "positives"):
        items = [
            l for l in rows
            if l.get("stage") in ["positive", "site_visit", "booking", "loan", "registration"]
            and l.get("status") != "negative"
        ]
        return {"kind": "leads", "items": items}
    if key in ("negatives", "not_interested"):
        items = filter_employee_metric_leads(rows, "not_interested")
        return {"kind": "leads", "items": items}
    if key == "call_notes":
        items = [a for a in activities if "call" in str(a.get("type", "")).lower()]
        return {"kind": "activities", "items": items}
    if key == "bookings_done":
        items = [a for a in activities if "booking" in str(a.get("type", "")).lower()]
        return {"kind": "activities", "items": items}
    if key == "loans_done":
        items = [a for a in activities if "loan" in str(a.get("type", "")).lower()]
        return {"kind": "activities", "items": items}
    if key in ("missed_leads", "missed"):
        return {"kind": "leads", "items": filter_employee_missed_leads(rows)}
    if key in EMPLOYEE_METRIC_KEYS:
        return {"kind": "leads", "items": filter_employee_metric_leads(rows, key)}
    raise HTTPException(400, detail=f"Unknown activity metric: {key}")


@api_router.get("/stats/me/activity/{metric_key}")
async def stats_me_activity(metric_key: str, limit: int = 500, cu: User = Depends(get_current_user)):
    """My Dashboard — tap an activity box to see the matching lead/activity list."""
    key = (metric_key or "").strip().lower()
    if key not in PERSONAL_ACTIVITY_METRICS:
        raise HTTPException(400, detail=f"metric must be one of: {', '.join(PERSONAL_ACTIVITY_METRICS)}")
    limit = min(max(limit, 1), 500)
    emp_id = resolve_employee_id(cu)
    if not emp_id and cu.role != "admin":
        raise HTTPException(400, detail="No employee profile linked to this login.")

    employee_row = None
    if emp_id:
        emps = sb_select("employees", {"employee_id": f"eq.{emp_id}", "select": "role", "limit": "1"})
        if emps:
            employee_row = emps[0]
    emp_role = (employee_row or {}).get("role") or cu.role

    activity_keys = actor_activity_keys(cu)
    db_activities = sb_select("activities", {
        "select": "activity_id,user_id,created_at,type,text,lead_id",
        "order": "created_at.desc",
        "limit": "1500",
    })
    cache_act_ids = {a.get("activity_id") for a in SESSION_CACHE["activities"]}
    all_activities = SESSION_CACHE["activities"] + [a for a in db_activities if a.get("activity_id") not in cache_act_ids]
    if cu.role == "admin" and not emp_id:
        activities = all_activities
        emp_leads = fetch_all_leads_merged(EMPLOYEE_WORKFLOW_LEAD_SELECT)
    else:
        activities = [a for a in all_activities if a.get("user_id") in activity_keys]
        emp_leads = fetch_employee_assigned_leads(cu, EMPLOYEE_WORKFLOW_LEAD_SELECT) if emp_id else []

    result = filter_personal_activity(key, emp_leads, activities, emp_role)
    lead_map = {l.get("lead_id"): l for l in emp_leads if l.get("lead_id")}
    if result["kind"] == "leads":
        items = enrich_leads_display_fields(
            sorted(result["items"], key=lambda l: l.get("created_at") or "", reverse=True)[:limit]
        )
    else:
        items = sorted(result["items"], key=lambda a: a.get("created_at") or "", reverse=True)[:limit]
        for act in items:
            lid = act.get("lead_id")
            if lid and lid in lead_map:
                act["lead_name"] = lead_map[lid].get("name")
                act["lead_phone"] = lead_map[lid].get("phone")

    return {
        "metric": key,
        "label": PERSONAL_ACTIVITY_LABELS.get(key, key.replace("_", " ").title()),
        "kind": result["kind"],
        "total": len(result["items"]),
        "items": items,
    }


@api_router.get("/leads/employee/{employee_id}/metric/{metric_key}")
async def list_employee_metric_leads(
    employee_id: str,
    metric_key: str,
    limit: int = 500,
    cu: User = Depends(get_current_user),
):
    """Drill-down list for dashboard employee performance boxes."""
    if not _employee_may_view_stats(cu, employee_id):
        raise HTTPException(403, detail="Not allowed to view this employee's leads.")
    key = (metric_key or "").strip().lower()
    if key not in EMPLOYEE_METRIC_KEYS:
        raise HTTPException(400, detail=f"metric must be one of: {', '.join(EMPLOYEE_METRIC_KEYS)}")
    limit = min(max(limit, 1), 500)
    all_leads = fetch_all_leads_merged(EMPLOYEE_WORKFLOW_LEAD_SELECT)
    employees = sb_select("employees", {"employee_id": f"eq.{employee_id}", "select": "employee_id,name,user_id", "limit": "1"})
    employee = employees[0] if employees else {"employee_id": employee_id}
    emp_leads = leads_for_employee_record(employee, all_leads)
    if key in ("missed_leads", "missed"):
        filtered = filter_employee_missed_leads(emp_leads)
    else:
        filtered = filter_employee_metric_leads(emp_leads, key)
    filtered.sort(key=lambda l: l.get("created_at") or "", reverse=True)
    return {
        "employee_id": employee_id,
        "metric": key,
        "total": len(filtered),
        "leads": enrich_leads_display_fields(filtered[:limit]),
    }


@api_router.get("/stats/employees")
async def stats_employees(cu: User=Depends(get_current_user)):
    return await asyncio.to_thread(_stats_employees_sync)

# ---- Health & Wiring ----
@api_router.get("/")
async def root(): return {"app": "Umang Hometech LLP CRM", "status": "ok", "database": "supabase"}

app.include_router(api_router)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')

_housing_sync_task: Optional[asyncio.Task] = None
_facebook_sync_task: Optional[asyncio.Task] = None


def run_facebook_sync_background(reason: str = "auto") -> Dict[str, Any]:
    """Resync pending Meta webhooks + import leads from the last 24 hours."""
    if not FACEBOOK_PAGE_ACCESS_TOKEN:
        return {"status": "skipped", "reason": "FACEBOOK_PAGE_ACCESS_TOKEN not configured"}
    try:
        resync = _facebook_resync_pending_impl()
    except HTTPException as exc:
        return {"status": "error", "detail": str(exc.detail)[:180]}
    import_result: Dict[str, Any] = {"fetched": 0, "created": 0, "duplicates": 0}
    try:
        import_result = _facebook_import_impl(days=1, limit=100)
    except HTTPException as exc:
        import_result = {"error": str(exc.detail)[:180]}
    except Exception as exc:
        import_result = {"error": str(exc)[:180]}
    return {
        "status": "success",
        "reason": reason,
        "resync": resync,
        "import_recent": {
            "fetched": import_result.get("fetched", 0),
            "created": import_result.get("created", 0),
            "duplicates": import_result.get("duplicates", 0),
            "error": import_result.get("error"),
        },
    }


async def _facebook_background_sync_loop() -> None:
    """Retry pending Meta webhooks and import recent Lead Ad submissions."""
    await asyncio.sleep(45)
    while True:
        try:
            if FACEBOOK_AUTO_SYNC_ENABLED and FACEBOOK_PAGE_ACCESS_TOKEN:
                result = await asyncio.to_thread(run_facebook_sync_background, "auto")
                created = (result.get("resync") or {}).get("created", 0)
                created += (result.get("import_recent") or {}).get("created", 0)
                if created:
                    logging.info("Facebook auto-sync imported %s new lead(s)", created)
                elif result.get("status") == "error":
                    logging.warning("Facebook auto-sync error: %s", result.get("detail"))
        except asyncio.CancelledError:
            break
        except Exception:
            logging.exception("Facebook background sync loop error")
        await asyncio.sleep(max(FACEBOOK_AUTO_SYNC_INTERVAL_SEC, 120))


async def _housing_background_sync_loop() -> None:
    """Pull new Housing.com leads every HOUSING_AUTO_SYNC_INTERVAL_SEC (default 5 min)."""
    await asyncio.sleep(30)
    while True:
        try:
            if HOUSING_AUTO_SYNC_ENABLED and HOUSING_PROFILE_ID and HOUSING_ENCRYPTION_KEY:
                result = await asyncio.to_thread(run_housing_sync_background, "auto")
                created = result.get("created") or []
                if isinstance(created, list) and created:
                    logging.info("Housing auto-sync imported %s new lead(s)", len(created))
                elif result.get("status") == "error":
                    logging.warning("Housing auto-sync error: %s", result.get("detail"))
        except asyncio.CancelledError:
            break
        except Exception:
            logging.exception("Housing background sync loop error")
        await asyncio.sleep(max(HOUSING_AUTO_SYNC_INTERVAL_SEC, 60))


@app.on_event("startup")
async def start_housing_background_sync() -> None:
    global _housing_sync_task, _facebook_sync_task

    async def _warm_caches() -> None:
        await asyncio.sleep(2)
        try:
            await asyncio.to_thread(fetch_all_leads_merged, LEADS_CANONICAL_SELECT)
            await asyncio.to_thread(_get_dashboard_stats_cached)
            logging.info("Startup cache warm completed")
        except Exception:
            logging.exception("Startup cache warm failed")

    asyncio.create_task(_warm_caches())

    if FACEBOOK_AUTO_SYNC_ENABLED and FACEBOOK_PAGE_ACCESS_TOKEN:
        _facebook_sync_task = asyncio.create_task(_facebook_background_sync_loop())
        logging.info(
            "Facebook auto-sync started — every %ss (resync webhooks + import last 24h)",
            FACEBOOK_AUTO_SYNC_INTERVAL_SEC,
        )
    elif not FACEBOOK_PAGE_ACCESS_TOKEN:
        logging.warning("Facebook auto-sync skipped — set FACEBOOK_PAGE_ACCESS_TOKEN on Render")
    else:
        logging.info("Facebook auto-sync disabled (FACEBOOK_AUTO_SYNC_ENABLED=false)")

    if not HOUSING_AUTO_SYNC_ENABLED:
        logging.info("Housing auto-sync disabled (HOUSING_AUTO_SYNC_ENABLED=false)")
        return
    if not HOUSING_PROFILE_ID or not HOUSING_ENCRYPTION_KEY:
        logging.warning("Housing auto-sync skipped — set HOUSING_PROFILE_ID and HOUSING_ENCRYPTION_KEY")
        return
    _housing_sync_task = asyncio.create_task(_housing_background_sync_loop())
    logging.info(
        "Housing auto-sync started — every %ss, initial window %ss",
        HOUSING_AUTO_SYNC_INTERVAL_SEC,
        HOUSING_POLL_INITIAL_WINDOW_SEC,
    )


@app.on_event("shutdown")
async def stop_housing_background_sync() -> None:
    global _housing_sync_task, _facebook_sync_task
    if _facebook_sync_task:
        _facebook_sync_task.cancel()
        try:
            await _facebook_sync_task
        except asyncio.CancelledError:
            pass
        _facebook_sync_task = None
    if _housing_sync_task:
        _housing_sync_task.cancel()
        try:
            await _housing_sync_task
        except asyncio.CancelledError:
            pass
        _housing_sync_task = None
