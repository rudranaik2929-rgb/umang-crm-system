"""Team Lead hierarchy: Admin pool -> Team Lead -> Employee.

Covers the ownership rule (no lead is ever owned by nobody or by two teams),
backend team scoping, assignment permissions, the emergency transfer, employee
removal, monthly exports and the security fixes that came with the change.
"""
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import main as core

T = "2026-10-0{}T06:30:00+00:00"


def _emp(eid, role, team_lead_id=None, regions=None, **extra):
    row = {
        "employee_id": eid, "name": extra.pop("name", eid.upper()), "email": f"{eid}@test.com",
        "role": role, "active": True, "user_id": f"user_{eid}", "department": role,
        "location": "Virar", "team_lead_id": team_lead_id, "regions": regions or [],
        "allowed_pages": None,
    }
    row.update(extra)
    return row


def _lead(lid, day=5, **extra):
    row = {
        "lead_id": lid, "name": f"Lead {lid}", "phone": f"98{abs(hash(lid)) % 10**8:08d}", "email": None,
        "source": "manual", "stage": "new", "status": "active", "lead_type": "standard",
        "priority": None, "call_status": None, "budget": None, "location": None, "property_type": None,
        "assigned_to": None, "assigned_at": None, "assigned_by": None, "last_employee_action_at": None,
        "site_visitor_id": None, "site_visit_assigned_by": None, "site_visit_assigned_at": None,
        "follow_up_at": None, "created_at": T.format(day), "updated_at": T.format(day),
        "starred": False, "brokerage_amount": None, "raw_payload": None, "external_lead_id": None,
        "team_lead_id": None, "team_lead_assigned_at": None, "team_lead_assigned_by": None,
    }
    row.update(extra)
    return row


EMPLOYEES = [
    _emp("tl_a", "team_lead", regions=["Virar"], name="Asha Lead"),
    _emp("tl_b", "team_lead", regions=["Vasai"], name="Bharat Lead"),
    _emp("e1", "telecaller", team_lead_id="tl_a"),
    _emp("e2", "telecaller", team_lead_id="tl_b"),
    _emp("e3", "telecaller"),  # no team yet
]

LEADS = [
    _lead("L1", day=1, location="Virar West"),                       # admin pool
    _lead("L2", day=1, location="Vasai East"),                       # admin pool
    _lead("L3", day=2, team_lead_id="tl_a"),                         # A pool
    _lead("L4", day=2, assigned_to="e1", stage="assigned"),          # A via legacy employee
    _lead("L5", day=3, team_lead_id="tl_b", assigned_to="e2", stage="assigned"),  # B
    _lead("L6", day=3, assigned_to="e3", stage="assigned"),          # employee with no team
    _lead("L7", day=4, team_lead_id="tl_b"),                         # B pool
]


class _Req:
    def __init__(self, headers=None):
        self.headers = headers or {}


def _user(role, eid=None, email=None):
    return core.User(
        user_id=f"user_{eid or role}", email=email or f"{eid or role}@test.com", name=(eid or role).upper(),
        role=role, employee_id=eid, acting_as_employee_id=eid, created_at=core.now_utc(),
    )


def _login(user, headers=None):
    # Same finishing step as production get_current_user (scope + act-as policy).
    async def _dep():
        return core._finish_request_user(user.model_dump(), _Req(headers))
    core.app.dependency_overrides[core.get_current_user] = _dep


@pytest.fixture
def db(fake_db):
    fake_db.reset(seed={
        "employees": EMPLOYEES,
        "leads": LEADS,
        "users": [{"user_id": "user_admin", "email": "admin@test.com", "name": "ADMIN", "role": "admin"}],
        "regions": [dict(r) for r in core.DEFAULT_REGIONS],
    })
    core.invalidate_team_directory()
    return fake_db


@pytest.fixture
def admin(db):
    _login(_user("admin", email="admin@test.com"))
    return TestClient(core.app)


@pytest.fixture
def lead_a(db):
    _login(_user("team_lead", "tl_a"))
    return TestClient(core.app)


def _row(db, table, key, value):
    return next(r for r in db.table(table) if r.get(key) == value)


# ---- ownership rule -------------------------------------------------------

def test_every_lead_has_exactly_one_owner(db):
    d = core.get_team_directory(force=True)
    owners = {l["lead_id"]: d.lead_team_lead_id(l) for l in LEADS}
    assert owners == {"L1": None, "L2": None, "L3": "tl_a", "L4": "tl_a", "L5": "tl_b", "L6": None, "L7": "tl_b"}
    facets = core.compute_team_facets(LEADS, d)
    assert facets["admin_pool"] == 2
    assert facets["no_team_assigned"] == 1
    assert facets["teams"]["tl_a"] == {"total": 2, "pool": 1, "assigned": 1}
    assert facets["teams"]["tl_b"] == {"total": 2, "pool": 1, "assigned": 1}
    total = facets["admin_pool"] + facets["no_team_assigned"] + sum(t["total"] for t in facets["teams"].values())
    assert total == facets["pipeline_total"] == len(LEADS)


def test_legacy_manager_role_behaves_as_team_lead(db):
    user = core._finish_request_user(_user("manager", "tl_a").model_dump(), _Req())
    assert user.role == "team_lead"
    core.ensure_roles(user, ["admin", "team_lead"])  # does not raise
    assert core.current_team_scope() == "tl_a"


# ---- scoping --------------------------------------------------------------

def test_team_lead_only_sees_own_team(lead_a):
    r = lead_a.get("/api/leads/assign-workspace", params={"limit": 100})
    assert r.status_code == 200, r.text
    body = r.json()
    assert {l["lead_id"] for l in body["leads"]} == {"L3", "L4"}
    assert {e["employee_id"] for e in body["assignable"]} == {"tl_a", "e1"}
    assert [t["employee_id"] for t in body["team_leads"]] == ["tl_a"]
    assert body["viewer"]["team_lead_id"] == "tl_a"


def test_team_lead_employee_list_is_scoped(lead_a):
    r = lead_a.get("/api/employees")
    assert r.status_code == 200
    assert {e["employee_id"] for e in r.json()} == {"tl_a", "e1"}


def test_team_lead_cannot_open_other_team_lead(lead_a):
    assert lead_a.get("/api/leads/L5").status_code == 403
    assert lead_a.get("/api/leads/L4").status_code == 200


def test_admin_team_and_region_filters(admin):
    pool = admin.get("/api/leads/assign-workspace", params={"team": "admin_pool", "team_view": "pool"}).json()
    assert {l["lead_id"] for l in pool["leads"]} == {"L1", "L2"}
    b_pool = admin.get("/api/leads/assign-workspace", params={"team": "tl_b", "team_view": "pool"}).json()
    assert {l["lead_id"] for l in b_pool["leads"]} == {"L7"}
    virar = admin.get("/api/leads/assign-workspace", params={"region": "Virar", "team": "admin_pool"}).json()
    assert {l["lead_id"] for l in virar["leads"]} == {"L1"}
    names = {l["lead_id"]: l["team_lead_name"] for l in admin.get("/api/leads/assign-workspace").json()["leads"]}
    assert names["L4"] == "Asha Lead" and names["L1"] is None


# ---- assignment permissions ----------------------------------------------

def test_admin_routes_leads_to_team_lead_pool(admin, db):
    r = admin.post("/api/leads/bulk-manage", json={"lead_ids": ["L1", "L6"], "team_lead_id": "tl_b"})
    assert r.status_code == 200, r.text
    assert r.json()["updated_count"] == 2
    l1, l6 = _row(db, "leads", "lead_id", "L1"), _row(db, "leads", "lead_id", "L6")
    assert l1["team_lead_id"] == "tl_b" and l1["assigned_to"] is None
    # e3 is not in Team B, so the lead lands in Team B's pool for re-assignment.
    assert l6["team_lead_id"] == "tl_b" and l6["assigned_to"] is None and l6["stage"] == "new"


def test_assigning_employee_stamps_their_team(admin, db):
    r = admin.post("/api/leads/bulk-manage", json={"lead_ids": ["L2"], "assigned_to": "e2"})
    assert r.status_code == 200, r.text
    l2 = _row(db, "leads", "lead_id", "L2")
    assert l2["assigned_to"] == "e2" and l2["team_lead_id"] == "tl_b"


def test_team_lead_assigns_within_team(lead_a, db):
    r = lead_a.post("/api/leads/bulk-manage", json={"lead_ids": ["L3"], "assigned_to": "e1"})
    assert r.status_code == 200, r.text
    l3 = _row(db, "leads", "lead_id", "L3")
    assert l3["assigned_to"] == "e1" and l3["team_lead_id"] == "tl_a"


def test_team_lead_cannot_assign_outside_team(lead_a, db):
    r = lead_a.post("/api/leads/bulk-manage", json={"lead_ids": ["L3"], "assigned_to": "e2"})
    assert r.status_code == 403
    assert _row(db, "leads", "lead_id", "L3")["assigned_to"] is None


def test_team_lead_cannot_touch_other_teams_leads(lead_a, db):
    r = lead_a.post("/api/leads/bulk-manage", json={"lead_ids": ["L5"], "assigned_to": "e1"})
    assert r.status_code == 200
    assert r.json()["skipped"] == [{"lead_id": "L5", "reason": "not_in_team"}]
    assert _row(db, "leads", "lead_id", "L5")["assigned_to"] == "e2"


def test_team_lead_cannot_move_leads_between_team_leads(lead_a):
    r = lead_a.post("/api/leads/bulk-manage", json={"lead_ids": ["L3"], "team_lead_id": "tl_b"})
    assert r.status_code == 403


def test_team_lead_can_return_lead_to_admin_pool(lead_a, db):
    r = lead_a.post("/api/leads/bulk-manage", json={"lead_ids": ["L4"], "return_to_admin": True})
    assert r.status_code == 200, r.text
    l4 = _row(db, "leads", "lead_id", "L4")
    assert l4["assigned_to"] is None and l4["team_lead_id"] is None and l4["stage"] == "new"


def test_team_lead_cannot_delete_leads(lead_a, db):
    assert lead_a.delete("/api/leads/L3").status_code == 403
    assert lead_a.post("/api/leads/bulk-delete", json={"lead_ids": ["L3"]}).status_code == 403
    assert any(l["lead_id"] == "L3" for l in db.table("leads"))


def test_failed_saves_are_reported_not_hidden(admin, db, monkeypatch):
    monkeypatch.setattr(core, "sb_patch_filter", lambda *a, **k: "leads: 500 boom")
    r = admin.post("/api/leads/bulk-manage", json={"lead_ids": ["L1", "L2"], "assigned_to": "e1"})
    body = r.json()
    assert body["status"] == "partial"
    assert body["failed_count"] == 2 and body["updated_count"] == 0
    assert {s["reason"] for s in body["skipped"]} == {"save_failed"}
    assert _row(db, "leads", "lead_id", "L1")["assigned_to"] is None


def test_patch_rejects_unknown_employee_and_unassigns_on_empty(admin, db):
    assert admin.patch("/api/leads/L1", json={"assigned_to": "ghost"}).status_code == 400
    r = admin.patch("/api/leads/L4", json={"assigned_to": ""})
    assert r.status_code == 200, r.text
    assert _row(db, "leads", "lead_id", "L4")["assigned_to"] is None


def test_admin_auto_distributes_by_region(admin, db):
    r = admin.post("/api/leads/assign-queue/auto")
    assert r.status_code == 200, r.text
    assert _row(db, "leads", "lead_id", "L1")["team_lead_id"] == "tl_a"  # Virar
    assert _row(db, "leads", "lead_id", "L2")["team_lead_id"] == "tl_b"  # Vasai


def test_region_routing_reads_housing_payload_and_skips_inactive_regions(admin, db):
    # Area only inside the Housing payload (no location typed).
    db.table("leads").append(_lead("H1", raw_payload={"locality_name": "Virar East"}))
    r = admin.get("/api/leads/assign-workspace", params={"region": "Virar", "team": "admin_pool"}).json()
    assert {l["lead_id"] for l in r["leads"]} == {"L1", "H1"}
    # Deactivate Vasai: Vasai leads are no longer routed to its Team Lead by area.
    admin.delete("/api/regions/reg_vasai")
    admin.post("/api/leads/assign-queue/auto")
    assert _row(db, "leads", "lead_id", "H1")["team_lead_id"] == "tl_a"
    assert _row(db, "leads", "lead_id", "L1")["team_lead_id"] == "tl_a"
    assert _row(db, "leads", "lead_id", "L2")["team_lead_id"] in ("tl_a", "tl_b")  # round-robin now


def test_team_lead_snapshot_refreshes_right_after_assign(lead_a, db):
    before = lead_a.get("/api/stats/assignment").json()["team"]["teams"]["tl_a"]
    assert before == {"total": 2, "pool": 1, "assigned": 1}
    lead_a.post("/api/leads/bulk-manage", json={"lead_ids": ["L3"], "assigned_to": "e1"})
    after = lead_a.get("/api/stats/assignment").json()["team"]["teams"]["tl_a"]
    assert after == {"total": 2, "pool": 0, "assigned": 2}


# ---- emergency transfer ---------------------------------------------------

def test_transfer_keeps_employees_working(admin, db):
    r = admin.post("/api/teams/transfer", json={"from_team_lead_id": "tl_a", "to_team_lead_id": "tl_b"})
    assert r.status_code == 200, r.text
    assert r.json()["moved_count"] == 2
    l3, l4 = _row(db, "leads", "lead_id", "L3"), _row(db, "leads", "lead_id", "L4")
    assert l3["team_lead_id"] == l4["team_lead_id"] == "tl_b"
    assert l4["assigned_to"] == "e1"


def test_transfer_return_to_pool_and_move_members(admin, db):
    r = admin.post("/api/teams/transfer", json={
        "from_team_lead_id": "tl_a", "to_team_lead_id": "tl_b",
        "mode": "return_to_pool", "move_members": True,
    })
    assert r.status_code == 200, r.text
    assert r.json()["members_moved"] == ["e1"]
    l4 = _row(db, "leads", "lead_id", "L4")
    assert l4["assigned_to"] is None and l4["team_lead_id"] == "tl_b"
    assert _row(db, "employees", "employee_id", "e1")["team_lead_id"] == "tl_b"


def test_overview_accounts_for_every_lead(admin):
    body = admin.get("/api/teams/overview").json()
    teams = {t["employee_id"]: t for t in body["team_leads"]}
    total = body["admin_pool"] + body["no_team_assigned"] + sum(t["leads"]["total"] for t in teams.values())
    assert total == body["pipeline_total"] == len(LEADS)
    assert [e["employee_id"] for e in body["unteamed_employees"]] == ["e3"]


# ---- people ---------------------------------------------------------------

def test_deleting_employee_returns_their_leads_to_team_pool(admin, db):
    r = admin.delete("/api/employees/e1")
    assert r.status_code == 200, r.text
    assert r.json()["leads_returned"] == 1
    l4 = _row(db, "leads", "lead_id", "L4")
    assert l4["assigned_to"] is None and l4["team_lead_id"] == "tl_a"


def test_team_lead_with_leads_cannot_be_removed(admin):
    assert admin.patch("/api/employees/tl_a", json={"role": "telecaller"}).status_code == 409
    assert admin.patch("/api/employees/tl_a", json={"active": False}).status_code == 409
    assert admin.delete("/api/employees/tl_a").status_code == 409


def test_moving_employee_moves_their_leads(admin, db):
    r = admin.post("/api/teams/members", json={"employee_ids": ["e3"], "team_lead_id": "tl_a"})
    assert r.status_code == 200, r.text
    assert r.json()["leads_moved"] == 1
    assert _row(db, "leads", "lead_id", "L6")["team_lead_id"] == "tl_a"


def test_admin_creates_team_lead_with_regions(admin, db, monkeypatch):
    monkeypatch.setattr(core, "_verify_login_password_saved", lambda *a, **k: None)
    r = admin.post("/api/employees", json={
        "name": "Nagpur Lead", "email": "nagpur@test.com", "role": "team_lead", "password": "1234",
        "location": "Nagpur", "regions": ["nagpur"],
    })
    assert r.status_code == 200, r.text
    assert r.json()["regions"] == ["Nagpur"] and r.json()["team_lead_id"] is None
    bad_region = admin.post("/api/employees", json={
        "name": "Y", "email": "y@test.com", "role": "team_lead", "password": "1234",
        "location": "Pune", "regions": ["Pune"],
    })
    assert bad_region.status_code == 400


def test_team_lead_can_only_add_workers_to_own_team(lead_a, db, monkeypatch):
    monkeypatch.setattr(core, "_verify_login_password_saved", lambda *a, **k: None)
    bad = lead_a.post("/api/employees", json={
        "name": "X", "email": "x@test.com", "role": "team_lead", "password": "1234", "location": "Virar",
    })
    assert bad.status_code == 403
    ok = lead_a.post("/api/employees", json={
        "name": "New", "email": "new@test.com", "role": "telecaller", "password": "1234",
        "location": "Virar", "team_lead_id": "tl_b",
    })
    assert ok.status_code == 200, ok.text
    assert ok.json()["team_lead_id"] == "tl_a"  # forced into their own team


# ---- exports --------------------------------------------------------------

def test_export_summary_counts_every_lead_per_day(admin, db):
    db.table("leads").append(_lead("BROKER", day=4, stage="broker", lead_type="brokerage"))
    db.table("leads").append(_lead("NEG", day=4, status="negative"))
    db.table("leads").append(_lead("SEPT", created_at="2026-09-30T10:00:00+00:00"))
    body = admin.get("/api/exports/leads/summary", params={"month": "2026-10"}).json()
    assert body["total"] == len(LEADS) + 2
    assert len(body["days"]) == 31
    by_day = {d["date"]: d["count"] for d in body["days"]}
    assert by_day["2026-10-01"] == 2 and by_day["2026-10-04"] == 3
    assert sum(by_day.values()) == body["total"]
    assert body["by_status"]["negative"] == 1


def test_team_lead_export_is_scoped(lead_a):
    body = lead_a.get("/api/exports/leads/summary", params={"month": "2026-10"}).json()
    assert body["total"] == 2 and body["scope"] == "team:tl_a"


def test_export_downloads_and_is_archived(admin, db, monkeypatch):
    stored = {}
    monkeypatch.setattr(core, "sb_storage_upload", lambda bucket, path, content, ctype, upsert=True: stored.update({path: content}) or True)
    r = admin.post("/api/exports/leads", json={"month": "2026-10"})
    assert r.status_code == 200, r.text
    assert r.headers["x-export-stored"] == "true"
    assert int(r.headers["x-export-rows"]) == len(LEADS)
    wb = openpyxl.load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Leads", "Daily Summary", "By Team"]
    assert wb["Leads"].max_row == len(LEADS) + 1
    record = db.table("lead_exports")[0]
    assert record["stored"] is True and record["file_path"] in stored
    assert admin.get(f"/api/exports/leads/history").json()["exports"][0]["export_id"] == record["export_id"]


def test_export_still_downloads_when_storage_fails(admin, monkeypatch):
    monkeypatch.setattr(core, "sb_storage_upload", lambda *a, **k: False)
    monkeypatch.setattr(core, "sb_storage_ensure_bucket", lambda *a, **k: False)
    r = admin.post("/api/exports/leads", json={"month": "2026-10"})
    assert r.status_code == 200
    assert r.headers["x-export-stored"] == "false"
    assert len(r.content) > 0


def test_export_month_is_validated(admin):
    assert admin.get("/api/exports/leads/summary", params={"month": "2026-13"}).status_code == 400


# ---- security + lead-loss fixes ------------------------------------------

def test_non_admin_cannot_promote_themselves(db):
    _login(_user("telecaller", "e1"))
    r = TestClient(core.app).post("/api/auth/set-role", json={"role": "admin"})
    assert r.status_code == 403


def test_acting_as_header_only_for_admins(db):
    worker = core._finish_request_user(_user("telecaller", "e1").model_dump(), _Req({"X-Acting-As": "e2"}))
    assert worker.acting_as_employee_id == "e1"
    boss = core._finish_request_user(_user("admin", email="admin@test.com").model_dump(), _Req({"X-Acting-As": "e2"}))
    assert boss.acting_as_employee_id == "e2"


def test_former_manager_email_no_longer_bypasses_roles():
    rohit = _user("team_lead", email="rohitsingh241993@gmail.com")
    with pytest.raises(core.HTTPException):
        core.ensure_roles(rohit, ["admin"])


def test_full_table_read_raises_instead_of_truncating(monkeypatch):
    calls = []

    def flaky(table, params):
        calls.append(params)
        if params["offset"] == "0":
            return [{"lead_id": str(i)} for i in range(1000)]
        raise core.SupabaseReadError("leads: HTTP 500")

    monkeypatch.setattr(core, "sb_select_strict", flaky)
    monkeypatch.setattr(core.time, "sleep", lambda s: None)
    with pytest.raises(core.SupabaseReadError):
        core.sb_select_all("leads", {"select": "lead_id", "order": "created_at.desc"})
    assert calls[0]["order"] == "created_at.desc,lead_id.desc"
